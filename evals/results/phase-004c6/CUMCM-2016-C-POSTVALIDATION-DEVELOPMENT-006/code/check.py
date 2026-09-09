"""Independent numerical checker: scalar Hermite construction and bisection.

Does not import producer, SciPy interpolation/optimization, or shared solution helpers.
Only writes --output. It does not read final labels or evaluate a held-out test.
"""

import argparse
import bisect
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook


def average(a):
    return sum(a) / len(a)


def coefficients(points, cubic):
    x = [p[0] for p in points]
    y = [p[1] for p in points]
    h = [b - a for a, b in zip(x[:-1], x[1:], strict=True)]
    d = [(b - a) / z for a, b, z in zip(y[:-1], y[1:], h, strict=True)]
    if not cubic:
        return [[0.0, 0.0, z, a] for z, a in zip(d, y[:-1], strict=True)]
    slopes = [0.0] * len(x)
    for i in range(1, len(x) - 1):
        if d[i - 1] * d[i] > 0:
            w1 = 2 * h[i] + h[i - 1]
            w2 = h[i] + 2 * h[i - 1]
            slopes[i] = (w1 + w2) / (w1 / d[i - 1] + w2 / d[i])

    def edge(h0, h1, d0, d1):
        m = ((2 * h0 + h1) * d0 - h0 * d1) / (h0 + h1)
        if m * d0 <= 0:
            return 0.0
        if d0 * d1 < 0 and abs(m) > 3 * abs(d0):
            return 3 * d0
        return m

    slopes[0] = edge(h[0], h[1], d[0], d[1])
    slopes[-1] = edge(h[-1], h[-2], d[-1], d[-2])
    return [
        [
            (slopes[i] + slopes[i + 1] - 2 * d[i]) / h[i] ** 2,
            (3 * d[i] - 2 * slopes[i] - slopes[i + 1]) / h[i],
            slopes[i],
            y[i],
        ]
        for i in range(len(h))
    ]


def evaluate(x, c, t):
    k = min(max(0, bisect.bisect_right(x, t) - 1), len(c) - 1)
    v = t - x[k]
    a, b, d, e = c[k]
    return ((a * v + b) * v + d) * v + e


def inverse(fun, end, voltage):
    if abs(fun(end) - voltage) < 1e-9:
        return end
    # A sum of scaled cubic pieces can turn inside a merged knot interval.
    # Split at every derivative root before choosing the last downward bracket.
    knots = fun.knots
    boundaries = list(knots)
    for left, right in zip(knots[:-1], knots[1:], strict=True):
        width = right - left
        eps = min(width / 1000, 1e-6)
        d0 = fun.derivative(left + eps)
        dm = fun.derivative((left + right) / 2)
        d1 = fun.derivative(right - eps)
        # Interpolate the quadratic derivative at interior coordinates, avoiding
        # the one-sided slope ambiguity of piecewise-linear knot derivatives.
        z0, zm, z1 = eps, width / 2, width - eps
        a = ((d1 - dm) / (z1 - zm) - (dm - d0) / (zm - z0)) / (z1 - z0)
        b = (dm - d0) / (zm - z0) - a * (zm + z0)
        c = d0 - a * z0 * z0 - b * z0
        roots = []
        if abs(a) < 1e-16:
            if abs(b) > 1e-16:
                roots = [-c / b]
        elif b * b - 4 * a * c >= 0:
            disc = math.sqrt(b * b - 4 * a * c)
            roots = [(-b - disc) / (2 * a), (-b + disc) / (2 * a)]
        boundaries.extend(left + z for z in roots if 0 < z < width)
    points = sorted(set(boundaries))
    brackets = [
        (a, b)
        for a, b in zip(points[:-1], points[1:], strict=True)
        if fun(a) >= voltage >= fun(b) and fun(a) > fun(b)
    ]
    if not brackets:
        raise ValueError("CHECK_MODEL_DOWNWARD_CROSSING_MISSING")
    lo, hi = brackets[-1]
    for _ in range(75):
        mid = (lo + hi) / 2
        if fun(mid) > voltage:
            lo = mid
        else:
            hi = mid
    return (lo + hi) / 2


def sample_time(points, v):
    for i in range(len(points) - 2, -1, -1):
        t0, u0 = points[i]
        t1, u1 = points[i + 1]
        if u0 >= v >= u1:
            return t1 if u0 == u1 else t0 + (u0 - v) * (t1 - t0) / (u0 - u1)
    raise ValueError("CHECK_RAW_CROSSING_MISSING")


def error(a, b):
    return average([abs(x - y) / y for x, y in zip(a, b, strict=True)])


def difference(a, b):
    return max([abs(x - y) for x, y in zip(a, b, strict=True)] + [0.0])


def regression(reference, target, affine):
    x = reference[: len(target)]
    if affine:
        xm = average(x)
        ym = average(target)
        b = sum((u - xm) * (v - ym) for u, v in zip(x, target, strict=True)) / sum(
            (u - xm) ** 2 for u in x
        )
        a = ym - b * xm
    else:
        a = 0.0
        b = sum(u * v for u, v in zip(x, target, strict=True)) / sum(u * u for u in x)
    return a, b, [a + b * u for u in reference]


def residual(value, tolerance=1e-7):
    return {"value": float(value), "relation": "LE", "limit": 0.0, "tolerance": tolerance}


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--case-root", type=Path, required=True)
    p.add_argument("--run-id", required=True)
    p.add_argument("--output", type=Path, required=True)
    a = p.parse_args()
    output_path = a.case_root / "runs" / a.run_id / "output.json"
    out = json.loads(output_path.read_text())
    cubic = out["candidate_id"] != "BASELINE"
    wb = load_workbook(a.case_root / "data/raw/appendix.xlsx", read_only=True, data_only=True)
    rows = list(wb["附件1"].values)
    curves = {
        i: [(float(r[0]), float(r[j])) for r in rows[2:] if isinstance(r[j], (int, float))]
        for j, i in enumerate(range(20, 101, 10), 1)
    }
    polys = {}
    durations = {}
    coefficient_residual = 0.0
    model_time_residual = 0.0
    q1_metrics = {}
    q1_mre = []
    grid = [9 + k * 0.005 for k in range(231)]
    for i, points in curves.items():
        n = len(points)
        count = 33 if cubic else 17
        indices = sorted(
            set(list(range(min(31, n))) + [round(j * (n - 1) / (count - 1)) for j in range(count)])
        )
        knots = [points[j] for j in indices]
        x = [v[0] for v in knots]
        c = coefficients(knots, cubic)
        durations[i] = points[-1][0]
        record = out["Q1"]["curve_models"][str(i)]
        coefficient_residual = max(coefficient_residual, difference(x, record["time_knots_min"]))
        coefficient_residual = max(
            coefficient_residual,
            max(
                abs(c[j][k] - record["coefficients_descending"][k][j])
                for j in range(len(c))
                for k in range(4)
            ),
        )

        def fun(t, x=x, c=c):
            return evaluate(x, c, t)

        def derivative(t, x=x, c=c):
            j = min(max(0, bisect.bisect_right(x, t) - 1), len(c) - 1)
            z = t - x[j]
            a, b, d, _ = c[j]
            return (3 * a * z + 2 * b) * z + d

        fun.knots = x
        fun.derivative = derivative
        polys[i] = fun
        values = [inverse(fun, durations[i], u) for u in grid]
        actual = [sample_time(points, u) for u in grid]
        m = error(values, actual)
        q1_mre.append(m)
        assert abs(m - out["Q1"]["per_current_MRE"][str(i)]) < 1e-9
        q1_metrics[f"q1_MRE_{i}A"] = m
        original = next(r for r in out["Q1"]["mre_evaluation"] if r["current_A"] == i)
        model_time_residual = max(
            model_time_residual, difference(values, original["model_elapsed_at_voltage_min"])
        )
        if 30 <= i <= 70:
            q1_metrics[f"q1_remaining_{i}A_min"] = durations[i] - inverse(fun, durations[i], 9.8)
    q1_metrics["q1_mean_fitted_MRE"] = average(q1_mre)

    def surface(current, omit=None):
        keys = [i for i in polys if i != omit]
        lo = max(i for i in keys if i < current)
        hi = min(i for i in keys if i > current)
        w = math.log(current / lo) / math.log(hi / lo) if cubic else (current - lo) / (hi - lo)
        T = (
            math.exp((1 - w) * math.log(durations[lo]) + w * math.log(durations[hi]))
            if cubic
            else (1 - w) * durations[lo] + w * durations[hi]
        )

        def combined(t):
            return (1 - w) * polys[lo](t * durations[lo] / T) + w * polys[hi](t * durations[hi] / T)

        combined.knots = sorted(
            set(
                [t * T / durations[lo] for t in polys[lo].knots]
                + [t * T / durations[hi] for t in polys[hi].knots]
            )
        )
        combined.derivative = lambda t: (
            (1 - w) * polys[lo].derivative(t * durations[lo] / T) * durations[lo] / T
            + w * polys[hi].derivative(t * durations[hi] / T) * durations[hi] / T
        )
        return combined, T

    q2_errors = []
    q2_vector_residual = 0.0
    for i in range(30, 91, 10):
        f, T = surface(i, i)
        v = [inverse(f, T, u) for u in grid]
        observed = [sample_time(curves[i], u) for u in grid]
        q2_errors.append(error(v, observed))
        row = next(r for r in out["Q2"]["leave_one_current_out"] if r["current_A"] == i)
        q2_vector_residual = max(q2_vector_residual, difference(v, row["predicted_elapsed_min"]))
    f55, T55 = surface(55)
    for j, row in enumerate(out["Q2"]["table55"]):
        t = j * T55 / 100
        q2_vector_residual = max(
            q2_vector_residual, abs(t - row["time_min"]), abs(f55(t) - row["voltage_V"])
        )
    sheet = wb["附件2"]
    voltage = []
    new = []
    old = []
    state2full = []
    state2 = []
    target = []
    for r in range(3, 304):
        u = sheet.cell(r, 1).value
        if not isinstance(u, (int, float)):
            continue
        voltage.append(float(u))
        new.append(float(sheet.cell(r, 2).value))
        old.append(float(sheet.cell(r, 3).value))
        state2full.append(float(sheet.cell(r, 4).value))
        if u >= 9.765 - 1e-10:
            state2.append(float(sheet.cell(r, 4).value))
            target.append(float(sheet.cell(r, 5).value))
    wb.close()
    ref = [(x + y + z) / 3 for x, y, z in zip(new, old, state2full, strict=True)]
    a3, b3, pred3 = regression(ref, target, cubic)
    a2, b2, pred2 = regression([(x + y) / 2 for x, y in zip(new, old, strict=True)], state2, cubic)
    _, _, dev = regression(new, old[: len(target)], cubic)
    val = [j for j, u in enumerate(voltage) if u < 9.765 - 1e-10 and u >= 9.030 - 1e-10]
    assert all(j < len(old) for j in val)
    remaining = pred3[-1] - target[-1]
    q3_vector_residual = max(
        difference(pred3, out["Q3"]["predicted_elapsed_min"]),
        difference(pred2, out["Q3"]["state2_development_prediction_min"]),
        difference(dev, out["Q3"]["state1_validation_prediction_min"]),
        abs(a3 - out["Q3"]["a_min"]),
        abs(b3 - out["Q3"]["b"]),
    )
    for target2, record in zip(
        ([v + 1 for v in target], [v - 1 for v in target], target[:-10]),
        out["robustness_evidence"]["perturbations"],
        strict=True,
    ):
        _, _, pred = regression(ref, target2, cubic)
        q3_vector_residual = max(q3_vector_residual, abs(pred[-1] - target[-1] - record["result"]))
    for reference, record in zip(
        (new, old, state2full), out["Q3"]["reference_specific_remaining_min"], strict=True
    ):
        _, _, pred = regression(reference, target, cubic)
        q3_vector_residual = max(q3_vector_residual, abs(pred[-1] - target[-1] - record))
    q2_metrics = {"q2_loco_MRE": average(q2_errors)}
    history = []
    matrix = [new, old, state2full]
    for state in (1, 2):
        reference = [average(row) for row in zip(*matrix[:state], strict=True)]
        for cut in (9.95, 9.85, 9.765):
            size = sum(u >= cut - 1e-10 for u in voltage)
            prefix = matrix[state][:size]
            _, _, prediction = regression(reference, prefix, cubic)
            history.append(
                {
                    "sample_id": f"S{state}-CUT-{cut}",
                    "target": "remaining_time",
                    "unit": "min",
                    "origin": state * 10000 + prefix[-1],
                    "predicted_end_time": state * 10000 + prediction[-1],
                    "observed_end_time": state * 10000 + matrix[state][-1],
                }
            )
    for expected, observed in zip(history, out["Q3"]["rolling_origin_development"], strict=True):
        assert set(expected) == set(observed)
        assert all(expected[k] == observed[k] for k in expected if k != "predicted_end_time")
        q3_vector_residual = max(
            q3_vector_residual, abs(expected["predicted_end_time"] - observed["predicted_end_time"])
        )
    q3_remaining_mre = average(
        [
            abs(r["predicted_end_time"] - r["observed_end_time"])
            / (r["observed_end_time"] - r["origin"])
            for r in history
        ]
    )
    q3_metrics = {"q3_remaining_MRE": q3_remaining_mre, "q3_remaining_min": remaining}
    for current in range(30, 71, 10):
        assert (
            abs(
                q1_metrics[f"q1_remaining_{current}A_min"]
                - out["Q1"]["remaining_at_9_8V_min"][str(current)]
            )
            < 1e-7
        )
    q1_metrics = {"q1_mean_fitted_MRE": q1_metrics["q1_mean_fitted_MRE"]}
    results = {}
    for rid, metrics, recalc in (
        (
            "REQ-Q1",
            q1_metrics,
            {
                "all_coefficients": residual(coefficient_residual, 1e-9),
                "all_2079_inverse_times_min": residual(model_time_residual),
            },
        ),
        ("REQ-Q2", q2_metrics, {"all_LOCO_and_55A_vectors": residual(q2_vector_residual)}),
        (
            "REQ-Q3",
            q3_metrics,
            {"all_aging_vectors_and_perturbations_min": residual(q3_vector_residual)},
        ),
    ):
        # Report independently recalculated values; machine gate compares each metric too.
        recalc["metric_max_difference"] = residual(
            max(abs(v - out["final_metrics"][k]) for k, v in metrics.items()), 1e-9
        )
        constraint = residual(max(0.0, -remaining) if rid == "REQ-Q3" else 0.0)
        results[rid] = {
            "metric_values": metrics,
            "recalculation_residuals": recalc,
            "feasible": remaining >= 0 if rid == "REQ-Q3" else True,
            "constraint_residuals": {
                "nonnegative_remaining_time_min"
                if rid == "REQ-Q3"
                else "bounded_interpolation_domain": constraint
            },
            "domain_interpretation": (
                "Numerical/domain consistency onl"
                "y; no validation of state3 futur"
                "e truth or aging invariance."
            ),
        }
    plan = json.loads((a.case_root / "experiments/experiment_plan.json").read_text())["content"]
    samples = out["metric_samples"]
    for metric in ("q1_mean_fitted_MRE", "q2_loco_MRE"):
        actual_values = []
        for row in samples[metric]:
            if metric == "q1_mean_fitted_MRE":
                current, j = map(int, row["sample_id"][1:].split("-U"))
                prediction = inverse(polys[current], durations[current], grid[j])
            else:
                current, j = map(int, row["sample_id"][5:].split("-U"))
                fun, duration = surface(current, current)
                prediction = inverse(fun, duration, grid[j])
            truth = sample_time(curves[current], grid[j])
            assert row["target"] == "elapsed_time_at_voltage" and row["unit"] == "min"
            assert abs(row["truth"] - truth) < 1e-9 and abs(row["prediction"] - prediction) < 1e-7
            actual_values.append(abs(prediction - truth) / truth)
        assert len(actual_values) == (2079 if metric == "q1_mean_fitted_MRE" else 1617)
        assert abs(average(actual_values) - out["final_metrics"][metric]) < 1e-9
    for supplied, expected in zip(samples["q3_remaining_MRE"], history, strict=True):
        assert all(supplied[k] == expected[k] for k in expected if k != "predicted_end_time")
        assert abs(supplied["predicted_end_time"] - expected["predicted_end_time"]) < 1e-7
    assert samples["q3_remaining_min"] == [
        {
            "sample_id": "S3-LAST",
            "target": "future_remaining_time",
            "unit": "min",
            "value": out["Q3"]["predicted_remaining_min"],
        }
    ]
    assert abs(samples["q3_remaining_min"][0]["value"] - remaining) < 1e-7
    evidence2 = out["scientific_evidence"]["REQ-Q2"]["prediction_evidence"]
    evidence3 = out["scientific_evidence"]["REQ-Q3"]["prediction_evidence"]
    assert abs(evidence2["predictions"][0]["value"] - T55) < 1e-7
    assert abs(evidence3["predictions"][0]["value"] - remaining) < 1e-7
    assert (
        evidence3["uncertainty"]["variant_values"] == out["Q3"]["reference_specific_remaining_min"]
    )
    results["REQ-Q2"]["prediction_evidence"] = evidence2
    results["REQ-Q3"]["prediction_evidence"] = evidence3
    payload = {
        "run_id": a.run_id,
        "output_sha256": hashlib.sha256(output_path.read_bytes()).hexdigest(),
        "requirements": results,
        "metric_definitions": plan["metric_definitions"],
        "metric_samples": samples,
        "temporal_lineage": plan["temporal_design"],
        "independence": (
            "Manual scalar PCHIP coefficients"
            ", bisection, closed-form regress"
            "ion; no producer imports; same m"
            "odel assumptions are not indepen"
            "dently established."
        ),
        "final_test_access": False,
    }
    a.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n")


if __name__ == "__main__":
    main()
