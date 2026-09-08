"""Independent numerical checker: scalar Hermite construction and bisection.

Does not import producer, SciPy interpolation/optimization, or shared solution helpers.
Only writes --output. It does not read final labels or evaluate a held-out test.
"""

import argparse
import bisect
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook


def average(a):
    return sum(a) / len(a)


def coefficients(points, cubic):
    x = [p[0] for p in points]
    y = [p[1] for p in points]
    h = [b - a for a, b in zip(x[:-1], x[1:], strict=True)]
    d = [(b - a) / z for a, b, z in zip(y[:-1], y[1:], h, strict=True)]
    if not cubic:
        return [[0.0, 0.0, z, a] for z, a in zip(d, y[:-1], strict=True)]
    slopes = [0.0] * len(x)
    for i in range(1, len(x) - 1):
        if d[i - 1] * d[i] > 0:
            w1 = 2 * h[i] + h[i - 1]
            w2 = h[i] + 2 * h[i - 1]
            slopes[i] = (w1 + w2) / (w1 / d[i - 1] + w2 / d[i])

    def edge(h0, h1, d0, d1):
        m = ((2 * h0 + h1) * d0 - h0 * d1) / (h0 + h1)
        if m * d0 <= 0:
            return 0.0
        if d0 * d1 < 0 and abs(m) > 3 * abs(d0):
            return 3 * d0
        return m

    slopes[0] = edge(h[0], h[1], d[0], d[1])
    slopes[-1] = edge(h[-1], h[-2], d[-1], d[-2])
    return [
        [
            (slopes[i] + slopes[i + 1] - 2 * d[i]) / h[i] ** 2,
            (3 * d[i] - 2 * slopes[i] - slopes[i + 1]) / h[i],
            slopes[i],
            y[i],
        ]
        for i in range(len(h))
    ]


def evaluate(x, c, t):
    k = min(max(0, bisect.bisect_right(x, t) - 1), len(c) - 1)
    v = t - x[k]
    a, b, d, e = c[k]
    return ((a * v + b) * v + d) * v + e


def inverse(fun, end, voltage):
    if abs(fun(end) - voltage) < 1e-9:
        return end
    lo = 0.0
    hi = end
    for _ in range(75):
        mid = (lo + hi) / 2
        if fun(mid) > voltage:
            lo = mid
        else:
            hi = mid
    return (lo + hi) / 2


def sample_time(points, v):
    for i in range(len(points) - 2, -1, -1):
        t0, u0 = points[i]
        t1, u1 = points[i + 1]
        if u0 >= v >= u1:
            return t1 if u0 == u1 else t0 + (u0 - v) * (t1 - t0) / (u0 - u1)
    raise ValueError("CHECK_RAW_CROSSING_MISSING")


def error(a, b):
    return average([abs(x - y) / y for x, y in zip(a, b, strict=True)])


def difference(a, b):
    return max([abs(x - y) for x, y in zip(a, b, strict=True)] + [0.0])


def regression(reference, target, affine):
    x = reference[: len(target)]
    if affine:
        xm = average(x)
        ym = average(target)
        b = sum((u - xm) * (v - ym) for u, v in zip(x, target, strict=True)) / sum(
            (u - xm) ** 2 for u in x
        )
        a = ym - b * xm
    else:
        a = 0.0
        b = sum(u * v for u, v in zip(x, target, strict=True)) / sum(u * u for u in x)
    return a, b, [a + b * u for u in reference]


def residual(value, tolerance=1e-7):
    return {"value": float(value), "relation": "LE", "limit": 0.0, "tolerance": tolerance}


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--case-root", type=Path, required=True)
    p.add_argument("--run-id", required=True)
    p.add_argument("--output", type=Path, required=True)
    a = p.parse_args()
    output_path = a.case_root / "runs" / a.run_id / "output.json"
    out = json.loads(output_path.read_text())
    cubic = out["candidate_id"] != "BASELINE"
    wb = load_workbook(a.case_root / "data/raw/appendix.xlsx", read_only=True, data_only=True)
    rows = list(wb["附件1"].values)
    curves = {
        i: [(float(r[0]), float(r[j])) for r in rows[2:] if isinstance(r[j], (int, float))]
        for j, i in enumerate(range(20, 101, 10), 1)
    }
    polys = {}
    durations = {}
    coefficient_residual = 0.0
    model_time_residual = 0.0
    q1_metrics = {}
    q1_mre = []
    grid = [9 + k * 0.005 for k in range(231)]
    for i, points in curves.items():
        n = len(points)
        count = 33 if cubic else 17
        indices = sorted(
            set(list(range(min(31, n))) + [round(j * (n - 1) / (count - 1)) for j in range(count)])
        )
        knots = [points[j] for j in indices]
        x = [v[0] for v in knots]
        c = coefficients(knots, cubic)
        durations[i] = points[-1][0]
        record = out["Q1"]["curve_models"][str(i)]
        coefficient_residual = max(coefficient_residual, difference(x, record["time_knots_min"]))
        coefficient_residual = max(
            coefficient_residual,
            max(
                abs(c[j][k] - record["coefficients_descending"][k][j])
                for j in range(len(c))
                for k in range(4)
            ),
        )

        def fun(t, x=x, c=c):
            return evaluate(x, c, t)

        polys[i] = fun
        values = [inverse(fun, durations[i], u) for u in grid]
        actual = [sample_time(points, u) for u in grid]
        m = error(values, actual)
        q1_mre.append(m)
        q1_metrics[f"q1_MRE_{i}A"] = m
        original = next(r for r in out["Q1"]["mre_evaluation"] if r["current_A"] == i)
        model_time_residual = max(
            model_time_residual, difference(values, original["model_elapsed_at_voltage_min"])
        )
        if 30 <= i <= 70:
            q1_metrics[f"q1_remaining_{i}A_min"] = durations[i] - inverse(fun, durations[i], 9.8)
    q1_metrics["q1_mean_fitted_MRE"] = average(q1_mre)

    def surface(current, omit=None):
        keys = [i for i in polys if i != omit]
        lo = max(i for i in keys if i < current)
        hi = min(i for i in keys if i > current)
        w = math.log(current / lo) / math.log(hi / lo) if cubic else (current - lo) / (hi - lo)
        T = (
            math.exp((1 - w) * math.log(durations[lo]) + w * math.log(durations[hi]))
            if cubic
            else (1 - w) * durations[lo] + w * durations[hi]
        )
        return (
            lambda t: (
                (1 - w) * polys[lo](t * durations[lo] / T) + w * polys[hi](t * durations[hi] / T)
            ),
            T,
        )

    q2_errors = []
    q2_vector_residual = 0.0
    for i in range(30, 91, 10):
        f, T = surface(i, i)
        v = [inverse(f, T, u) for u in grid]
        observed = [sample_time(curves[i], u) for u in grid]
        q2_errors.append(error(v, observed))
        row = next(r for r in out["Q2"]["leave_one_current_out"] if r["current_A"] == i)
        q2_vector_residual = max(q2_vector_residual, difference(v, row["predicted_elapsed_min"]))
    f55, T55 = surface(55)
    for j, row in enumerate(out["Q2"]["table55"]):
        t = j * T55 / 100
        q2_vector_residual = max(
            q2_vector_residual, abs(t - row["time_min"]), abs(f55(t) - row["voltage_V"])
        )
    sheet = wb["附件2"]
    voltage = []
    new = []
    old = []
    state2full = []
    state2 = []
    target = []
    for r in range(3, 304):
        u = sheet.cell(r, 1).value
        if not isinstance(u, (int, float)):
            continue
        voltage.append(float(u))
        new.append(float(sheet.cell(r, 2).value))
        old.append(float(sheet.cell(r, 3).value))
        state2full.append(float(sheet.cell(r, 4).value))
        if u >= 9.765 - 1e-10:
            state2.append(float(sheet.cell(r, 4).value))
            target.append(float(sheet.cell(r, 5).value))
    wb.close()
    ref = [(x + y + z) / 3 for x, y, z in zip(new, old, state2full, strict=True)]
    a3, b3, pred3 = regression(ref, target, cubic)
    a2, b2, pred2 = regression([(x + y) / 2 for x, y in zip(new, old, strict=True)], state2, cubic)
    _, _, dev = regression(new, old[: len(target)], cubic)
    val = [j for j, u in enumerate(voltage) if u < 9.765 - 1e-10 and u >= 9.030 - 1e-10]
    q3_mre = error([dev[j] for j in val], [old[j] for j in val])
    remaining = pred3[-1] - target[-1]
    q3_vector_residual = max(
        difference(pred3, out["Q3"]["predicted_elapsed_min"]),
        difference(pred2, out["Q3"]["state2_development_prediction_min"]),
        difference(dev, out["Q3"]["state1_validation_prediction_min"]),
        abs(a3 - out["Q3"]["a_min"]),
        abs(b3 - out["Q3"]["b"]),
    )
    for target2, record in zip(
        ([v + 1 for v in target], [v - 1 for v in target], target[:-10]),
        out["robustness_evidence"]["perturbations"],
        strict=True,
    ):
        _, _, pred = regression(ref, target2, cubic)
        q3_vector_residual = max(q3_vector_residual, abs(pred[-1] - target[-1] - record["result"]))
    for reference, record in zip(
        (new, old, state2full), out["Q3"]["reference_specific_remaining_min"], strict=True
    ):
        _, _, pred = regression(reference, target, cubic)
        q3_vector_residual = max(q3_vector_residual, abs(pred[-1] - target[-1] - record))
    q2_metrics = {"q2_loco_MRE": average(q2_errors)}
    q3_metrics = {"q3_validation_tail_MRE": q3_mre, "q3_remaining_min": remaining}
    results = {}
    for rid, metrics, recalc in (
        (
            "REQ-Q1",
            q1_metrics,
            {
                "all_coefficients": residual(coefficient_residual, 1e-9),
                "all_2079_inverse_times_min": residual(model_time_residual),
            },
        ),
        ("REQ-Q2", q2_metrics, {"all_LOCO_and_55A_vectors": residual(q2_vector_residual)}),
        (
            "REQ-Q3",
            q3_metrics,
            {"all_aging_vectors_and_perturbations_min": residual(q3_vector_residual)},
        ),
    ):
        # Report independently recalculated values; machine gate compares each metric too.
        recalc["metric_max_difference"] = residual(
            max(abs(v - out["final_metrics"][k]) for k, v in metrics.items()), 1e-9
        )
        constraint = residual(max(0.0, -remaining) if rid == "REQ-Q3" else 0.0)
        results[rid] = {
            "metric_values": metrics,
            "recalculation_residuals": recalc,
            "feasible": remaining >= 0 if rid == "REQ-Q3" else True,
            "constraint_residuals": {
                "nonnegative_remaining_time_min"
                if rid == "REQ-Q3"
                else "bounded_interpolation_domain": constraint
            },
            "domain_interpretation": (
                "Numerical/domain consistency onl"
                "y; no validation of state3 futur"
                "e truth or aging invariance."
            ),
        }
    payload = {
        "run_id": a.run_id,
        "output_sha256": hashlib.sha256(output_path.read_bytes()).hexdigest(),
        "requirements": results,
        "independence": (
            "Manual scalar PCHIP coefficients"
            ", bisection, closed-form regress"
            "ion; no producer imports; same m"
            "odel assumptions are not indepen"
            "dently established."
        ),
        "final_test_access": False,
    }
    a.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n")


if __name__ == "__main__":
    main()
