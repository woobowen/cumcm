"Prepare case artifacts without running models; bind only to main-confirmed commit."

import argparse
import importlib.util
import json
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[5]
CASE = ROOT / ".cache/pr12-rc8/fresh/2016/case-r2"
OWN = ROOT / ("evals/results/phase-004c5/CUMCM-2016-C-VALIDATION-004")
CORE = ROOT / (".agents/skills/cumcm-modeling-evidence/scripts/cumcm_case.py")
spec = importlib.util.spec_from_file_location("case_core", CORE)
core = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = core
spec.loader.exec_module(core)


def put(key, content):
    core.write_json(CASE / core.ARTIFACT_PATHS[key], core.artifact(key, content))


def dump(path, value):
    core.write_json(path, value)


def prepare():
    from openpyxl import load_workbook

    original = ROOT / ".cache/pr12-rc8/fresh/2016/official/c_files"
    mapping = {
        "CUMCM2016-Problem-C-Chinese-version.docx": "data/raw/problem.docx",
        "CUMCM2016-C-Appendix-Chinese.xlsx": "data/raw/appendix.xlsx",
    }
    for name, target in mapping.items():
        if not (CASE / target).exists():
            shutil.copyfile(original / name, CASE / target)
            (CASE / target).chmod(0o444)
        assert core.file_hash(original / name) == core.file_hash(CASE / target)
    for name in ("produce.py", "check.py"):
        shutil.copyfile(OWN / "code" / name, CASE / "models" / name)
    ids = ["REQ-Q1", "REQ-Q2", "REQ-Q3"]
    texts = [
        (
            "用初等函数表示附件1九条放电曲线，分别计算题定MRE；计算30、"
            "40、50、60、70A在9.8V的剩余放电时间。"
        ),
        "建立20至100A任意恒流曲线模型，按MRE评估精度，交付55A表格和曲线图形数据。",
        "预测附件2衰减状态3在最后观测时的剩余放电时间，并明确预测的证据和不确定性。",
    ]
    reqs = []
    for rid, text in zip(ids, texts, strict=True):
        reqs.append(
            {
                "requirement_id": rid,
                "text": text,
                "role": "PRIMARY",
                "required_evidence_classes": ["PROVIDED_EMPIRICAL"],
                "allowed_evidence_classes": ["PROVIDED_EMPIRICAL"],
                "minimum_data_fields": ["current_A", "voltage_V", "elapsed_min"]
                + (["state3_terminal_elapsed_min"] if rid == "REQ-Q3" else []),
                "required_time_scope": ["PROVIDED_DISCHARGE_RECORDS"],
                "required_entity_scope": ["PROVIDED_BATTERY_CURVES"],
                "external_data_allowed": False,
                "external_data_required": False,
                "simulation_substitution_allowed": False,
                "partial_completion_allowed": True,
                "dependency_requirements": ["REQ-Q1"] if rid == "REQ-Q2" else [],
                "completion_rule": (
                    "REQUIRED_MODEL_DELIVERABLES_AND_APPROPRIATE_SCIENTIFIC_VALIDATION"
                ),
                "scientific_facts_required": True,
                "claim_type": "EMPIRICAL" if rid == "REQ-Q1" else "PREDICTIVE",
                "accuracy_scope": (
                    "Q3 actual target error cannot be"
                    " established without future term"
                    "inal truth; no full completion c"
                    "laim."
                ),
            }
        )
    put(
        "problem_requirements",
        {
            "contract_version": "requirement-evidence/v1",
            "case_id": "CUMCM-2016-C-VALIDATION-004",
            "title": "电池剩余放电时间预测",
            "requirements": reqs,
            "ambiguities": [
                "初等函数允许分段多项式；若要求一个低参数全局公式，本方案不满足该更强解释。",
                (
                    "附件MRE使用231点，按9+0.005k(k=0..230)构"
                    "造公共电压格；采样时间用最后向下线性穿越得到。"
                ),
            ],
        },
    )
    source = {
        "source_id": "SRC-CUMCM-2016-C-WORKBOOK",
        "supports_requirement_ids": ids,
        "evidence_class": "PROVIDED_EMPIRICAL",
        "provenance": (
            "Official C-only workbook; archiv"
            "e lineage in this case registrat"
            "ion/official_retrieval.json; imm"
            "utable byte-identical case copy."
        ),
        "authority": "CUMCM_OFFICIAL_PROVIDED_INPUT",
        "retrieval_time": "2026-09-08T20:41:34.881829+00:00",
        "license_or_usage_status": (
            "OFFICIAL_PROBLEM_INPUT_LOCAL_ANALYSIS_ONLY_NO_RAW_REPUBLICATION"
        ),
        "geographic_scope": [],
        "time_scope": ["PROVIDED_DISCHARGE_RECORDS"],
        "entity_scope": ["PROVIDED_BATTERY_CURVES"],
        "field_schema": ["current_A", "voltage_V", "elapsed_min"],
        "hash": core.file_hash(CASE / "data/raw/appendix.xlsx"),
        "freshness": "IMMUTABLE_COMPETITION_INPUT",
        "limitations": [
            ("Attachment2 current magnitude unspecified; same current across states is stated."),
            ("Only prefix state3 observed; future terminal label unavailable."),
            ("Repeated voltage samples are not independent battery entities."),
        ],
    }
    put(
        "source_ledger",
        {
            "contract_version": "requirement-evidence/v1",
            "sources": [source],
            "answer_access_status": "NOT_ACCESSED",
        },
    )
    search = {
        "query": ("site:docs.scipy.org PchipInterpolator monotonic piecewise cubic"),
        "source_id": "SRC-SCIPY-PCHIP-DOCS",
        "url": (
            "https://docs.scipy.org/doc/scipy"
            "/reference/generated/scipy.inter"
            "polate.PchipInterpolator.html"
        ),
        "purpose": (
            "Verify coefficient ordering, shape preservation and independent Hermite formula."
        ),
        "accessed_at": "2026-09-08",
        "status": "READ_OFFICIAL_DOCUMENTATION",
        "local_version_checked": "scipy1.17.1 PchipInterpolator.__doc__",
        "no_solution_search": True,
    }
    put(
        "research_plan",
        {
            "mode": "OFFLINE_OR_REGISTERED_SEARCH",
            "questions": [
                (
                    "Can piecewise elementary polynom"
                    "ials describe the nonmonotone re"
                    "laxation plus monotone terminal "
                    "branch?"
                ),
                ("Can adjacent-current interpolation transfer to55A?"),
                (
                    "Does affine time transfer from c"
                    "omplete aging reference curves e"
                    "xtrapolate a truncated state?"
                ),
            ],
            "external_search": True,
            "search_ledger": [search],
            "source_fact_boundary": (
                "Documentation supports numerical method only, not battery transfer assumptions."
            ),
            "answer_access_status": "NOT_ACCESSED",
        },
    )
    assumption = {
        "assumptions": [
            {
                "id": "A1",
                "statement": (
                    "Within each observed current a p"
                    "iecewise polynomial interpolant "
                    "represents measured voltage; ini"
                    "tial relaxation is retained rath"
                    "er than forcing false global mon"
                    "otonicity."
                ),
                "challenge": (
                    "Strongly oscillatory or high-curvature gaps may be poorly represented."
                ),
            },
            {
                "id": "A2",
                "statement": (
                    "Between adjacent measured curren"
                    "ts, normalized voltage shape and"
                    " total discharge time vary smoot"
                    "hly."
                ),
                "challenge": (
                    "No55A observation; evaluate via leave-one-interior-current-out diagnostics."
                ),
            },
            {
                "id": "A3",
                "statement": (
                    "At fixed current, aging-state el"
                    "apsed time at equal voltage is a"
                    "pproximately scale or affine tra"
                    "nsform of complete reference cur"
                    "ves."
                ),
                "challenge": (
                    "Only one battery; future state3 behavior and error unknown; transfer may fail."
                ),
            },
        ],
        "symbols": {
            "t": {"meaning": "elapsed discharge time", "unit": "min"},
            "U": {"meaning": "terminal voltage", "unit": "V"},
            "I": {"meaning": "constant current", "unit": "A"},
            "T": {"meaning": "time at9V", "unit": "min"},
            "s": {"meaning": "t/T", "unit": "1"},
            "a": {"meaning": "aging time offset", "unit": "min"},
            "b": {"meaning": "aging time scale", "unit": "1"},
        },
        "formulas": [
            {
                "formula_id": "F-Q1",
                "requirement_ids": ["REQ-Q1"],
                "expression": (
                    "P_I(t)=c3_j(t-x_j)^3+c2_j(t-x_j)"
                    "^2+c1_j(t-x_j)+c0_j on [x_j,x_{j"
                    "+1}]; R_I(9.8)=T_I-P_I^{-1}(9.8)"
                    "."
                ),
            },
            {
                "formula_id": "F-MRE",
                "requirement_ids": ["REQ-Q1", "REQ-Q2"],
                "expression": (
                    "MRE=(1/231)sum_k abs(t_model(9+0"
                    ".005k)-t_sample(9+0.005k))/t_sam"
                    "ple(9+0.005k), k=0..230."
                ),
            },
            {
                "formula_id": "F-Q2",
                "requirement_ids": ["REQ-Q2"],
                "expression": (
                    "U(t,I)=(1-w)P_L(t*T_L/T_I)+w*P_H"
                    "(t*T_H/T_I). Baseline w=(I-L)/(H"
                    "-L), T arithmetic; candidate w=l"
                    "og(I/L)/log(H/L), T geometric."
                ),
            },
            {
                "formula_id": "F-Q3",
                "requirement_ids": ["REQ-Q3"],
                "expression": (
                    "t_3(U)=a+b*(t_new(U)+t_state1(U)"
                    "+t_state2(U))/3; R_3=t_3(9)-596."
                    "2. Fit only observed target pref"
                    "ix; baseline a=0, candidate free"
                    " a."
                ),
            },
        ],
    }
    put("assumptions_and_symbols", assumption)
    wb = load_workbook(CASE / "data/raw/appendix.xlsx", read_only=True, data_only=True)
    rows = list(wb["附件1"].values)
    audit_curves = []
    for j, i in enumerate(range(20, 101, 10), 1):
        vals = [(r[0], r[j]) for r in rows[2:] if isinstance(r[j], (int, float))]
        inc = [b[1] - a[1] for a, b in zip(vals[:-1], vals[1:], strict=True) if b[1] > a[1]]
        audit_curves.append(
            {
                "current_A": i,
                "observations": len(vals),
                "time_min_min": vals[0][0],
                "time_max_min": vals[-1][0],
                "voltage_increases": len(inc),
                "maximum_increase_V": max(inc, default=0),
                "trailing_blanks": "structural_end_of_discharge_not_imputed",
            }
        )
    wb.close()
    hashes = {rel: core.file_hash(CASE / rel) for rel in mapping.values()}
    access = {
        "initial_input_read": (
            "DOCX body; workbook sheet head/t"
            "ail; all worksheet rows material"
            "ized by metadata audit before sp"
            "lit design."
        ),
        "human_model_visible_values": (
            "Attachment1 first10 observations"
            " and20A terminal5 observations; "
            "Attachment2 first10 rows and9.02"
            "0..9.000 terminal5 rows, includi"
            "ng state2 terminal979.0; target "
            "state3 endpoint9.765V596.2min."
        ),
        "program_materialization": (
            "list(sheet.values) and numeric c"
            "ount iteration materialized ALL "
            "workbook labels. Not printed is "
            "not equivalent to not accessed."
        ),
        "answer_access": (
            "Official solutions remain SEALED and NOT_ACCESSED; no solution retrieval."
        ),
        "decision": (
            "No strict untouched prediction t"
            "est exists in this episode after"
            " this audit. Use DEVELOPMENT_NO_"
            "FINAL_EVALUATION; retain PREDICT"
            "IVE Q2/Q3 types. Formal Final/Cl"
            "aim/Handoff cannot claim full va"
            "lidation."
        ),
        "proposed_intermediate_tail_not_promoted": (
            "The idea of reserving state2 9.7"
            "55..9.030 was rejected before nu"
            "merical results because audited "
            "cells had already been materiali"
            "zed. It is not a blind test."
        ),
    }
    put(
        "data_audit",
        {
            "raw_immutable": True,
            "data_hashes": hashes,
            "data_dictionary": [
                {"field": "current_A", "unit": "A", "location": "attachment1 columns"},
                {
                    "field": "voltage_V",
                    "unit": "V",
                    "location": "attachment1 cells; attachment2 first column",
                },
                {
                    "field": "elapsed_min",
                    "unit": "min",
                    "location": ("attachment1 first column; attachment2 state columns"),
                },
            ],
            "quality_report": {
                "attachment1": audit_curves,
                "attachment2_counts": {"new": 301, "state1": 301, "state2": 301, "state3": 148},
                "target3_endpoint": {"voltage_V": 9.765, "elapsed_min": 596.2},
                "target3_missing_future_labels": 153,
            },
            "leakage_findings": [access],
            "acquisition_plans": [],
            "processed_lineage": [],
        },
    )
    assessments = []
    for rid in ids:
        assessments.append(
            {
                "requirement_id": rid,
                "data_sufficiency_status": "PARTIAL" if rid == "REQ-Q3" else "SUFFICIENT",
                "missing_fields": ["state3_terminal_elapsed_min"] if rid == "REQ-Q3" else [],
                "missing_entities": [],
                "missing_time_scope": [],
                "candidate_sources": [source["source_id"]],
                "acquisition_cost": "0; no additional source acquisition",
                "acquisition_time": "0",
                "allowed_substitutions": [],
                "forbidden_substitutions": [
                    "Fabricated future state3 labels",
                    ("Relabeling an audited data slice as untouched Final"),
                    "PREDICTIVE-to-DESCRIPTIVE substitution",
                ],
                "affected_downstream_stages": [
                    "FINAL_RUN",
                    "CLAIM_EVIDENCE_VALIDATION",
                    "MODELING_TO_PAPER_HANDOFF",
                ]
                if rid != "REQ-Q1"
                else [],
            }
        )
    put(
        "data_sufficiency",
        {
            "contract_version": "data-sufficiency/v1",
            "requirements": reqs,
            "sources": [source],
            "acquisition_plans": [],
            "aggregate_completion_claimed": False,
            "requirement_assessments": assessments,
            "coverage_mode_by_requirement": {
                rid: {"mode": "SINGLE_SOURCE", "source_id": source["source_id"]} for rid in ids
            },
            "source_compositions": [],
            "interpretation": (
                "PARTIAL aggregate: Q1/Q2 model i"
                "nputs exist, but primary Q3 nece"
                "ssary target terminal truth is m"
                "issing. Only limited model compu"
                "tation and forecast diagnosis ca"
                "n proceed; full acceptance is bl"
                "ocked."
            ),
        },
    )
    candidates = [
        {
            "candidate_id": "BASELINE",
            "baseline": True,
            "method": (
                "17 uniform-time knots plus first"
                "31 records; piecewise linear vol"
                "tage; arithmetic current/time in"
                "terpolation; zero-intercept agin"
                "g scale."
            ),
            "rationale": ("Simplest local interpolation and one-parameter aging transfer."),
        },
        {
            "candidate_id": "CUBIC_LOG_AFFINE",
            "baseline": False,
            "method": (
                "33 uniform-time knots plus first"
                "31 records; shape-preserving pie"
                "cewise cubic; logarithmic curren"
                "t/geometric duration interpolati"
                "on; affine aging transfer."
            ),
            "rationale": (
                "Resolves curvature and allows ag"
                "ing offset; extra flexibility ma"
                "y extrapolate poorly."
            ),
        },
    ]
    put(
        "model_candidates",
        {
            "candidates": candidates,
            "baseline_definition": (
                "Unique BASELINE; compare fixed f"
                "ull portfolios with equal weight"
                "s on two dimensionless developme"
                "nt MRE diagnostics."
            ),
        },
    )
    probe = {
        "candidate_id": "CONTRACT-PROBE",
        "status": "CONTRACT_PROBE",
        "probe_only": True,
        "ranking_eligible": False,
        "result_values_are_placeholders": True,
        "final_metrics": {"placeholder": 0.0},
        "claim_scope": "NONRESULT_STRUCTURE_ONLY",
        "requirement_claims": {
            rid: {
                "claim_id": "CLAIM-PROBE-" + rid[4:],
                "claim_text": "NONRESULT placeholder",
                "evidence_artifact_ids": ["PLACEHOLDER"],
            }
            for rid in ids
        },
        "figure_ready_data": [{"placeholder": True}],
        "uncertainty": {"placeholder": True},
        "limitations": ["NONRESULT placeholders"],
        "robustness_evidence": {
            "metric": "placeholder",
            "metric_direction": "MIN",
            "perturbations": [
                {
                    "perturbation_id": "P0",
                    "metric": "placeholder",
                    "result": 0.0,
                    "evidence": ("DETERMINISTIC_RECOMPUTATION_FROM_BOUND_INPUTS"),
                }
            ],
            "failure_cases": ["NONRESULT"],
        },
    }
    dump(CASE / ("experiments/selected_output_contract_probe.json"), probe)
    proposal = {
        "schema_version": "rc8-case-pre-run-proposal/v1",
        "case_id": "CUMCM-2016-C-VALIDATION-004",
        "status": "PROPOSED_BEFORE_NUMERICAL_MODEL_RESULTS",
        "created_at": core.utc_now(),
        "revision": {
            "number": 2,
            "active_case_root": ".cache/pr12-rc8/fresh/2016/case-r2",
            "superseded_case_root": ".cache/pr12-rc8/fresh/2016/case",
            "archive_index": "pre_run/revision_history.json",
            "same_episode": True,
            "prior_model_capture_count": 0,
        },
        "case_budget": {
            "anchor": "2026-09-08T20:37:30Z",
            "candidate_stop": "2026-09-08T22:17:30Z",
            "absolute_deadline": "2026-09-08T22:37:30Z",
            "wall_seconds": 7200,
        },
        "requirements": reqs,
        "data_audit": core.read_artifact(CASE, "data_audit")["content"],
        "access_ledger": access,
        "sources": [source],
        "research_log": [search],
        "assumptions": assumption,
        "candidates": candidates,
        "random_seeds": [1729],
        "metric": {
            "name": "joint_validation_MRE",
            "direction": "MIN",
            "formula": (
                "0.5*mean seven leave-one-interio"
                "r-current-out MRE +0.5*state1 te"
                "mporal-tail transfer MRE from ne"
                "w-state reference"
            ),
            "units": "dimensionless",
            "q1_fitted_MRE": "diagnostic only, not selection",
            "tie_rule": "ARGMIN_THEN_ID",
        },
        "validation_boundary": {
            "train": [
                "ATTACHMENT1_CURVES",
                "ATTACHMENT2_NEW_COMPLETE",
                "ATTACHMENT2_STATE1_PREFIX",
                "ATTACHMENT2_STATE2_COMPLETE",
                "ATTACHMENT2_STATE3_PREFIX",
            ],
            "validation": ["ATTACHMENT2_STATE1_TAIL_9.760_TO_9.030"],
            "test": [],
            "interpretation": (
                "Development cross-validation and"
                " temporal transfer diagnostics; "
                "no untouched heldout."
            ),
        },
        "evaluation_design": {
            "mode": "DEVELOPMENT_NO_FINAL_EVALUATION",
            "reason": (
                "All official workbook labels wer"
                "e materialized by intake audit; "
                "no later internal split can be a"
                "sserted untouched; actual future"
                " state3 label absent."
            ),
            "predictive_final_authorized": False,
        },
        "scientific_evidence_needs": {
            "Q1": (
                "Per-current full coefficient vec"
                "tors,231 inverse predictions and"
                " actual sample-crossing MRE; req"
                "uired5 remaining times."
            ),
            "Q2": (
                "Arbitrary-current formula, seven"
                " current-level LOCO errors and55"
                "A101-point table/figure data."
            ),
            "Q3": (
                "Prefix-only fitted aging transfe"
                "r, total and remaining forecast,"
                " state1 development-tail error, "
                "actual input perturbation and te"
                "mplate spread; state3 actual acc"
                "uracy UNKNOWN."
            ),
        },
        "independent_checker": {
            "path": "models/check.py",
            "algorithm": (
                "Manual scalar Hermite coefficien"
                "ts, binary-search inverses, clos"
                "ed-form OLS; never imports produ"
                "cer or common solution helper."
            ),
            "coverage": (
                "All polynomial coefficients; nin"
                "e231-point inverse vectors; seve"
                "n231-point LOCO vectors;55A101 r"
                "ows; state3/state2/reference pre"
                "dictions; all perturbation resul"
                "ts; per-requirement metrics."
            ),
            "tolerances": {
                "coefficient_absolute": 1e-9,
                "time_vector_min_absolute": 1e-7,
                "metric_absolute": 1e-9,
            },
            "domain": (
                "Nonnegative predicted remaining "
                "duration separate from arithmeti"
                "c agreement; does not establish "
                "future truth or aging invariance"
                "."
            ),
        },
        "robustness": [
            "State3 prefix times+1min and-1min",
            ("Drop last10 state3 prefix observations and refit"),
            ("Each of all3 complete references separately; range is not a confidence interval"),
        ],
        "execution_budget": {
            "fixed_candidate_attempts": 2,
            ("maximum_captured_attempts_entire_episode_all_revisions"): 4,
            "remaining_repair_captures_after_two_planned": 2,
            "seeds_each": 1,
            "timeout_seconds_each": 300,
            "maximum_repairs_same_root_cause": 3,
            "repair_rule": (
                "No erased attempts. No default p"
                "ost-result model changes. Any ne"
                "cessary code correction requires"
                " main Git delivery and retained "
                "failed capture; unresolved freez"
                "e incompatibility yields PARTIAL"
                "/BLOCK."
            ),
            "final_access_budget": 0,
        },
        "stop_rule": (
            "At most4 total model captures ac"
            "ross every revision, normally2 f"
            "ixed candidates, plus independen"
            "t checkers; stop main numeric wo"
            "rk on completion or22:17:30Z, wh"
            "ichever earlier. No re-run after"
            " candidate terminal freeze. No F"
            "inal test authorized; preserve s"
            "ubsequent scientific BLOCK."
        ),
        "anticipated_limits": [
            ("No full scientific PASS can be asserted merely from computation/CLI success."),
            (
                "Q2 andQ3 remain PREDICTIVE; abse"
                "nce of valid heldout and state3 "
                "future truth retained."
            ),
            (
                "Ordinary source/metadata audit r"
                "ead too much for later untouched"
                " holdout; permanent access histo"
                "ry retained."
            ),
        ],
        "required_code_repository_paths": [
            str((OWN / "code" / n).relative_to(ROOT))
            for n in ("produce.py", "check.py", "prepare.py")
        ],
        "case_artifact_hashes": {
            key: core.file_hash(CASE / core.ARTIFACT_PATHS[key])
            for key in (
                "problem_requirements",
                "research_plan",
                "source_ledger",
                "assumptions_and_symbols",
                "data_audit",
                "data_sufficiency",
                "model_candidates",
            )
        },
    }
    dump(OWN / "pre_run/pre_run_proposal.json", proposal)
    dump(OWN / "pre_run/read_and_search_ledger.json", {"access": access, "search": [search]})
    print(
        json.dumps(
            {
                "status": "PREPARED_NO_MODEL_RUN",
                "proposal": str((OWN / "pre_run/pre_run_proposal.json").relative_to(ROOT)),
            }
        )
    )


def bind(commit):
    proposal = core.load_json(OWN / "pre_run/pre_run_proposal.json")
    records = []
    for filename in ("produce.py", "check.py"):
        path = CASE / "models" / filename
        records.append(
            {
                "scope": "CASE_ROOT",
                "path": "models/" + filename,
                "repository_path": str((OWN / "code" / filename).relative_to(ROOT)),
                "sha256": core.file_hash(path),
            }
        )
    records.append(
        {
            "scope": "SKILL_ROOT",
            "path": "scripts/cumcm_case.py",
            "repository_path": str(CORE.relative_to(ROOT)),
            "sha256": core.file_hash(CORE),
        }
    )
    plan = {
        "preregistered": True,
        "execution_prepared": True,
        "candidate_ids": ["BASELINE", "CUBIC_LOG_AFFINE"],
        "baseline_id": "BASELINE",
        "metric": "joint_validation_MRE",
        "metric_direction": "MIN",
        "aggregation_rule": "MEAN_PER_CANDIDATE_THEN_DIRECTION_THEN_ID",
        "selection_rule": "ARGMIN_THEN_ID",
        "random_seeds": [1729],
        "splits": {k: proposal["validation_boundary"][k] for k in ("train", "validation", "test")},
        "required_input_hashes": core.read_artifact(CASE, "data_audit")["content"]["data_hashes"],
        "required_code_files": records,
        "code_commit": commit,
        "stop_rule": proposal["stop_rule"],
        "handoff_generated_at": "2026-09-08T22:17:30Z",
        "evaluation_design": proposal["evaluation_design"],
    }
    plan["trusted_freeze_registry"] = {
        "candidate_set": core.canonical_hash(plan["candidate_ids"]),
        "metric": core.canonical_hash(
            {
                "name": plan["metric"],
                "direction": plan["metric_direction"],
                "aggregation_rule": plan["aggregation_rule"],
                "selection_rule": plan["selection_rule"],
            }
        ),
        "seed_schedule": core.canonical_hash(plan["random_seeds"]),
        "split_assignment": core.canonical_hash(plan["splits"]),
        "baseline": core.canonical_hash(plan["baseline_id"]),
        "input_set": core.canonical_hash(plan["required_input_hashes"]),
        "execution_policy": core.canonical_hash(
            core.execution_policy_payload(
                plan["stop_rule"], plan["handoff_generated_at"], plan["evaluation_design"]
            )
        ),
        "code_set": core.canonical_hash(records),
        "code_commit": core.canonical_hash(commit),
    }
    put("experiment_plan", plan)
    print(
        json.dumps(
            {
                "status": "BOUND",
                "code_commit": commit,
                "verified_freezes": core.trusted_freezes(CASE),
            }
        )
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--bind-commit")
    args = parser.parse_args()
    if args.bind_commit:
        bind(args.bind_commit)
    else:
        prepare()
