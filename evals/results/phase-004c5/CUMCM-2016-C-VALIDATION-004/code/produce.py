"""First-party, preregistered battery curve and aging transfer models.

No reference solutions. This episode authorizes no Final evaluation or held-out claims.
"""

import argparse
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator, PPoly
from scipy.optimize import brentq

REQS = ["REQ-Q1", "REQ-Q2", "REQ-Q3"]
SOURCE = "SRC-CUMCM-2016-C-WORKBOOK"
RAW = "data/raw/appendix.xlsx"


def write(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n")


def inputs(root):
    workbook = load_workbook(root / RAW, data_only=True, read_only=True)
    curves = {}
    rows = list(workbook["附件1"].values)
    for j, current in enumerate(range(20, 101, 10), 1):
        pairs = [(r[0], r[j]) for r in rows[2:] if isinstance(r[j], (int, float))]
        curves[current] = np.array(pairs, dtype=float)
    # Development-only: use all three complete references; no held-out claim.
    sheet = workbook["附件2"]
    voltage, refs, state2_prefix, state3_prefix = [], [], [], []
    for row in range(3, 304):
        u = sheet.cell(row, 1).value
        if not isinstance(u, (int, float)):
            continue
        voltage.append(float(u))
        refs.append([float(sheet.cell(row, j).value) for j in (2, 3, 4)])
        if u >= 9.765 - 1e-10:
            state2_prefix.append(float(sheet.cell(row, 4).value))
            state3_prefix.append(float(sheet.cell(row, 5).value))
    workbook.close()
    return (
        curves,
        np.array(voltage),
        np.array(refs),
        np.array(state2_prefix),
        np.array(state3_prefix),
    )


def polynomial(curve, candidate):
    n = len(curve)
    count = 17 if candidate == "BASELINE" else 33
    indices = np.unique(
        np.r_[np.arange(min(31, n)), np.rint(np.linspace(0, n - 1, count)).astype(int)]
    )
    t, u = curve[indices].T
    if candidate == "BASELINE":
        coefficients = np.vstack(
            [np.zeros(len(t) - 1), np.zeros(len(t) - 1), np.diff(u) / np.diff(t), u[:-1]]
        )
        p = PPoly(coefficients, t, extrapolate=False)
    else:
        p = PchipInterpolator(t, u, extrapolate=False)
    return p, {
        "time_knots_min": t.tolist(),
        "coefficients_descending": p.c.tolist(),
        "knot_indices": indices.tolist(),
        "duration_min": float(t[-1]),
    }


def crossing(poly, voltage):
    roots = poly.solve(float(voltage), extrapolate=False)
    roots = roots[np.isfinite(roots) & (roots >= -1e-9) & (roots <= poly.x[-1] + 1e-9)]
    if len(roots) == 0:
        if abs(float(poly(poly.x[-1])) - voltage) <= 1e-9:
            return float(poly.x[-1])
        raise ValueError("MODEL_VOLTAGE_OUTSIDE_DOMAIN")
    return float(np.max(roots))


def raw_cross(curve, voltage):
    # Last downward crossing: early relaxation rebound is retained in the forward fit.
    pairs = np.where((curve[:-1, 1] >= voltage) & (curve[1:, 1] <= voltage))[0]
    if not len(pairs):
        raise ValueError("RAW_VOLTAGE_OUTSIDE_DOMAIN")
    k = int(pairs[-1])
    a, b = curve[k], curve[k + 1]
    return float(b[0] if a[1] == b[1] else a[0] + (a[1] - voltage) / (a[1] - b[1]) * (b[0] - a[0]))


def surface(polys, durations, current, candidate):
    keys = sorted(polys)
    lo = max(x for x in keys if x < current)
    hi = min(x for x in keys if x > current)
    if candidate == "BASELINE":
        w = (current - lo) / (hi - lo)
        duration = (1 - w) * durations[lo] + w * durations[hi]
    else:
        w = np.log(current / lo) / np.log(hi / lo)
        duration = np.exp((1 - w) * np.log(durations[lo]) + w * np.log(durations[hi]))

    def voltage(s):
        return (1 - w) * polys[lo](np.asarray(s) * durations[lo]) + w * polys[hi](
            np.asarray(s) * durations[hi]
        )

    def time_at(u):
        # All evaluation voltages are below the transient/plateau.
        if abs(float(voltage(1)) - u) < 1e-9:
            return float(duration)
        return float(duration * brentq(lambda s: float(voltage(s)) - u, 0, 1, xtol=1e-13))

    return (
        voltage,
        time_at,
        float(duration),
        {"left_current_A": lo, "right_current_A": hi, "weight": float(w)},
    )


def transfer(reference, prefix, candidate):
    x = reference[: len(prefix)]
    if candidate == "BASELINE":
        a, b = 0.0, float(np.dot(x, prefix) / np.dot(x, x))
    else:
        a, b = np.linalg.lstsq(np.column_stack([np.ones(len(x)), x]), prefix, rcond=None)[0]
    return float(a), float(b), a + b * reference


def mre(predicted, actual):
    return float(np.mean(np.abs(np.asarray(predicted) - actual) / actual))


def develop(root, candidate):
    curves, voltage, refs, state2, state3 = inputs(root)
    grid = 9 + np.arange(231) * 0.005
    polys, records, durations = {}, {}, {}
    for current, curve in curves.items():
        polys[current], records[str(current)] = polynomial(curve, candidate)
        durations[current] = float(curve[-1, 0])
    per_current, rows = {}, []
    for current in sorted(curves):
        actual = np.array([raw_cross(curves[current], v) for v in grid])
        predicted = np.array([crossing(polys[current], v) for v in grid])
        per_current[str(current)] = mre(predicted, actual)
        rows.append(
            {
                "current_A": current,
                "MRE": per_current[str(current)],
                "model_elapsed_at_voltage_min": predicted.tolist(),
                "voltage_grid_V": grid.tolist(),
            }
        )
    remaining = {str(i): durations[i] - crossing(polys[i], 9.8) for i in range(30, 71, 10)}
    loco = []
    for current in range(30, 91, 10):
        keep = {i: p for i, p in polys.items() if i != current}
        _, inv, _, _ = surface(keep, durations, current, candidate)
        estimated = np.array([inv(v) for v in grid])
        observed = np.array([raw_cross(curves[current], v) for v in grid])
        loco.append(
            {
                "current_A": current,
                "MRE": mre(estimated, observed),
                "predicted_elapsed_min": estimated.tolist(),
            }
        )
    q2_error = float(np.mean([r["MRE"] for r in loco]))
    forward55, _, duration55, blend55 = surface(polys, durations, 55, candidate)
    s = np.linspace(0, 1, 101)
    table55 = [{"time_min": float(z * duration55), "voltage_V": float(forward55(z))} for z in s]
    cutoff = len(state3)
    # Development transfer uses state1 tail; state3 uses all three complete references.
    _, _, dev_prediction = transfer(refs[:, 0], refs[:cutoff, 1], candidate)
    val_mask = (voltage < 9.765 - 1e-10) & (voltage >= 9.030 - 1e-10)
    q3_error = mre(dev_prediction[val_mask], refs[val_mask, 1])
    reference = refs.mean(axis=1)
    a3, b3, prediction3 = transfer(reference, state3, candidate)
    a2, b2, prediction2 = transfer(refs[:, :2].mean(axis=1), state2, candidate)
    last_time = float(state3[-1])
    life = float(prediction3[-1] - last_time)
    alternatives = []
    for j in range(3):
        _, _, pred = transfer(refs[:, j], state3, candidate)
        alternatives.append(float(pred[-1] - last_time))
    perturbations = []
    detail = []
    # Actual perturbations to measured prefix times and cutoff, refit each from inputs.
    for label, prefix in (
        ("PREFIX_TIME_PLUS_1_MIN", state3 + 1),
        ("PREFIX_TIME_MINUS_1_MIN", state3 - 1),
        ("PREFIX_DROP_LAST_10", state3[:-10]),
    ):
        _, _, pred = transfer(reference, prefix, candidate)
        value = float(pred[-1] - last_time)
        perturbations.append(
            {
                "perturbation_id": label,
                "metric": "q3_remaining_min",
                "result": value,
                "evidence": ("DETERMINISTIC_RECOMPUTATION_FROM_BOUND_INPUTS"),
            }
        )
        detail.append({"id": label, "remaining_min": value, "delta_min": value - life})
    final_metrics = {
        "q1_mean_fitted_MRE": float(np.mean(list(per_current.values()))),
        "q2_loco_MRE": q2_error,
        "q3_validation_tail_MRE": q3_error,
        "q3_remaining_min": life,
    }
    final_metrics.update({f"q1_MRE_{i}A": value for i, value in per_current.items()})
    final_metrics.update({f"q1_remaining_{i}A_min": value for i, value in remaining.items()})
    metrics_by_req = {
        "REQ-Q1": {k: v for k, v in final_metrics.items() if k.startswith("q1_")},
        "REQ-Q2": {"q2_loco_MRE": q2_error},
        "REQ-Q3": {"q3_validation_tail_MRE": q3_error, "q3_remaining_min": life},
    }
    claims = {
        "REQ-Q1": (
            "九条分段初等多项式放电曲线、231电压格点拟合MRE及30至70"
            "A在9.8V的模型剩余时间；拟合误差不等于新电池泛化误差。"
        ),
        "REQ-Q2": (
            "20至100A内按相邻电流和归一化放电时间插值，给出55A曲线；"
            "精度证据为七个中间电流的留一电流诊断。"
        ),
        "REQ-Q3": (
            "衰减状态3预测剩余时间为模型迁移外推结果，真实终止标签缺失；相邻"
            "衰减状态验证不证明状态3实际误差。"
        ),
    }
    evidence = {}
    for rid in REQS:
        evidence[rid] = {
            "generation_method": "PREDICTION" if rid != "REQ-Q1" else "EMPIRICAL_ANALYSIS",
            "source_ids": [SOURCE],
            "scope": {
                "fields": ["current_A", "voltage_V", "elapsed_min"],
                "time": ["PROVIDED_DISCHARGE_RECORDS"],
                "entities": ["PROVIDED_BATTERY_CURVES"],
            },
            "metric_values": metrics_by_req[rid],
            "status": "INSUFFICIENT" if rid == "REQ-Q3" else "COMPUTED",
        }
    return {
        "candidate_id": candidate,
        "status": "SUCCESS",
        "validation_metrics": {"joint_validation_MRE": (q2_error + q3_error) / 2},
        "final_metrics": final_metrics,
        "claim_scope": (
            "逐问模型与有限内部诊断；状态3终点预测尚无目标真实标签验证，不能宣称全题科学PASS。"
        ),
        "requirement_claims": {
            rid: {
                "claim_id": "CLAIM-2016-" + rid[4:],
                "claim_text": claims[rid],
                "evidence_artifact_ids": ["output.json#" + rid],
            }
            for rid in REQS
        },
        "scientific_evidence": evidence,
        "Q1": {
            "curve_models": records,
            "per_current_MRE": per_current,
            "remaining_at_9_8V_min": remaining,
            "mre_evaluation": rows,
            "definition": (
                "mean(abs(t_model(U_k)-t_sample(U"
                "_k))/t_sample(U_k)), U_k=9+0.005"
                "k,k=0..230; sample inverse is la"
                "st downward piecewise linear cro"
                "ssing"
            ),
        },
        "Q2": {
            "model": ("U(t,I)=(1-w)P_L(t*T_L/T_I)+w*P_H(t*T_H/T_I); I in [L,H]"),
            "blend": "arithmetic I and T" if candidate == "BASELINE" else "log I and geometric T",
            "duration55_min": duration55,
            "blend55": blend55,
            "table55": table55,
            "leave_one_current_out": loco,
        },
        "Q3": {
            "transfer_equation": ("t_state3(U)=a+b*(t_new(U)+t_state1(U)+t_state2(U))/3"),
            "a_min": a3,
            "b": b3,
            "current_voltage_V": float(voltage[cutoff - 1]),
            "current_elapsed_min": last_time,
            "predicted_total_min": float(prediction3[-1]),
            "predicted_remaining_min": life,
            "prediction_voltage_V": voltage.tolist(),
            "predicted_elapsed_min": prediction3.tolist(),
            "reference_specific_remaining_min": alternatives,
            "state2_development_prediction_min": prediction2.tolist(),
            "state2_transfer": {"a_min": a2, "b": b2},
            "state1_validation_prediction_min": dev_prediction.tolist(),
            "target_terminal_label_available": False,
            "final_evaluation_authorized": False,
        },
        "figure_ready_data": [
            {"figure_id": "FIG-55A", "x": "time_min", "y": "voltage_V", "data": table55}
        ],
        "uncertainty": {
            "kind": ("MODEL_AND_PREFIX_SENSITIVITY_NOT_CONFIDENCE_INTERVAL"),
            "reference_spread_min": alternatives,
            "prefix_perturbations": detail,
            "independent_battery_count": 1,
            "target_error": "UNKNOWN",
        },
        "limitations": [
            ("State3 true 9V termination is absent; prediction accuracy unknown."),
            (
                "Only one battery across aging st"
                "ates; voltage rows are correlate"
                "d, not independent battery repli"
                "cates."
            ),
            (
                "Initial voltage relaxation viola"
                "tes literal global monotonicity;"
                " inverse uses last downward cros"
                "sing."
            ),
            (
                "Q1 piecewise polynomial represen"
                "tation uses 17 or33 uniform-time"
                " knots plus first31 sample point"
                "s, not one compact global elemen"
                "tary formula."
            ),
            ("Q2 interpolates within20–100A only; no observation at55A exists."),
            (
                "All workbook rows were materiali"
                "zed by intake audit; no untouche"
                "d Final labels are claimed. All "
                "reported cross-validation errors"
                " are Development diagnostics."
            ),
        ],
        "robustness_evidence": {
            "metric": "q3_remaining_min",
            "metric_direction": "MIN",
            "perturbations": perturbations,
            "failure_cases": [
                ("No state3 target error can be established without its future termination."),
                (
                    "Transfer shape may change with a"
                    "ging; sensitivity spread is not "
                    "calibrated uncertainty."
                ),
            ],
        },
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--case-root", type=Path, required=True)
    parser.add_argument("--candidate-id", choices=["BASELINE", "CUBIC_LOG_AFFINE"], required=True)
    parser.add_argument("--seed", type=int, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--final-evaluation", action="store_true")
    parser.add_argument("--authorization-hash")
    parser.add_argument("--final-output", type=Path)
    args = parser.parse_args()
    if args.final_evaluation:
        raise SystemExit("FINAL_EVALUATION_NOT_AUTHORIZED_BY_DEVELOPMENT_ONLY_PROTOCOL")
    else:
        write(args.output, develop(args.case_root, args.candidate_id))


if __name__ == "__main__":
    main()
