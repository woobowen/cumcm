#!/usr/bin/env python3
"""Independent raw-workbook/plan recomputation. Never imports producer or its helper.

Only W001-W216 numeric cells are requested. Nominal constraints, future assumptions,
and Development historical stresses are distinct. Optimizer self-reports cannot
replace raw balances. A simple optimistic capacity lower bound can certify count
when it meets the feasible plan count; otherwise optimality remains unverified.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

CONSUMPTION = {"A": 0.60, "B": 0.66, "C": 0.72}
PRICES = {"A": 1.20, "B": 1.10, "C": 1.00}
RAW = {
    "raw/case_files/附件1 近5年402家供应商的相关数据.xlsx": (
        "1b93a7598b8f0489e3c054439967bca654a83654837908c833752617cf950e1b"
    ),
    "raw/case_files/附件2 近5年8家转运商的相关数据.xlsx": (
        "29e1499b133aa88d765902081bfb827f78ffb4c091f5f75341fc37d844344685"
    ),
}
TOLERANCE = 0.05


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def residual(
    value: float, limit: float = 0.0, relation: str = "LE", tolerance: float = TOLERANCE
) -> dict[str, Any]:
    if not math.isfinite(float(value)) or not math.isfinite(float(limit)):
        raise ValueError("NONFINITE_CHECK_RESIDUAL")
    return {
        "value": float(value),
        "limit": float(limit),
        "relation": relation,
        "tolerance": float(tolerance),
    }


def passes(rows: dict[str, dict[str, Any]]) -> bool:
    for row in rows.values():
        value, limit, tol = row["value"], row["limit"], row["tolerance"]
        if row["relation"] == "LE" and value > limit + tol:
            return False
        if row["relation"] == "GE" and value < limit - tol:
            return False
        if row["relation"] == "EQ" and abs(value - limit) > tol:
            return False
    return True


def read_inputs(root: Path) -> dict[str, Any]:
    observed = {name: digest(root / name) for name in RAW}
    if observed != RAW:
        raise ValueError("OFFICIAL_WORKBOOK_HASH_MISMATCH")
    names = list(RAW)
    book = load_workbook(root / names[0], read_only=True, data_only=True)
    order = list(book["企业的订货量（m³）"].iter_rows(max_col=218, values_only=True))
    supplied = list(book["供应商的供货量（m³）"].iter_rows(max_col=218, values_only=True))
    book.close()
    book = load_workbook(root / names[1], read_only=True, data_only=True)
    losses = list(book["运输损耗率（%）"].iter_rows(max_col=217, values_only=True))
    book.close()
    weeks = [f"W{i:03d}" for i in range(1, 217)]
    if (
        list(order[0][2:]) != weeks
        or list(supplied[0][2:]) != weeks
        or list(losses[0][1:]) != weeks
    ):
        raise ValueError("AUTHORIZED_WEEK_HEADERS_INVALID")
    order = [r for r in order[1:] if r[0] is not None]
    supplied = [r for r in supplied[1:] if r[0] is not None]
    losses = [r for r in losses[1:] if r[0] is not None]
    if len(order) != 402 or len(supplied) != 402 or len(losses) != 8:
        raise ValueError("OFFICIAL_INPUT_ENTITY_COUNTS_INVALID")
    if [(r[0], r[1]) for r in order] != [(r[0], r[1]) for r in supplied]:
        raise ValueError("SUPPLIER_ORDER_SUPPLY_IDENTITY_MISMATCH")
    values = {
        "order": np.asarray([r[2:] for r in order], float),
        "supply": np.asarray([r[2:] for r in supplied], float),
        "loss_pct": np.asarray([r[1:] for r in losses], float),
    }
    if any(not np.all(np.isfinite(v)) or np.any(v < 0) for v in values.values()):
        raise ValueError("INVALID_RAW_NUMERIC_VALUES")
    if any(str(r[1]) not in CONSUMPTION for r in order):
        raise ValueError("MATERIAL_TYPE_UNKNOWN")
    return {
        **values,
        "supplier_ids": [str(r[0]) for r in order],
        "types": [str(r[1]) for r in order],
        "carrier_ids": [str(r[0]) for r in losses],
        "input_hashes": observed,
    }


def derive_parameters(raw: dict[str, Any], candidate: str, seed: int) -> dict[str, Any]:
    if candidate not in {
        "BASELINE_MEAN_GREEDY",
        "ROBUST_QUANTILE_LEXICOGRAPHIC",
        "SCENARIO_CVAR_PORTFOLIO",
    }:
        raise ValueError("CANDIDATE_UNKNOWN")
    rng = np.random.default_rng(seed)
    ratios = []
    caps = []
    features = []
    for i, typ in enumerate(raw["types"]):
        order = raw["order"][i, :168]
        supply = raw["supply"][i, :168]
        active = order > 0
        fulfillment = np.minimum(2.0, np.maximum(0.0, supply[active] / order[active]))
        caps.append(float(np.quantile(order[active], 0.75)) if np.any(active) else 0.0)
        if not len(fulfillment):
            ratio = 0.0
        elif candidate == "BASELINE_MEAN_GREEDY":
            ratio = float(fulfillment.mean())
        elif candidate == "ROBUST_QUANTILE_LEXICOGRAPHIC":
            ratio = float(np.quantile(fulfillment, 0.2))
        else:
            samples = np.asarray(
                [rng.choice(fulfillment, len(fulfillment), replace=True).mean() for _ in range(64)]
            )
            ratio = float(samples[samples <= np.quantile(samples, 0.2)].mean())
        ratios.append(min(1.5, max(0.0, ratio)))
        positive = supply[supply > 0]
        features.append(
            [
                float(supply.mean() / CONSUMPTION[typ]),
                float(np.mean(fulfillment >= 0.95)) if len(fulfillment) else 0.0,
                1 / (1 + float(positive.std()) / max(float(positive.mean()), 1e-8))
                if len(positive)
                else 0.0,
                float(np.mean(supply > 0)),
            ]
        )
    loss = []
    for row in raw["loss_pct"][:, :168]:
        positive = row[row > 0] / 100.0
        if not len(positive):
            raise ValueError("NO_OBSERVED_CARRIER_LOSS")
        if candidate == "BASELINE_MEAN_GREEDY":
            value = positive.mean()
        elif candidate == "ROBUST_QUANTILE_LEXICOGRAPHIC":
            value = np.quantile(positive, 0.75)
        else:
            value = positive[positive >= np.quantile(positive, 0.8)].mean()
        loss.append(float(value))
    return {
        "ratio": np.asarray(ratios),
        "order_capacity": np.asarray(caps),
        "loss": np.asarray(loss),
        "supply_capacity": np.asarray(ratios) * np.asarray(caps),
        "features": np.asarray(features),
    }


def decode_plan(plan: dict[str, Any], raw: dict[str, Any]) -> tuple[np.ndarray, np.ndarray]:
    supplier = {sid: i for i, sid in enumerate(raw["supplier_ids"])}
    carrier = {cid: j for j, cid in enumerate(raw["carrier_ids"])}
    orders = np.zeros(len(supplier))
    transport = np.zeros((len(supplier), len(carrier)))
    order_seen = set()
    route_seen = set()
    for row in plan["orders"]:
        sid = row["supplier_id"]
        value = row["volume_m3"]
        if (
            sid not in supplier
            or sid in order_seen
            or isinstance(value, bool)
            or not math.isfinite(value)
            or value < 0
        ):
            raise ValueError("ORDER_RECORD_INVALID_OR_DUPLICATE")
        orders[supplier[sid]] = value
        order_seen.add(sid)
    for row in plan["transport"]:
        sid, cid = row["supplier_id"], row["carrier_id"]
        value = row["volume_m3"]
        if (
            sid not in supplier
            or cid not in carrier
            or (sid, cid) in route_seen
            or isinstance(value, bool)
            or not math.isfinite(value)
            or value < 0
        ):
            raise ValueError("TRANSPORT_RECORD_INVALID_OR_DUPLICATE")
        transport[supplier[sid], carrier[cid]] = value
        route_seen.add((sid, cid))
    if plan.get("repeat_for_weeks") != list(range(1, 25)):
        raise ValueError("PLAN_24_WEEK_COVERAGE_INVALID")
    return orders, transport


def recompute_plan(
    plan: dict[str, Any],
    raw: dict[str, Any],
    parameters: dict[str, Any],
    question: int,
    *,
    legacy: bool = False,
) -> dict[str, Any]:
    orders, transport = decode_plan(plan, raw)
    consumption = np.asarray([CONSUMPTION[t] for t in raw["types"]])
    price = np.asarray([PRICES[t] for t in raw["types"]])
    actual_supply = orders * parameters["ratio"]
    dispatched = transport.sum(axis=1)
    loads = transport.sum(axis=0)
    raw_loss = float(np.sum(transport * parameters["loss"][None, :]))
    production = float(np.sum(transport * (1 - parameters["loss"][None, :]) / consumption[:, None]))
    expected_demand = production if question == 4 else 28200.0
    demand = float(plan.get("weekly_demand_product_m3", expected_demand))
    initial = float(plan.get("initial_inventory_product_m3", 2 * demand))
    floor = float(plan.get("inventory_floor_product_m3", 2 * demand))
    inventory = []
    stock = initial
    for _ in range(24):
        stock += production - demand
        inventory.append(float(stock))
    checks = {
        "supplier_all_supply_transport_balance_max_abs_m3": residual(
            float(np.max(np.abs(actual_supply - dispatched)))
        ),
        "order_capacity_max_excess_m3": residual(
            float(np.max(orders - parameters["order_capacity"]))
        ),
        "supply_capacity_max_excess_m3": residual(
            float(np.max(actual_supply - parameters["supply_capacity"]))
        ),
        "carrier_capacity_max_excess_m3": residual(float(np.max(loads - 6000.0))),
        "weekly_production_margin_product_m3": residual(production - demand, 0, "GE"),
        "declared_demand_matches_question_product_m3": residual(demand - expected_demand, 0, "EQ"),
        "declared_initial_inventory_two_demand_weeks": residual(initial - 2 * demand, 0, "EQ"),
        "declared_inventory_floor_two_demand_weeks": residual(floor - 2 * demand, 0, "EQ"),
        "minimum_nominal_inventory_margin_product_m3": residual(
            min(inventory) - 2 * demand, 0, "GE"
        ),
        "minimum_initial_inventory_margin_product_m3": residual(initial - 2 * demand, 0, "GE"),
    }
    material = {
        typ: float(actual_supply[np.asarray(raw["types"]) == typ].sum()) for typ in CONSUMPTION
    }
    return {
        "feasible": passes(checks),
        "constraint_residuals": checks,
        "metric_values": {
            "weekly_all_supply_purchase_cost": float(actual_supply @ price),
            "weekly_transport_loss_m3": raw_loss,
            "weekly_effective_received_product_m3": production,
            "weekly_received_raw_m3": float(dispatched.sum() - raw_loss),
            "used_supplier_count": int(np.sum(orders > 1e-7)),
            "maximum_carrier_load_m3": float(loads.max()),
            "minimum_inventory_product_m3": min(inventory),
            "material_A_supply_m3": material["A"],
            "material_C_supply_m3": material["C"],
            "split_supplier_count": int(np.sum(np.sum(transport > 1e-7, axis=1) > 1)),
        },
        "inventory_trace_product_m3": inventory,
        "material_mix": material,
        "inventory_initial_stock_observed": False,
        "nominal_supply": actual_supply.tolist(),
    }


def recompute_stress(
    plan: dict[str, Any], raw: dict[str, Any], p: dict[str, Any], question: int
) -> dict[str, Any]:
    orders, transport = decode_plan(plan, raw)
    totals = transport.sum(axis=1)
    shares = np.zeros_like(transport)
    for i, total in enumerate(totals):
        if total > 1e-8:
            shares[i] = transport[i] / total
    conversion = np.asarray([CONSUMPTION[t] for t in raw["types"]])
    prices = np.asarray([PRICES[t] for t in raw["types"]])
    nominal = float(np.sum(transport * (1 - p["loss"][None, :]) / conversion[:, None]))
    demand = nominal if question == 4 else 28200.0
    records = []
    stock = 2 * demand
    for w in range(168, 216):
        if w in [168, 192]:
            stock = 2 * demand
        ratio = p["ratio"].copy()
        ordered = raw["order"][:, w] > 0
        ratio[ordered] = raw["supply"][ordered, w] / raw["order"][ordered, w]
        supply = orders * ratio
        shipment = supply[:, None] * shares
        loss = raw["loss_pct"][:, w] / 100.0
        loss = np.minimum(0.999999, np.where(loss > 0, loss, p["loss"]))
        received = float(np.sum(shipment * (1 - loss[None, :]) / conversion[:, None]))
        stock += received - demand
        excess = float(np.maximum(shipment.sum(axis=0) - 6000.0, 0.0).sum())
        deficit = float(max(0.0, 2 * demand - stock))
        records.append(
            {
                "historical_week": f"W{w + 1:03d}",
                "end_inventory_product_m3": float(stock),
                "product_received_m3": received,
                "all_supply_purchase_cost": float(supply @ prices),
                "carrier_capacity_excess_m3": excess,
                "inventory_floor_shortfall_product_m3": deficit,
            }
        )
    value = float(
        np.mean(
            [
                r["all_supply_purchase_cost"] / demand
                + 100 * r["inventory_floor_shortfall_product_m3"] / (2 * demand)
                + 100 * r["carrier_capacity_excess_m3"] / 48000.0
                for r in records
            ]
        )
    )
    residuals = {
        "minimum_historical_inventory_margin_product_m3": residual(
            min(r["end_inventory_product_m3"] for r in records) - 2 * demand, 0, "GE"
        ),
        "maximum_historical_carrier_capacity_excess_m3": residual(
            max(r["carrier_capacity_excess_m3"] for r in records)
        ),
    }
    return {
        "feasible": passes(residuals),
        "constraint_residuals": residuals,
        "evaluation_boundary": "DEVELOPMENT_CONDITIONAL_HISTORICAL_STRESS",
        "metric_values": {
            "validation_penalized_cost_per_effective_m3": value,
            "mean_all_supply_purchase_cost": float(
                np.mean([r["all_supply_purchase_cost"] for r in records])
            ),
            "stress_inventory_floor_violation_weeks": sum(
                r["inventory_floor_shortfall_product_m3"] > 0.05 for r in records
            ),
            "stress_carrier_capacity_violation_weeks": sum(
                r["carrier_capacity_excess_m3"] > 0.05 for r in records
            ),
        },
        "weeks": records,
        "stochastic_feasibility_established": False,
    }


def compare_parameters(
    output: dict[str, Any], raw: dict[str, Any], p: dict[str, Any]
) -> dict[str, Any]:
    emitted = output.get("parameters")
    if emitted is None:
        return {"producer_parameter_bundle_available": residual(0, 1, "EQ", 0)}
    checks = {}
    for key, expected in [
        ("supplier_delivery_ratios", p["ratio"]),
        ("supplier_order_capacity_m3", p["order_capacity"]),
        ("supplier_supply_capacity_m3", p["supply_capacity"]),
        ("carrier_loss_rates", p["loss"]),
    ]:
        ids = raw["carrier_ids"] if key == "carrier_loss_rates" else raw["supplier_ids"]
        actual = emitted.get(key, {})
        checks[f"{key}_entity_count_difference"] = residual(len(actual) - len(ids), 0, "EQ", 0)
        if set(actual) != set(ids):
            raise ValueError("PARAMETER_ENTITY_SET_INVALID")
        checks[f"{key}_max_abs_difference"] = residual(
            max(abs(float(actual[i]) - float(v)) for i, v in zip(ids, expected, strict=True)),
            0,
            "EQ",
            1e-7,
        )
    checks["purchase_price_mapping_equal"] = residual(
        float(emitted.get("purchase_prices") == PRICES), 1, "EQ", 0
    )
    checks["consumption_mapping_equal"] = residual(
        float(emitted.get("consumption") == CONSUMPTION), 1, "EQ", 0
    )
    checks["carrier_capacity_equal"] = residual(
        float(emitted.get("carrier_capacity_m3", -1)), 6000, "EQ", 0
    )
    checks["supplier_material_mapping_equal"] = residual(
        float(
            emitted.get("supplier_types")
            == dict(zip(raw["supplier_ids"], raw["types"], strict=True))
        ),
        1,
        "EQ",
        0,
    )
    return checks


def check_q1(output: dict[str, Any], raw: dict[str, Any], p: dict[str, Any]) -> dict[str, Any]:
    features = p["features"]
    normalized = (features - features.min(axis=0)) / np.maximum(np.ptp(features, axis=0), 1e-8)
    scores = normalized @ np.asarray([0.45, 0.25, 0.15, 0.15])
    ranking = sorted(range(len(scores)), key=lambda i: (-float(scores[i]), raw["supplier_ids"][i]))[
        :50
    ]
    wanted = [raw["supplier_ids"][i] for i in ranking]
    rows = output.get("question_1", {}).get("top_50_suppliers", [])
    checks = {
        "top50_count": residual(len(rows), 50, "EQ", 0),
        "raw_recomputed_top50_ranking_matches": residual(
            float([r.get("supplier_id") for r in rows] == wanted), 1, "EQ", 0
        ),
        "score_max_abs_difference": residual(
            max(
                [
                    abs(float(r["importance_score"]) - scores[i])
                    for r, i in zip(rows, ranking, strict=False)
                ]
                + [0.0]
            ),
            0,
            "EQ",
            1e-7,
        ),
    }
    reference = set(wanted)
    overlaps = []
    for factor in range(4):
        for scale in [0.8, 1.2]:
            weights = np.asarray([0.45, 0.25, 0.15, 0.15])
            weights[factor] *= scale
            weights /= weights.sum()
            score = normalized @ weights
            top = sorted(
                range(len(score)), key=lambda i: (-float(score[i]), raw["supplier_ids"][i])
            )[:50]
            overlaps.append(len(reference & {raw["supplier_ids"][i] for i in top}))
    emitted = (
        output.get("question_1", {})
        .get("weight_sensitivity", {})
        .get("minimum_top50_overlap_count")
    )
    if emitted is not None:
        checks["weight_sensitivity_overlap_equal"] = residual(emitted, min(overlaps), "EQ", 0)
    return {
        "feasible": passes(checks),
        "constraint_residuals": checks,
        "metric_values": {
            "top50_supplier_count": len(rows),
            "minimum_weight_sensitivity_top50_overlap": min(overlaps),
        },
    }


def check_tables(root: Path, output: dict[str, Any], raw: dict[str, Any]) -> dict[str, Any]:
    failures = []
    verified = 0
    artifacts = output.get("plan_table_artifacts", [])
    for artifact in artifacts:
        path = (root / artifact["artifact_id"]).resolve()
        path.relative_to(root.resolve())
        if not path.is_file() or digest(path) != artifact["sha256"]:
            failures.append("TABLE_HASH_MISMATCH")
            continue
        question = artifact["question"]
        kind = artifact["kind"]
        orders, transport = decode_plan(
            output[f"question_{question}"]["weekly_plan_repeated_for_24_weeks"], raw
        )
        with path.open(encoding="utf-8", newline="") as stream:
            rows = list(csv.reader(stream))
        expected = (
            ["supplier_id"]
            + ([] if kind == "orders" else ["carrier_id"])
            + [f"FUTURE_W{w:02d}" for w in range(1, 25)]
        )
        if not rows or rows[0] != expected:
            failures.append("TABLE_HEADER_INVALID")
            continue
        seen = set()
        for row in rows[1:]:
            sid = row[0]
            i = raw["supplier_ids"].index(sid)
            if kind == "orders":
                identifier = (sid,)
                value = orders[i]
                numeric = row[1:]
            else:
                cid = row[1]
                j = raw["carrier_ids"].index(cid)
                identifier = (sid, cid)
                value = transport[i, j]
                numeric = row[2:]
            if (
                identifier in seen
                or len(numeric) != 24
                or any(abs(float(v) - value) > 1e-8 for v in numeric)
            ):
                failures.append("TABLE_VALUE_OR_DUPLICATE_INVALID")
            seen.add(identifier)
        expected_rows = len(raw["supplier_ids"]) * (
            len(raw["carrier_ids"]) if kind == "transport" else 1
        )
        if len(seen) != expected_rows:
            failures.append("TABLE_ENTITY_COVERAGE_INVALID")
        verified += 1
    checks = {
        "derived_plan_table_count": residual(verified, 6, "EQ", 0),
        "derived_plan_table_failure_count": residual(len(failures), 0, "EQ", 0),
        "official_named_templates_provided": residual(0, 1, "EQ", 0),
    }
    return {
        "feasible": False,
        "constraint_residuals": checks,
        "failures": sorted(set(failures)),
        "derived_tables_consistent": verified == 6 and not failures,
        "metric_values": {"verified_derived_plan_table_count": verified},
        "limitation": (
            "Derived CSVs do not complete official named attachment A/B template submission."
        ),
    }


def check(root: Path, run_id: str, model_output: Path | None = None) -> dict[str, Any]:
    path = root / "runs" / run_id / "output.json" if model_output is None else model_output
    path = path.resolve()
    path.relative_to(root.resolve())
    output = json.loads(path.read_text(encoding="utf-8"))
    output_hash = digest(path)
    raw = read_inputs(root)
    p = derive_parameters(raw, output["candidate_id"], int(output["seed"]))
    legacy = output.get("implementation_revision") != "RC8_2021_SCIENTIFIC_REPAIR_V6"
    parameters = compare_parameters(output, raw, p)
    assumptions = root / "models/assumptions_and_symbols.json"
    assumption_hash = digest(assumptions) if assumptions.is_file() else None
    if not legacy:
        parameters["assumption_artifact_hash_equal"] = residual(
            float(
                output.get("assumption_artifact_sha256") == assumption_hash
                and assumption_hash is not None
            ),
            1,
            "EQ",
            0,
        )
    requirements = {}
    plans = {}
    stress = {}
    q1 = check_q1(output, raw, p)
    requirements["REQ-Q1-IMPORTANCE-MODEL"] = q1
    requirements["REQ-Q1-TOP50"] = q1
    optimistic = np.sort(
        p["supply_capacity"]
        * (1 - p["loss"].min())
        / np.asarray([CONSUMPTION[t] for t in raw["types"]])
    )[::-1]
    lower_bound = int(np.searchsorted(np.cumsum(optimistic), 28200.0, side="left") + 1)
    for question in [2, 3, 4]:
        key = f"question_{question}"
        record = output.get(key, {})
        if "weekly_plan_repeated_for_24_weeks" not in record:
            plans[key] = {
                "feasible": False,
                "constraint_residuals": {"plan_available": residual(0, 1, "EQ", 0)},
                "metric_values": {},
            }
            stress[key] = plans[key]
        else:
            plan = record["weekly_plan_repeated_for_24_weeks"]
            plans[key] = recompute_plan(plan, raw, p, question, legacy=legacy)
            stress[key] = recompute_stress(plan, raw, p, question)
            stress[key]["metric_values"].update(plans[key]["metric_values"])
            if question == 4:
                capacity = plans[key]["metric_values"]["weekly_effective_received_product_m3"]
                plans[key]["metric_values"].update(
                    {
                        "conditional_capacity_product_m3": capacity,
                        "capacity_increase_product_m3": capacity - 28200.0,
                    }
                )
            if not legacy:
                plans[key]["constraint_residuals"].update(parameters)
                emitted = record.get("material_mix", {})
                plans[key]["constraint_residuals"]["reported_purchase_cost_error"] = residual(
                    float(emitted.get("all_supply_purchase_cost", -1))
                    - plans[key]["metric_values"]["weekly_all_supply_purchase_cost"],
                    0,
                    "EQ",
                )
                plans[key]["constraint_residuals"]["reported_transport_loss_error"] = residual(
                    float(emitted.get("transport_loss_m3", -1))
                    - plans[key]["metric_values"]["weekly_transport_loss_m3"],
                    0,
                    "EQ",
                )
                plans[key]["feasible"] = passes(plans[key]["constraint_residuals"])
        for suffix in ["ORDER-PLAN", "TRANSPORT-PLAN"]:
            requirements[f"REQ-Q{question}-{suffix}"] = plans[key]
        if question in [2, 3]:
            requirements[f"REQ-Q{question}-EFFECT"] = {
                **stress[key],
                "status": "DEVELOPMENT_STRESS_RESULT_WITH_FAILURES_PRESERVED",
            }
    q2 = plans["question_2"]
    count = q2["metric_values"].get("used_supplier_count", 0)
    count_checks = {
        **q2["constraint_residuals"],
        "actual_used_count_equals_optimistic_lower_bound": residual(count, lower_bound, "EQ", 0),
        "reported_minimum_count_equals_actual_used": residual(
            output["question_2"].get("minimum_supplier_count", 0), count, "EQ", 0
        ),
    }
    requirements["REQ-Q2-MINIMUM-SUPPLIERS"] = {
        "feasible": passes(count_checks),
        "constraint_residuals": count_checks,
        "metric_values": {
            **q2["metric_values"],
            "minimum_supplier_count": count,
            "optimistic_cardinality_lower_bound": lower_bound,
        },
        "independent_optimality_certificate": "FEASIBLE_COUNT_MATCHES_RAW_OPTIMISTIC_LOWER_BOUND"
        if passes(count_checks)
        else "NOT_ESTABLISHED",
        "solver_reported_dual_bound": output["question_2"].get("cardinality_dual_bound"),
        "solver_reported_gap": output["question_2"].get("cardinality_gap"),
    }
    if q2["feasible"]:
        requirements["REQ-Q2-MINIMUM-SUPPLIERS"]["optimality_certificate"] = {
            "proof_kind": "INDEPENDENT_BOUND",
            "lower_bound": float(lower_bound),
            "upper_bound": float(count),
            "objective_value": float(count),
            "tolerance": 1e-6,
            "scope": "REGISTERED_CAPACITY_SCENARIO",
        }
    requirements["REQ-Q3-MATERIAL-PREFERENCE"] = plans["question_3"]
    q4 = {**plans["question_4"], "metric_values": dict(plans["question_4"]["metric_values"])}
    q4["metric_values"].update(
        {
            "conditional_capacity_product_m3": q4["metric_values"].get(
                "weekly_effective_received_product_m3", 0
            ),
            "capacity_increase_product_m3": q4["metric_values"].get(
                "weekly_effective_received_product_m3", 0
            )
            - 28200.0,
        }
    )
    q4["optimality_independently_proven"] = False
    q4["capacity_scope"] = (
        "FEASIBLE_STATIONARY_CONDITIONAL_PLAN; LP optimum from solver stage not"
        " independently dual-certified."
    )
    requirements["REQ-Q4-CAPACITY-INCREASE"] = q4
    inventory_checks = {
        f"{name}_{k}": v
        for name, record in plans.items()
        for k, v in record["constraint_residuals"].items()
        if "inventory" in k or "production" in k
    }
    if len(plans) != 3 or any(not p["feasible"] for p in plans.values()):
        inventory_checks["all_nominal_plans_feasible"] = residual(0, 1, "EQ", 0)
    requirements["REQ-INVENTORY-PRODUCTION"] = {
        "feasible": passes(inventory_checks),
        "constraint_residuals": inventory_checks,
        "historical_stress_feasible": all(s["feasible"] for s in stress.values()),
        "metric_values": {
            **q2["metric_values"],
            "stress_inventory_floor_violation_weeks": stress["question_2"]
            .get("metric_values", {})
            .get("stress_inventory_floor_violation_weeks", 48),
            "stress_carrier_capacity_violation_weeks": stress["question_2"]
            .get("metric_values", {})
            .get("stress_carrier_capacity_violation_weeks", 48),
        },
    }
    requirements["REQ-TRANSPORT-BUSINESS-RULES"] = {
        "feasible": all(p["feasible"] for p in plans.values()),
        "constraint_residuals": {
            f"{name}_{key}": value
            for name, record in plans.items()
            for key, value in record["constraint_residuals"].items()
            if "carrier_capacity" in key or "supply_transport" in key
        },
        "soft_preference": "One carrier per supplier is preferred, not enforced as a hard rule.",
        "metric_values": dict(q2["metric_values"]),
    }
    tables = check_tables(root, output, raw)
    requirements["REQ-OUTPUT-ATTACHMENT-A"] = tables
    requirements["REQ-OUTPUT-ATTACHMENT-B"] = tables
    metric_values = stress["question_2"].get("metric_values", {})
    if not legacy and "validation_penalized_cost_per_effective_m3" in metric_values:
        score = metric_values["validation_penalized_cost_per_effective_m3"]
        metric_consistent = abs(float(output["metric_value"]) - score) <= 1e-7
    else:
        metric_consistent = None
    metric_values = dict(metric_values)
    for req, record in list(requirements.items()):
        prefix = (
            "ATTACHMENTS"
            if req.startswith("REQ-OUTPUT-")
            else (
                "Q1"
                if req.startswith("REQ-Q1")
                else "Q3"
                if "Q3" in req
                else "Q4"
                if "Q4" in req
                else "Q2"
            )
        )
        scoped_metrics = {
            f"{prefix}.{key}": value for key, value in record.get("metric_values", {}).items()
        }
        requirements[req] = {**record, "metric_values": scoped_metrics}
        metric_values.update(scoped_metrics)
    return {
        "schema_version": "c2021-independent-check/v1",
        "run_id": run_id,
        "output_sha256": output_hash,
        "input_hashes": raw["input_hashes"],
        "assumption_artifact_sha256": assumption_hash,
        "checker_sha256": digest(Path(__file__)),
        "legacy_baseline_recalculation": legacy,
        "requirements": requirements,
        "plans": plans,
        "historical_stress": stress,
        "metric_values": metric_values,
        "producer_primary_metric_matches_independent": metric_consistent,
        "parameter_residuals": parameters,
        "optimistic_cardinality_capacity_prefix_product_m3": np.cumsum(optimistic).tolist(),
        "scientific_coverage_status": "PARTIAL_SCIENTIFIC_COVERAGE",
        "whole_problem_scientifically_complete": False,
        "independence": "SEPARATE_FIRST_PARTY_IMPLEMENTATION_NO_PRODUCER_OR_HELPER_IMPORT",
        "limitations": [
            "Independent Python recomputation is not an independent agent audit.",
            (
                "Q2 count can be certified only when feasible count matches the raw "
                "optimistic lower bound."
            ),
            "Lexicographic and Q4 optimality solver claims are not independent dual certificates.",
            (
                "Future inventory assumes two demand-weeks of initial stock; historical"
                " stress is conditional Development evidence."
            ),
            "No stochastic feasibility or official attachment-template completion is claimed.",
        ],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--case-root", required=True)
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--model-output")
    args = parser.parse_args()
    root = Path(args.case_root).resolve()
    output_path = Path(args.output)
    if not output_path.is_absolute():
        output_path = root / output_path
    output_path = output_path.resolve()
    output_path.relative_to(root)
    model_output = None if args.model_output is None else (root / args.model_output).resolve()
    result = check(root, args.run_id, model_output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(result, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "run_id": args.run_id,
                "output_sha256": result["output_sha256"],
                "scientific_coverage_status": result["scientific_coverage_status"],
            },
            sort_keys=True,
        )
    )
    return 0  # Successful checking can contain negative scientific findings.


if __name__ == "__main__":
    raise SystemExit(main())
