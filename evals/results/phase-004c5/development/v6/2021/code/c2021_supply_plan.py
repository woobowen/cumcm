#!/usr/bin/env python3
"""First-party Development scientific repair for CUMCM 2021 problem C.

The program uses only the official case inputs.  Candidate generation and scoring use
the preregistered train/validation periods; W217-W240 are never inspected here.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
from c2021_feasibility import CONSUMPTION, PURCHASE_PRICE, verify_plan
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, linprog, milp

CANDIDATES = {
    "BASELINE_MEAN_GREEDY": {
        "capacity_quantile": 0.75,
        "ratio_method": "MEAN",
        "loss_method": "MEAN",
        "allocation": "LEXICOGRAPHIC_MILP_LP",
    },
    "ROBUST_QUANTILE_LEXICOGRAPHIC": {
        "capacity_quantile": 0.75,
        "ratio_method": "Q20",
        "loss_method": "Q75",
        "allocation": "LEXICOGRAPHIC_MILP_LP",
    },
    "SCENARIO_CVAR_PORTFOLIO": {
        "capacity_quantile": 0.75,
        "ratio_method": "BOOTSTRAP_LOWER_TAIL",
        "loss_method": "UPPER_TAIL_MEAN",
        "allocation": "LEXICOGRAPHIC_MILP_LP",
    },
}
SUPPLIER_FILE = "raw/case_files/附件1 近5年402家供应商的相关数据.xlsx"
CARRIER_FILE = "raw/case_files/附件2 近5年8家转运商的相关数据.xlsx"
TRAIN = slice(0, 168)
VALIDATION = slice(168, 216)
WEEKLY_DEMAND = 28200.0
CARRIER_CAPACITY = 6000.0
EPS = 1e-8


@dataclass(frozen=True)
class CaseData:
    supplier_ids: list[str]
    supplier_types: list[str]
    order: np.ndarray
    supply: np.ndarray
    carrier_ids: list[str]
    loss_pct: np.ndarray


def _sheet_matrix(
    path: Path, sheet: str, first_numeric_column: int
) -> tuple[list[str], list[str], np.ndarray]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet]
    rows = [
        row
        for row in worksheet.iter_rows(max_col=first_numeric_column + 216, values_only=True)
        if row[0] is not None
    ]
    headers = [str(value) for value in rows[0]]
    identifiers = [str(row[0]) for row in rows[1:]]
    matrix = np.asarray(
        [[float(value) for value in row[first_numeric_column:]] for row in rows[1:]],
        dtype=float,
    )
    workbook.close()
    return headers, identifiers, matrix


def load_case_data(case_root: Path) -> CaseData:
    verify_input_hashes(case_root)
    supplier_path = case_root / SUPPLIER_FILE
    carrier_path = case_root / CARRIER_FILE
    order_headers, order_ids, order = _sheet_matrix(supplier_path, "企业的订货量（m³）", 2)
    supply_headers, supply_ids, supply = _sheet_matrix(supplier_path, "供应商的供货量（m³）", 2)
    workbook = load_workbook(supplier_path, read_only=True, data_only=True)
    worksheet = workbook["企业的订货量（m³）"]
    supplier_types = [
        str(row[1])
        for row in list(worksheet.iter_rows(max_col=2, values_only=True))[1:]
        if row[0] is not None
    ]
    workbook.close()
    carrier_headers, carrier_ids, loss_pct = _sheet_matrix(carrier_path, "运输损耗率（%）", 1)
    expected_weeks = [f"W{week:03d}" for week in range(1, 217)]
    if (
        order_ids != supply_ids
        or order_headers[2:] != expected_weeks
        or supply_headers[2:] != expected_weeks
        or carrier_headers[1:] != expected_weeks
        or order.shape != (402, 216)
        or supply.shape != (402, 216)
        or loss_pct.shape != (8, 216)
        or any(material not in CONSUMPTION for material in supplier_types)
        or np.any(~np.isfinite(order))
        or np.any(~np.isfinite(supply))
        or np.any(~np.isfinite(loss_pct))
        or np.any(order < 0)
        or np.any(supply < 0)
        or np.any(loss_pct < 0)
    ):
        raise ValueError("OFFICIAL_INPUT_SCHEMA_INVALID")
    return CaseData(order_ids, supplier_types, order, supply, carrier_ids, loss_pct)


def _safe_quantile(values: np.ndarray, quantile: float, default: float = 0.0) -> float:
    values = values[np.isfinite(values)]
    return default if values.size == 0 else float(np.quantile(values, quantile))


def estimate_suppliers(data: CaseData, candidate_id: str, seed: int) -> dict[str, np.ndarray]:
    config = CANDIDATES[candidate_id]
    rng = np.random.default_rng(seed)
    order = data.order[:, TRAIN]
    supply = data.supply[:, TRAIN]
    count = len(data.supplier_ids)
    order_capacity = np.zeros(count)
    ratio = np.zeros(count)
    reliability = np.zeros(count)
    stability = np.zeros(count)
    activity = np.zeros(count)
    mean_product = np.zeros(count)
    for index in range(count):
        ordered = order[index] > 0
        positive_orders = order[index, ordered]
        ratios = np.clip(supply[index, ordered] / positive_orders, 0.0, 2.0)
        order_capacity[index] = _safe_quantile(positive_orders, float(config["capacity_quantile"]))
        if ratios.size:
            if config["ratio_method"] == "MEAN":
                ratio[index] = float(np.mean(ratios))
            elif config["ratio_method"] == "Q20":
                ratio[index] = _safe_quantile(ratios, 0.20)
            else:
                boot = np.asarray(
                    [np.mean(rng.choice(ratios, size=ratios.size, replace=True)) for _ in range(64)]
                )
                cutoff = np.quantile(boot, 0.20)
                ratio[index] = float(np.mean(boot[boot <= cutoff]))
            reliability[index] = float(np.mean(ratios >= 0.95))
        positives = supply[index, supply[index] > 0]
        activity[index] = float(np.mean(supply[index] > 0))
        if positives.size:
            stability[index] = 1.0 / (
                1.0 + float(np.std(positives)) / max(float(np.mean(positives)), EPS)
            )
        mean_product[index] = (
            float(np.mean(supply[index])) / CONSUMPTION[data.supplier_types[index]]
        )
    ratio = np.clip(ratio, 0.0, 1.5)
    delivered_capacity = order_capacity * ratio
    return {
        "order_capacity": order_capacity,
        "delivery_ratio": ratio,
        "delivered_capacity": delivered_capacity,
        "reliability": reliability,
        "stability": stability,
        "activity": activity,
        "mean_product": mean_product,
    }


def estimate_carrier_losses(data: CaseData, candidate_id: str) -> np.ndarray:
    method = CANDIDATES[candidate_id]["loss_method"]
    losses = np.zeros(len(data.carrier_ids))
    for index, row in enumerate(data.loss_pct[:, TRAIN]):
        positive = row[row > 0] / 100.0
        if positive.size == 0:
            raise ValueError(f"CARRIER_POSITIVE_LOSS_MISSING:{data.carrier_ids[index]}")
        if method == "MEAN":
            losses[index] = float(np.mean(positive))
        elif method == "Q75":
            losses[index] = _safe_quantile(positive, 0.75)
        else:
            cutoff = np.quantile(positive, 0.80)
            losses[index] = float(np.mean(positive[positive >= cutoff]))
    return losses


def importance_scores(
    data: CaseData, estimates: dict[str, np.ndarray]
) -> tuple[np.ndarray, list[dict[str, Any]]]:
    feature_names = ["mean_product", "reliability", "stability", "activity"]
    weights = np.asarray([0.45, 0.25, 0.15, 0.15])
    normalized: list[np.ndarray] = []
    for name in feature_names:
        values = estimates[name]
        span = float(np.max(values) - np.min(values))
        normalized.append(
            np.zeros_like(values) if span <= EPS else (values - np.min(values)) / span
        )
    score = np.vstack(normalized).T @ weights
    ranked = sorted(
        range(len(data.supplier_ids)), key=lambda i: (-float(score[i]), data.supplier_ids[i])
    )
    top50 = [
        {
            "rank": rank,
            "supplier_id": data.supplier_ids[index],
            "material_type": data.supplier_types[index],
            "importance_score": round(float(score[index]), 9),
        }
        for rank, index in enumerate(ranked[:50], 1)
    ]
    return score, top50


def _effective_coefficients(data: CaseData, losses: np.ndarray) -> np.ndarray:
    consumption = np.asarray([CONSUMPTION[value] for value in data.supplier_types])
    return (1.0 - losses[None, :]) / consumption[:, None]


OFFICIAL_HASHES = {
    SUPPLIER_FILE: "1b93a7598b8f0489e3c054439967bca654a83654837908c833752617cf950e1b",
    CARRIER_FILE: "29e1499b133aa88d765902081bfb827f78ffb4c091f5f75341fc37d844344685",
}


def verify_input_hashes(case_root: Path) -> dict[str, str]:
    actual = {
        name: hashlib.sha256((case_root / name).read_bytes()).hexdigest()
        for name in OFFICIAL_HASHES
    }
    if actual != OFFICIAL_HASHES:
        raise ValueError("OFFICIAL_WORKBOOK_HASH_MISMATCH")
    return actual


def _number(value: Any) -> float | None:
    return float(value) if value is not None and math.isfinite(float(value)) else None


def _stage_record(result: Any, name: str, *, integer: bool) -> dict[str, Any]:
    return {
        "stage": name,
        "solver": "SCIPY_HIGHS_MILP" if integer else "SCIPY_HIGHS_LP",
        "status_code": int(result.status),
        "message": str(result.message),
        "has_incumbent": result.x is not None,
        "optimal": bool(result.success),
        "incumbent_objective": _number(getattr(result, "fun", None)),
        "dual_bound": _number(getattr(result, "mip_dual_bound", None))
        if integer
        else (_number(getattr(result, "fun", None)) if result.success else None),
        "relative_gap": _number(getattr(result, "mip_gap", None))
        if integer
        else (0.0 if result.success else None),
        "node_count": getattr(result, "mip_node_count", None),
    }


def _linear_rows(
    data: CaseData,
    estimates: dict[str, np.ndarray],
    losses: np.ndarray,
    *,
    integer: bool,
    demand: float | None,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    n, m = len(data.supplier_ids), len(data.carrier_ids)
    offset = n if integer else 0
    rows = np.zeros((n + m + (demand is not None), offset + n * m))
    upper = np.r_[
        np.zeros(n) if integer else estimates["delivered_capacity"],
        np.full(m, CARRIER_CAPACITY),
        [] if demand is None else [demand],
    ]
    lower = np.full(len(upper), -np.inf)
    for i in range(n):
        rows[i, offset + i * m : offset + (i + 1) * m] = 1.0
        if integer:
            rows[i, i] = -estimates["delivered_capacity"][i]
    for j in range(m):
        rows[n + j, offset + j :: m] = 1.0
    if demand is not None:
        rows[-1, offset:] = _effective_coefficients(data, losses).ravel()
        lower[-1] = demand  # No excess production may game A-max or loss priorities.
    return rows, lower, upper


def optimize_lexicographic(
    data: CaseData,
    estimates: dict[str, np.ndarray],
    losses: np.ndarray,
    question: int,
    time_limit: float = 100.0,
) -> dict[str, Any]:
    """Conditional stationary optimization; preserve every solver-stage termination."""
    n, m = len(data.supplier_ids), len(data.carrier_ids)
    integer = question == 2
    offset = n if integer else 0
    size = offset + n * m
    consumption = np.asarray([CONSUMPTION[t] for t in data.supplier_types])
    price = np.repeat([PURCHASE_PRICE[t] for t in data.supplier_types], m)
    purchase = np.r_[np.zeros(offset), price]
    loss = np.r_[np.zeros(offset), np.tile(losses, n)]
    effective = _effective_coefficients(data, losses).ravel()
    rows, lower, upper = _linear_rows(
        data, estimates, losses, integer=integer, demand=None if question == 4 else WEEKLY_DEMAND
    )
    bounds_upper = np.full(size, np.inf)
    if integer:
        bounds_upper[:n] = (estimates["delivered_capacity"] > EPS).astype(float)
    objectives: list[tuple[str, np.ndarray]]
    if question == 2:
        objectives = [
            ("MINIMUM_SUPPLIER_CARDINALITY", np.r_[np.ones(n), np.zeros(n * m)]),
            ("MINIMUM_ALL_SUPPLY_PURCHASE_COST", purchase),
            ("MINIMUM_TRANSPORT_LOSS_AT_FROZEN_COST", loss),
        ]
    elif question == 3:
        c_material = np.repeat([float(t == "C") for t in data.supplier_types], m)
        a_material = -np.repeat([float(t == "A") for t in data.supplier_types], m)
        objectives = [
            ("MINIMUM_C_MATERIAL", c_material),
            ("MAXIMUM_A_MATERIAL", a_material),
            ("MINIMUM_ALL_SUPPLY_PURCHASE_COST", purchase),
            ("MINIMUM_TRANSPORT_LOSS_AT_FROZEN_COST", loss),
        ]
    else:
        objectives = [
            ("MAXIMUM_CONDITIONAL_WEEKLY_CAPACITY", -effective),
            ("MINIMUM_ALL_SUPPLY_PURCHASE_COST", purchase),
            ("MINIMUM_TRANSPORT_LOSS_AT_FROZEN_COST", loss),
        ]
    stages: list[dict[str, Any]] = []
    current = None
    for name, objective in objectives:
        if integer:
            result = milp(
                objective,
                integrality=np.r_[np.ones(n), np.zeros(n * m)],
                bounds=Bounds(np.zeros(size), bounds_upper),
                constraints=LinearConstraint(rows, lower, upper),
                options={"time_limit": time_limit, "mip_rel_gap": 0.0},
            )
        else:
            equal = np.isfinite(lower) & (lower == upper)
            inequality = ~equal
            result = linprog(
                objective,
                A_ub=rows[inequality],
                b_ub=upper[inequality],
                A_eq=rows[equal] if np.any(equal) else None,
                b_eq=upper[equal] if np.any(equal) else None,
                bounds=(0.0, None),
                method="highs",
                options={"time_limit": time_limit},
            )
        stage = _stage_record(result, name, integer=integer)
        stages.append(stage)
        if result.x is None:
            break
        current = np.asarray(result.x)
        if not result.success:
            # A timed incumbent remains evidence but cannot authorize later lexicographic stages.
            break
        value = float(objective @ current)
        tolerance = 0.0 if name == "MINIMUM_SUPPLIER_CARDINALITY" else max(1e-7, abs(value) * 1e-9)
        bound = float(round(value)) if name == "MINIMUM_SUPPLIER_CARDINALITY" else value + tolerance
        stage["frozen_upper_bound"] = bound
        stage["lexicographic_absolute_tolerance"] = tolerance
        rows = np.vstack([rows, objective])
        lower = np.r_[lower, -np.inf]
        upper = np.r_[upper, bound]
    allocation = None if current is None else np.maximum(current[offset:].reshape(n, m), 0.0)
    optimistic = estimates["delivered_capacity"] * (1 - float(np.min(losses))) / consumption
    sorted_capacity = np.sort(optimistic)[::-1]
    cumulative = np.cumsum(sorted_capacity)
    lower_bound = int(np.searchsorted(cumulative, WEEKLY_DEMAND, side="left") + 1)
    return {
        "allocation": allocation,
        "stages": stages,
        "lexicographic_complete": len(stages) == len(objectives)
        and all(s["optimal"] for s in stages),
        "candidate_supplier_pool_count": int(np.sum(estimates["delivered_capacity"] > EPS)),
        "actual_used_supplier_count": None
        if allocation is None
        else int(np.sum(allocation.sum(axis=1) > 1e-7)),
        "optimistic_cardinality_lower_bound": lower_bound
        if cumulative[-1] >= WEEKLY_DEMAND
        else None,
        "optimistic_top_k_capacity_product_m3": cumulative.tolist(),
        "optimistic_bound_ignores_individual_carrier_capacity": True,
    }


def sparse_plan(
    data: CaseData, estimates: dict[str, np.ndarray], allocation: np.ndarray, demand: float
) -> dict[str, Any]:
    supplied = np.sum(allocation, axis=1)
    ratio = estimates["delivery_ratio"]
    return {
        "orders": [
            {"supplier_id": data.supplier_ids[i], "volume_m3": float(supplied[i] / ratio[i])}
            for i in range(len(supplied))
            if supplied[i] > EPS and ratio[i] > EPS
        ],
        "transport": [
            {
                "supplier_id": data.supplier_ids[i],
                "carrier_id": data.carrier_ids[j],
                "volume_m3": float(allocation[i, j]),
            }
            for i in range(len(supplied))
            for j in range(len(data.carrier_ids))
            if allocation[i, j] > EPS
        ],
        "repeat_for_weeks": list(range(1, 25)),
        "weekly_demand_product_m3": float(demand),
        "initial_inventory_product_m3": float(2 * demand),
        "inventory_floor_product_m3": float(2 * demand),
        "inventory_unit": "PRODUCT_EQUIVALENT_M3",
        "inventory_assumption": "ALL_MATERIALS_SUBSTITUTE_AT_DECLARED_CONSUMPTION_FACTORS",
    }


def historical_stress(
    data: CaseData,
    plan: dict[str, Any],
    estimates: dict[str, np.ndarray],
    losses: np.ndarray,
    *,
    ratio_scale: float = 1.0,
    loss_addition: float = 0.0,
    carrier_scale: float = 1.0,
) -> dict[str, Any]:
    """Two chronological 24-week Development replays; all supply is paid and routed."""
    ids = {sid: i for i, sid in enumerate(data.supplier_ids)}
    carriers = {cid: j for j, cid in enumerate(data.carrier_ids)}
    n, m = len(ids), len(carriers)
    orders = np.zeros(n)
    planned = np.zeros((n, m))
    for row in plan["orders"]:
        orders[ids[row["supplier_id"]]] += row["volume_m3"]
    for row in plan["transport"]:
        planned[ids[row["supplier_id"]], carriers[row["carrier_id"]]] += row["volume_m3"]
    totals = planned.sum(axis=1)
    shares = np.divide(
        planned, totals[:, None], out=np.zeros_like(planned), where=totals[:, None] > EPS
    )
    consumption = np.asarray([CONSUMPTION[t] for t in data.supplier_types])
    price = np.asarray([PURCHASE_PRICE[t] for t in data.supplier_types])
    demand = plan["weekly_demand_product_m3"]
    records: list[dict[str, Any]] = []
    inventory = plan["initial_inventory_product_m3"]
    missing = 0
    for week in range(168, 216):
        if (week - 168) % 24 == 0:
            inventory = plan["initial_inventory_product_m3"]
        observed_order = data.order[:, week]
        observed_supply = data.supply[:, week]
        available = observed_order > 0
        # Do not clip realized fulfillment: surplus must be paid and tested against carrier limits.
        actual_ratio = np.divide(
            observed_supply, observed_order, out=estimates["delivery_ratio"].copy(), where=available
        )
        missing_count = int(np.sum((orders > EPS) & ~available))
        missing += missing_count
        supply = orders * actual_ratio * ratio_scale
        shipments = supply[:, None] * shares
        observed_loss = data.loss_pct[:, week] / 100.0
        used_loss = np.minimum(
            0.999999, np.where(observed_loss > 0, observed_loss, losses) + loss_addition
        )
        received = float(np.sum(shipments * (1 - used_loss[None, :]) / consumption[:, None]))
        inventory += received - demand
        capacity_excess = float(
            np.maximum(shipments.sum(axis=0) - CARRIER_CAPACITY * carrier_scale, 0).sum()
        )
        records.append(
            {
                "historical_week": f"W{week + 1:03d}",
                "stress_block": 1 + (week - 168) // 24,
                "product_received_m3": received,
                "all_supply_purchase_cost": float(supply @ price),
                "transport_loss_m3": float(np.sum(shipments * used_loss[None, :])),
                "carrier_capacity_excess_m3": capacity_excess,
                "end_inventory_product_m3": float(inventory),
                "inventory_floor_shortfall_product_m3": float(max(0, 2 * demand - inventory)),
                "production_shortage_product_m3": float(max(0, -inventory)),
                "fallback_supplier_ratio_count": missing_count,
                "zero_loss_fallback_carrier_count": int(
                    np.sum((shipments.sum(axis=0) > EPS) & (observed_loss == 0))
                ),
            }
        )
    costs = [r["all_supply_purchase_cost"] / demand for r in records]
    floor = [r["inventory_floor_shortfall_product_m3"] / (2 * demand) for r in records]
    excess = [
        r["carrier_capacity_excess_m3"] / (CARRIER_CAPACITY * carrier_scale * m) for r in records
    ]
    shortage = [max(0, demand - r["product_received_m3"]) / demand for r in records]
    # Preserve the legacy metric name for the capture interface; explicitly change its accounting.
    score = float(np.mean(costs) + 100 * np.mean(floor) + 100 * np.mean(excess))
    return {
        "evaluation_boundary": "DEVELOPMENT_CONDITIONAL_HISTORICAL_STRESS",
        "time": ["W169-W192", "W193-W216"],
        "inventory_reset": "TWO_TARGET_WEEKS_AT_START_OF_EACH_24_WEEK_BLOCK",
        "purchase_accounting": "PRICE_TIMES_ALL_REALIZED_SUPPLY_WITHOUT_RATIO_CLIPPING",
        "validation_penalized_cost_per_effective_m3": score,
        "metric_formula": (
            "mean_purchase_over_demand + "
            "100*mean_inventory_floor_shortfall_fraction + "
            "100*mean_carrier_excess_fraction"
        ),
        "normalized_purchase_cost": float(np.mean(costs)),
        "mean_shortage_fraction": float(np.mean(shortage)),
        "maximum_shortage_fraction": float(np.max(shortage)),
        "mean_inventory_floor_shortfall_fraction": float(np.mean(floor)),
        "mean_carrier_excess_fraction": float(np.mean(excess)),
        "maximum_carrier_excess_m3": max(r["carrier_capacity_excess_m3"] for r in records),
        "minimum_inventory_product_m3": min(r["end_inventory_product_m3"] for r in records),
        "inventory_floor_violation_weeks": sum(
            r["inventory_floor_shortfall_product_m3"] > 0.05 for r in records
        ),
        "carrier_capacity_violation_weeks": sum(
            r["carrier_capacity_excess_m3"] > 0.05 for r in records
        ),
        "fallback_supplier_ratio_observations": missing,
        "stochastic_feasibility_established": False,
        "conditional_historical_feasible": all(
            r["inventory_floor_shortfall_product_m3"] <= 0.05
            and r["carrier_capacity_excess_m3"] <= 0.05
            for r in records
        ),
        "weeks": records,
    }


def weight_sensitivity(
    data: CaseData, estimates: dict[str, np.ndarray], top50: list[dict[str, Any]]
) -> dict[str, Any]:
    names = ["mean_product", "reliability", "stability", "activity"]
    matrix = np.column_stack(
        [(estimates[k] - estimates[k].min()) / max(np.ptp(estimates[k]), EPS) for k in names]
    )
    base = np.asarray([0.45, 0.25, 0.15, 0.15])
    reference = {r["supplier_id"] for r in top50}
    variants = []
    for index, name in enumerate(names):
        for multiplier in [0.8, 1.2]:
            weights = base.copy()
            weights[index] *= multiplier
            weights /= weights.sum()
            scores = matrix @ weights
            ranked = sorted(
                range(len(scores)), key=lambda i: (-float(scores[i]), data.supplier_ids[i])
            )[:50]
            overlap = len(reference & {data.supplier_ids[i] for i in ranked})
            variants.append(
                {
                    "factor": name,
                    "multiplier": multiplier,
                    "weights": weights.tolist(),
                    "top50_overlap_count": overlap,
                    "top50_jaccard": overlap / (100 - overlap),
                }
            )
    return {
        "method": "ONE_FACTOR_PLUS_MINUS_20_PERCENT_RENORMALIZED",
        "variants": variants,
        "minimum_top50_overlap_count": min(v["top50_overlap_count"] for v in variants),
    }


def requirement_claims(output_path: str) -> dict[str, dict[str, Any]]:
    statements = {
        "REQ-Q1-IMPORTANCE-MODEL": (
            "Historical four-factor supplier importance with bounded weight sensitivity."
        ),
        "REQ-Q1-TOP50": "Fifty ranked suppliers under the declared min-max weights and tie rule.",
        "REQ-Q2-MINIMUM-SUPPLIERS": (
            "Conditional supplier cardinality with incumbent, dual bound and solver gap."
        ),
        "REQ-Q2-ORDER-PLAN": (
            "Conditional 24-week order plan, minimizing all-supply purchase cost "
            "after supplier count."
        ),
        "REQ-Q2-TRANSPORT-PLAN": (
            "Conditional transport loss minimization after frozen supplier count and purchase cost."
        ),
        "REQ-Q2-EFFECT": (
            "Development historical inventory and carrier-capacity stress, including failures."
        ),
        "REQ-Q3-MATERIAL-PREFERENCE": (
            "Conditional C-min then A-max priorities with reported material mix."
        ),
        "REQ-Q3-ORDER-PLAN": "Conditional question-3 24-week order plan and purchase cost.",
        "REQ-Q3-TRANSPORT-PLAN": "Conditional question-3 transport plan and loss.",
        "REQ-Q3-EFFECT": "Question-3 Development historical inventory and capacity stress.",
        "REQ-Q4-CAPACITY-INCREASE": (
            "Stationary conditional capacity bound, not a future capacity guarantee."
        ),
        "REQ-Q4-ORDER-PLAN": "Question-4 24-week order plan at increased modeled demand.",
        "REQ-Q4-TRANSPORT-PLAN": (
            "Question-4 transport and two-increased-demand-week inventory plan."
        ),
        "REQ-OUTPUT-ATTACHMENT-A": (
            "Derived numerical 24-week order tables; official named template submission incomplete."
        ),
        "REQ-OUTPUT-ATTACHMENT-B": (
            "Derived numerical 24-week transport tables; official named template "
            "submission incomplete."
        ),
        "REQ-INVENTORY-PRODUCTION": (
            "Nominal 24-week inventory and separate historical stress at each question's demand."
        ),
        "REQ-TRANSPORT-BUSINESS-RULES": (
            "All-supply routing, carrier capacity and soft supplier splitting diagnostics."
        ),
    }
    return {
        key: {
            "claim_id": f"CLAIM-C2021-{i:02d}",
            "claim_text": value,
            "evidence_artifact_ids": [output_path],
        }
        for i, (key, value) in enumerate(statements.items(), 1)
    }


def scientific_evidence(result: dict[str, Any], assumptions_hash: str) -> dict[str, Any]:
    supplier_fields = ["supplier_id", "material_type", "order_m3", "supply_m3"]
    carrier_fields = ["carrier_id", "loss_rate_pct"]
    records = {}
    for req in requirement_claims("placeholder"):
        q1 = req.startswith("REQ-Q1")
        question = "question_3" if "Q3" in req else "question_4" if "Q4" in req else "question_2"
        record = result[question]
        nominal = record.get("producer_nominal_check", {})
        metric = {}
        status = "SUPPORTED"
        if q1:
            metric = {
                "top50_supplier_count": len(result["question_1"]["top_50_suppliers"]),
                "minimum_weight_sensitivity_top50_overlap": result["question_1"][
                    "weight_sensitivity"
                ]["minimum_top50_overlap_count"],
            }
        else:
            metric = {
                "used_supplier_count": record.get("weekly_plan_supplier_count", 0),
                "weekly_all_supply_purchase_cost": record.get("material_mix", {}).get(
                    "all_supply_purchase_cost", 0
                ),
                "weekly_effective_received_product_m3": record.get(
                    "weekly_effective_received_product_m3", 0
                ),
                "weekly_transport_loss_m3": record.get("material_mix", {}).get(
                    "transport_loss_m3", 0
                ),
            }
            if req == "REQ-Q2-MINIMUM-SUPPLIERS":
                metric["minimum_supplier_count"] = record.get("minimum_supplier_count", 0)
                if record.get("cardinality_status") != "MILP_EXACT_CARDINALITY" or record.get(
                    "minimum_supplier_count"
                ) != record["optimization"].get("optimistic_cardinality_lower_bound"):
                    status = "INSUFFICIENT"
            if "Q3" in req:
                metric.update(
                    {
                        "material_A_supply_m3": record.get("material_mix", {}).get("A", 0),
                        "material_C_supply_m3": record.get("material_mix", {}).get("C", 0),
                    }
                )
            if "Q4" in req:
                metric["conditional_capacity_product_m3"] = record.get(
                    "weekly_capacity_product_m3", 0
                )
                metric["capacity_increase_product_m3"] = record.get(
                    "weekly_capacity_increase_product_m3", 0
                )
            if "EFFECT" in req or req == "REQ-INVENTORY-PRODUCTION":
                stress = record.get("historical_stress", {})
                metric.update(
                    {
                        "stress_inventory_floor_violation_weeks": stress.get(
                            "inventory_floor_violation_weeks", 48
                        ),
                        "stress_carrier_capacity_violation_weeks": stress.get(
                            "carrier_capacity_violation_weeks", 48
                        ),
                    }
                )
            if not nominal.get("feasible") or not record["optimization"]["lexicographic_complete"]:
                status = "INSUFFICIENT"
        if req.startswith("REQ-OUTPUT-"):
            status = "INSUFFICIENT"  # Derived CSV does not pretend to be missing official template.
            metric = {
                "verified_derived_plan_table_count": len(result.get("plan_table_artifacts", []))
            }
        if req == "REQ-Q4-CAPACITY-INCREASE":
            # Attained capacity is checked; independent maximum proof is absent.
            status = "INSUFFICIENT"
        prefix = (
            "ATTACHMENTS"
            if req.startswith("REQ-OUTPUT-")
            else ("Q1" if q1 else "Q3" if "Q3" in req else "Q4" if "Q4" in req else "Q2")
        )
        metric = {f"{prefix}.{key}": value for key, value in metric.items()}
        records[req] = {
            "generation_method": "CONDITIONAL_SIMULATION"
            if q1
            else ("CONDITIONAL_SIMULATION" if "EFFECT" in req else "OPTIMIZATION"),
            "source_ids": ["SRC-2021-SUPPLIER-WORKBOOK"]
            + ([] if q1 else ["SRC-2021-CARRIER-WORKBOOK"]),
            "scope": {
                "fields": supplier_fields + ([] if q1 else carrier_fields),
                "time": ["W001-W168", "W169-W216"],
                "entities": ["SUPPLIERS_402"] + ([] if q1 else ["CARRIERS_8"]),
            },
            "metric_values": metric,
            "assumption_artifact_sha256": assumptions_hash,
            "conditional_scope": {
                "fields": supplier_fields + ([] if q1 else carrier_fields),
                "time": ["W001-W168", "W169-W216"]
                if q1
                else ["FUTURE_24_WEEKS", "REGISTERED_CAPACITY_SCENARIO"],
                "entities": ["SUPPLIERS_402"] + ([] if q1 else ["CARRIERS_8"]),
            },
            "status": status,
            "scope_limitation": (
                "Development conditional model; excludes empirical future guarantee and"
                " global real-world optimum."
            ),
        }
    return records


def bind_scientific_metrics(result: dict[str, Any]) -> None:
    """Publish globally unique, measured metric IDs for runtime source binding."""
    for record in result["scientific_evidence"].values():
        result["final_metrics"].update(record["metric_values"])


def solve(case_root: Path, candidate_id: str, seed: int) -> dict[str, Any]:
    if candidate_id not in CANDIDATES:
        raise ValueError("CANDIDATE_ID_UNKNOWN")
    assumptions_path = case_root / "models/assumptions_and_symbols.json"
    assumptions_hash = hashlib.sha256(assumptions_path.read_bytes()).hexdigest()
    data = load_case_data(case_root)
    estimates = estimate_suppliers(data, candidate_id, seed)
    losses = estimate_carrier_losses(data, candidate_id)
    _, top50 = importance_scores(data, estimates)
    type_map = dict(zip(data.supplier_ids, data.supplier_types, strict=True))
    ratios = dict(zip(data.supplier_ids, map(float, estimates["delivery_ratio"]), strict=True))
    loss_map = dict(zip(data.carrier_ids, map(float, losses), strict=True))
    result = {
        "candidate_id": candidate_id,
        "seed": seed,
        "status": "SUCCESS",
        "implementation_revision": "RC8_2021_SCIENTIFIC_REPAIR_V6",
        "candidate_algorithm_change": (
            "All three historical candidate IDs now share repaired "
            "exact-cardinality and lexicographic objectives; parameter estimators "
            "retain their historical identity."
        ),
        "answer_access_status": "UNLOCKED_AFTER_FIRST_RUN",
        "test_accessed": False,
        "split_usage": {
            "train": "W001-W168",
            "validation": "W169-W216",
            "test": "W217-W240_UNACCESSED",
        },
        "input_hashes": dict(OFFICIAL_HASHES),
        "assumption_artifact_sha256": assumptions_hash,
        "parameters": {
            "supplier_types": type_map,
            "supplier_delivery_ratios": ratios,
            "supplier_order_capacity_m3": dict(
                zip(data.supplier_ids, map(float, estimates["order_capacity"]), strict=True)
            ),
            "supplier_supply_capacity_m3": dict(
                zip(data.supplier_ids, map(float, estimates["delivered_capacity"]), strict=True)
            ),
            "carrier_loss_rates": loss_map,
            "carrier_capacity_m3": CARRIER_CAPACITY,
            "consumption": CONSUMPTION,
            "purchase_prices": PURCHASE_PRICE,
            "training_weeks": "W001-W168",
            "candidate_estimator": CANDIDATES[candidate_id],
        },
        "question_1": {
            "importance_model": "MINMAX_WEIGHTED_CAPABILITY_RELIABILITY_STABILITY_ACTIVITY",
            "weights": [0.45, 0.25, 0.15, 0.15],
            "top_50_suppliers": top50,
            "weight_sensitivity": weight_sensitivity(data, estimates, top50),
        },
    }
    for question in [2, 3, 4]:
        key = f"question_{question}"
        optimized = optimize_lexicographic(data, estimates, losses, question)
        allocation = optimized.pop("allocation")
        record = {"optimization": optimized}
        if allocation is None:
            record.update({"status": "NO_FEASIBLE_INCUMBENT", "weekly_plan_supplier_count": 0})
            result[key] = record
            continue
        achieved = float(np.sum(allocation * _effective_coefficients(data, losses)))
        demand = achieved if question == 4 else WEEKLY_DEMAND
        plan = sparse_plan(data, estimates, allocation, demand)
        nominal = verify_plan(
            plan,
            supplier_types=type_map,
            delivery_ratios=ratios,
            carrier_loss_rates=loss_map,
            weekly_demand_product_m3=demand,
            carrier_capacity_m3=CARRIER_CAPACITY,
        )
        material = {
            t: float(np.sum(allocation[[i for i, v in enumerate(data.supplier_types) if v == t]]))
            for t in CONSUMPTION
        }
        material["transport_loss_m3"] = float(np.sum(allocation * losses[None, :]))
        material["all_supply_purchase_cost"] = sum(
            material[t] * PURCHASE_PRICE[t] for t in CONSUMPTION
        )
        record.update(
            {
                "status": "CONDITIONAL_PLAN_AVAILABLE",
                "weekly_plan_repeated_for_24_weeks": plan,
                "weekly_plan_supplier_count": optimized["actual_used_supplier_count"],
                "weekly_effective_received_product_m3": achieved,
                "producer_nominal_check": nominal,
                "material_mix": material,
                "historical_stress": historical_stress(data, plan, estimates, losses),
            }
        )
        if question == 2:
            first = optimized["stages"][0]
            record.update(
                {
                    "minimum_supplier_count": optimized["actual_used_supplier_count"],
                    "candidate_supplier_pool_count": optimized["candidate_supplier_pool_count"],
                    "minimum_supplier_ids": [r["supplier_id"] for r in plan["orders"]],
                    "cardinality_status": "MILP_EXACT_CARDINALITY"
                    if first["optimal"]
                    else "INCUMBENT_UPPER_BOUND",
                    "cardinality_claim_strength": "CONDITIONAL_STATIONARY_MODEL_ONLY",
                    "cardinality_incumbent": first["incumbent_objective"],
                    "cardinality_dual_bound": first["dual_bound"],
                    "cardinality_gap": first["relative_gap"],
                }
            )
        if question == 3:
            record["objective_interpretation"] = (
                "MIN_C_THEN_MAX_A_AT_FIXED_PRODUCTION_THEN_MIN_PURCHASE_THEN_MIN_LOSS"
            )
        if question == 4:
            record.update(
                {
                    "weekly_capacity_product_m3": achieved,
                    "weekly_capacity_increase_product_m3": achieved - WEEKLY_DEMAND,
                    "capacity_claim_strength": "CONDITIONAL_STATIONARY_LP_BOUND",
                    "initial_inventory_requirement_product_m3": 2 * achieved,
                }
            )
        result[key] = record
    primary = "validation_penalized_cost_per_effective_m3"
    if "historical_stress" in result["question_2"]:
        validation = result["question_2"]["historical_stress"]
        perturbations = []
        for name, params in [
            ("DELIVERY_RATIO_MINUS_10_PERCENT", {"ratio_scale": 0.9}),
            ("LOSS_PLUS_1_PERCENTAGE_POINT", {"loss_addition": 0.01}),
            ("CARRIER_CAPACITY_MINUS_10_PERCENT", {"carrier_scale": 0.9}),
        ]:
            replay = historical_stress(
                data,
                result["question_2"]["weekly_plan_repeated_for_24_weeks"],
                estimates,
                losses,
                **params,
            )
            perturbations.append(
                {
                    "perturbation_id": name,
                    "metric": primary,
                    "result": replay[primary],
                    "evidence": "DEVELOPMENT_COUNTERFACTUAL_HISTORICAL_REPLAY",
                }
            )
        result.update(
            {
                "validation": validation,
                "validation_metrics": {primary: validation[primary]},
                "metric_name": primary,
                "metric_value": validation[primary],
                "robustness_evidence": {
                    "metric": primary,
                    "metric_direction": "MIN",
                    "perturbations": perturbations,
                    "failure_cases": [
                        "Historical ratio stationarity can fail.",
                        "Observed surplus must all be purchased and can exceed carrier capacity.",
                        (
                            "No historical order requires explicit model-ratio fallback; not an "
                            "observed future supply."
                        ),
                    ],
                },
            }
        )
    else:
        # No finite placeholder score: runner must preserve the failed attempt.
        result["status"] = "FAILED_NO_Q2_INCUMBENT"
    result["claim_scope"] = (
        "Historical four-factor supplier importance with bounded weight sensitivity."
    )
    result["scientific_evidence"] = scientific_evidence(result, assumptions_hash)
    result["scientific_coverage_status"] = "PARTIAL_SCIENTIFIC_COVERAGE"
    result["whole_problem_scientifically_complete"] = False
    result["limitations"] = [
        (
            "All optimization statements are conditional on stationary "
            "training-derived capacity, fulfillment and loss parameters."
        ),
        (
            "Q2 cardinality is globally minimal only inside that conditional model "
            "when a zero-gap optimal solver stage exists."
        ),
        (
            "Q3 C-min then A-max is a declared priority interpretation, not a "
            "unique priority prescribed by the question."
        ),
        (
            "Normalized all-supply purchase prices exclude unspecified absolute "
            "transport/storage tariffs; total economic optimum is unidentified."
        ),
        (
            "Inventory is in product equivalents with initial stock assumed to "
            "equal two weeks of the question-specific target; Q4 needs additional "
            "initial stock."
        ),
        (
            "All positive supplied volume is purchased; stress routing follows "
            "fixed shares and reports overload rather than discarding surplus."
        ),
        (
            "Carrier splitting is a soft preference; split supplier counts are "
            "reported but not globally minimized."
        ),
        (
            "W169-W216 stresses are counterfactual Development replays, never Final"
            " evaluation or stochastic feasibility guarantees."
        ),
        (
            "Derived CSV plans are supplied; official attachment A/B template "
            "submission is incomplete."
        ),
    ]
    result["final_metrics"] = {
        primary: result.get("metric_value"),
        "question_2_supplier_count": result["question_2"].get("weekly_plan_supplier_count", 0),
        "question_4_weekly_capacity_product_m3": result["question_4"].get(
            "weekly_capacity_product_m3", 0
        ),
    }
    result["figure_ready_data"] = [{"figure_id": "FIG-SUPPLIER-TOP50", "rows": top50}]
    bind_scientific_metrics(result)
    return result


def write_plan_tables(
    case_root: Path, output_path: Path, result: dict[str, Any]
) -> list[dict[str, Any]]:
    artifacts = []
    supplier_ids = sorted(result["parameters"]["supplier_types"])
    carrier_ids = sorted(result["parameters"]["carrier_loss_rates"])
    for question in [2, 3, 4]:
        plan = result[f"question_{question}"].get("weekly_plan_repeated_for_24_weeks")
        if plan is None:
            continue
        orders = {r["supplier_id"]: r["volume_m3"] for r in plan["orders"]}
        transport = {(r["supplier_id"], r["carrier_id"]): r["volume_m3"] for r in plan["transport"]}
        for kind in ["orders", "transport"]:
            path = output_path.with_name(f"{output_path.stem}_q{question}_{kind}_24weeks.csv")
            with path.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(
                    (["supplier_id"] if kind == "orders" else ["supplier_id", "carrier_id"])
                    + [f"FUTURE_W{w:02d}" for w in range(1, 25)]
                )
                for supplier in supplier_ids:
                    if kind == "orders":
                        writer.writerow([supplier] + [orders.get(supplier, 0.0)] * 24)
                    else:
                        for carrier in carrier_ids:
                            writer.writerow(
                                [supplier, carrier] + [transport.get((supplier, carrier), 0.0)] * 24
                            )
            artifacts.append(
                {
                    "artifact_id": str(path.relative_to(case_root)),
                    "question": question,
                    "kind": kind,
                    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                    "format": "DERIVED_CSV_NOT_OFFICIAL_TEMPLATE",
                    "weeks": 24,
                    "data_rows": len(supplier_ids)
                    * (len(carrier_ids) if kind == "transport" else 1),
                }
            )
    return artifacts


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--case-root", required=True)
    parser.add_argument("--candidate-id", required=True, choices=sorted(CANDIDATES))
    parser.add_argument("--seed", required=True, type=int)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    case_root = Path(args.case_root).resolve()
    output_path = (case_root / args.output).resolve()
    output_path.relative_to(case_root)
    result = solve(case_root, args.candidate_id, args.seed)
    result["requirement_claims"] = requirement_claims(str(Path(args.output)))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    result["plan_table_artifacts"] = write_plan_tables(case_root, output_path, result)
    result["scientific_evidence"] = scientific_evidence(
        result, result["assumption_artifact_sha256"]
    )
    bind_scientific_metrics(result)
    output_path.write_text(
        json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    return 0 if result["status"] == "SUCCESS" else 2


if __name__ == "__main__":
    raise SystemExit(main())
