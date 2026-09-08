#!/usr/bin/env python3
"""Independent raw-workbook/derived-output checks, with explicit statistical limits.

No imports from the producer or a shared case helper. Classifier fits and clustering
fits are not independently repeated: their recorded scores/assignments get consistency
checks. Raw counts, composition, permutation tests, centers, perturbation transforms,
CV membership, applicability distances and association bootstraps are recomputed.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from scipy.stats import rankdata

WORKBOOK = "raw/case_files/附件.xlsx"
WORKBOOK_HASH = "ffb82a8e209a005f26883e115de3ddea42ab6e0a34986d312a52a3cea6b1063c"
COMPONENTS = [
    "SiO2",
    "Na2O",
    "K2O",
    "CaO",
    "MgO",
    "Al2O3",
    "Fe2O3",
    "CuO",
    "PbO",
    "BaO",
    "P2O5",
    "SrO",
    "SnO2",
    "SO2",
]
TYPES = {"高钾": 0, "铅钡": 1}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def residual(
    value: float, limit: float = 0, relation: str = "EQ", tolerance: float = 1e-7
) -> dict[str, Any]:
    if not math.isfinite(float(value)):
        raise ValueError("NONFINITE_INDEPENDENT_RESIDUAL")
    return {
        "value": float(value),
        "limit": float(limit),
        "relation": relation,
        "tolerance": float(tolerance),
    }


def check_record(
    rows: dict[str, dict[str, Any]], metrics: dict[str, float], **extra: Any
) -> dict[str, Any]:
    passed = True
    for row in rows.values():
        a, b, t = row["value"], row["limit"], row["tolerance"]
        if row["relation"] == "EQ":
            passed &= abs(a - b) <= t
        elif row["relation"] == "LE":
            passed &= a <= b + t
        elif row["relation"] == "GE":
            passed &= a >= b - t
        else:
            raise ValueError("UNKNOWN_RESIDUAL_RELATION")
    return {
        "feasible": bool(passed),
        "constraint_residuals": rows,
        "recalculation_residuals": {
            key: value
            for key, value in rows.items()
            if key not in {"unknown_accuracy_evaluation_available", "unknown_primary_accuracy_gap"}
        },
        "metric_values": metrics,
        **extra,
    }


def close(vector: np.ndarray) -> np.ndarray:
    values = np.asarray(vector, float)
    if (
        values.shape != (14,)
        or not np.all(np.isfinite(values))
        or np.any(values < 0)
        or values.sum() <= 0
    ):
        raise ValueError("RAW_COMPOSITION_INVALID")
    return values / values.sum()


def logratio(vector: np.ndarray) -> np.ndarray:
    values = close(vector)
    positive = values[values > 0]
    substituted = np.where(values > 0, values, float(positive.min()) * 0.5)
    logs = np.log(substituted / substituted.sum())
    return logs - logs.mean()


def features(vector: np.ndarray, candidate: str) -> np.ndarray:
    if candidate == "BASELINE_RAW_CENTROID":
        return close(vector)
    if candidate == "CLR_RIDGE_WARD":
        return logratio(vector)
    if candidate == "HELLINGER_KNN_COMPLETE":
        return np.sqrt(close(vector))
    raise ValueError("UNKNOWN_CANDIDATE")


def read_raw(root: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    path = root / WORKBOOK
    if digest(path) != WORKBOOK_HASH:
        raise ValueError("OFFICIAL_WORKBOOK_HASH_MISMATCH")
    book = load_workbook(path, read_only=True, data_only=True)
    metadata = []
    for row in book["表单1"].iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            metadata.append(
                {
                    "artifact_id": str(row[0]).zfill(2),
                    "pattern": row[1],
                    "glass_type": row[2],
                    "color": "未记录" if row[3] is None else row[3],
                    "surface_weathering": row[4],
                }
            )
    meta = {r["artifact_id"]: r for r in metadata}
    known = []
    for row in book["表单2"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        sample = str(row[0])
        prefix_match = re.match(r"^(\d+)", sample)
        if prefix_match is None:
            raise ValueError("ARTIFACT_PREFIX_MISSING")
        artifact = prefix_match.group(1).zfill(2)
        values = np.asarray([0.0 if v is None else float(v) for v in row[1:]])
        close(values)
        local = (
            "无风化"
            if "未风化点" in sample
            else "严重风化"
            if "严重风化点" in sample
            else meta[artifact]["surface_weathering"]
        )
        known.append(
            {
                **meta[artifact],
                "sample_id": sample,
                "composition": values,
                "valid": 85 <= values.sum() <= 105,
                "local_weathering": local,
                "total": float(values.sum()),
            }
        )
    unknown = []
    for row in book["表单3"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        values = np.asarray([0.0 if v is None else float(v) for v in row[2:]])
        close(values)
        unknown.append(
            {
                "sample_id": str(row[0]),
                "composition": values,
                "surface_weathering": row[1],
                "total": float(values.sum()),
                "valid": 85 <= values.sum() <= 105,
            }
        )
    book.close()
    if len(metadata) != 58 or len(unknown) != 8:
        raise ValueError("OFFICIAL_ENTITY_COUNTS_INVALID")
    return metadata, known, unknown


def split_groups(metadata: list[dict[str, Any]]) -> dict[str, list[str]]:
    by_type = defaultdict(list)
    for row in metadata:
        by_type[row["glass_type"]].append(row["artifact_id"])
    result = {"train": [], "validation": [], "test": []}
    for typ, ids in sorted(by_type.items()):
        ordered = sorted(
            ids,
            key=lambda identifier: hashlib.sha256(
                f"C-TARGET-BATCH-001-POSITION-1-SPLIT-V1:{typ}:{identifier}".encode()
            ).hexdigest(),
        )
        n_validation = max(2, round(len(ids) * 0.2))
        n_test = max(2, round(len(ids) * 0.2))
        n_train = len(ids) - n_validation - n_test
        result["train"] += ordered[:n_train]
        result["validation"] += ordered[n_train : n_train + n_validation]
        result["test"] += ordered[n_train + n_validation :]
    return {key: sorted(values) for key, values in result.items()}


def group_matrix(known: list[dict[str, Any]], candidate: str) -> tuple[list[str], np.ndarray]:
    grouped = defaultdict(list)
    for row in known:
        if row["valid"]:
            grouped[row["artifact_id"]].append(features(row["composition"], candidate))
    ids = sorted(grouped)
    return ids, np.asarray([np.median(grouped[i], axis=0) for i in ids])


def weathering_centers(known: list[dict[str, Any]]) -> dict[tuple[str, str], dict[str, np.ndarray]]:
    grouped = defaultdict(lambda: defaultdict(list))
    for row in known:
        if row["valid"]:
            status = "无风化" if row["local_weathering"] == "无风化" else "风化"
            grouped[(row["glass_type"], status)][row["artifact_id"]].append(
                logratio(row["composition"])
            )
    return {
        key: {artifact: np.mean(values, axis=0) for artifact, values in members.items()}
        for key, members in grouped.items()
    }


def spearman(values: np.ndarray) -> np.ndarray:
    ranked = np.asarray([rankdata(column, method="average") for column in values.T])
    ranked -= ranked.mean(axis=1, keepdims=True)
    length = np.sqrt(np.sum(ranked**2, axis=1))
    denominator = np.outer(length, length)
    return np.divide(
        ranked @ ranked.T, denominator, out=np.full((14, 14), np.nan), where=denominator > 1e-12
    )


def matrix_error(actual: Any, expected: np.ndarray) -> tuple[float, int]:
    if np.asarray(actual, dtype=object).shape != expected.shape:
        raise ValueError("MATRIX_SHAPE_MISMATCH")
    error = 0.0
    null_mismatches = 0
    for index in np.ndindex(expected.shape):
        value = actual[index[0]][index[1]]
        if not np.isfinite(expected[index]):
            null_mismatches += int(value is not None)
        elif value is None:
            null_mismatches += 1
        else:
            error = max(error, abs(float(value) - float(expected[index])))
    return error, null_mismatches


def verify_output(root: Path, run_id: str, model_path: Path | None = None) -> dict[str, Any]:
    path = root / "runs" / run_id / "output.json" if model_path is None else model_path
    path = path.resolve()
    path.relative_to(root.resolve())
    output = json.loads(path.read_text())
    metadata, known, unknown = read_raw(root)
    candidate = output["candidate_id"]
    seed = int(output["seed"])
    legacy = output.get("pipeline_version") != "2022-c-rc8-development-v6"
    ids, matrix = group_matrix(known, candidate)
    split = split_groups(metadata)
    if legacy:
        return legacy_check(path, run_id, output, metadata, known, unknown, ids, split)
    assumptions = digest(root / "models/assumptions_and_symbols.json")
    if assumptions != output.get("assumption_artifact_sha256"):
        raise ValueError("ASSUMPTION_HASH_MISMATCH")
    requirements = {}
    valid = [r for r in known if r["valid"]]
    invalid = sorted(r["sample_id"] for r in known if not r["valid"])
    scope = output["data_scope"]
    requirements["REQ-VALIDITY"] = check_record(
        {
            "known_row_count_error": residual(scope["known_rows"] - len(known)),
            "valid_known_count_error": residual(scope["valid_known_rows"] - len(valid)),
            "invalid_sample_ids_match": residual(
                float(sorted(scope["invalid_known_sample_ids"]) == invalid), 1
            ),
            "unknown_row_count_error": residual(scope["unknown_rows"] - len(unknown)),
        },
        {"DATA.known_rows": len(known), "DATA.valid_known_rows": len(valid)},
    )
    maximum_closure = max(
        abs(float(close(r["composition"]).sum()) - 1) for r in known + unknown if r["valid"]
    )
    requirements["REQ-COMPOSITION"] = check_record(
        {
            "maximum_closed_sum_error": residual(maximum_closure, 0, "LE"),
            "reported_closure_error_difference": residual(
                scope["maximum_closed_sum_error"] - maximum_closure
            ),
        },
        {"DATA.maximum_closed_sum_error": maximum_closure},
    )
    q1 = output["question_1"]
    table_residuals = {
        "association_field_coverage": residual(
            float(
                sorted(r["field"] for r in q1["weathering_associations"])
                == ["color", "glass_type", "pattern"]
            ),
            1,
        )
    }
    total_permutations = 0
    for record in q1["weathering_associations"]:
        field = record["field"]
        rows = [r for r in metadata if r[field] not in [None, "未记录"]]
        levels = sorted({r[field] for r in rows})
        category = np.asarray([levels.index(r[field]) for r in rows])
        weather = np.asarray([int(r["surface_weathering"] == "风化") for r in rows])
        table = np.zeros((len(levels), 2), int)
        np.add.at(table, (category, weather), 1)
        table_residuals[field + "_table_max_error"] = residual(
            float(np.max(np.abs(np.asarray(record["counts"]) - table)))
        )
        if len(levels) > 1 and np.all(table.sum(axis=0) > 0):
            expected = np.outer(table.sum(axis=1), table.sum(axis=0)) / table.sum()
            statistic = float(np.sum((table - expected) ** 2 / expected))
            rng = np.random.default_rng(record["permutation_seed"])
            exceed = 0
            for _ in range(999):
                perm = np.zeros_like(table)
                np.add.at(perm, (category, rng.permutation(weather)), 1)
                exceed += int(float(np.sum((perm - expected) ** 2 / expected)) >= statistic - 1e-12)
            table_residuals[field + "_permutation_p_error"] = residual(
                record["permutation_p_value"] - (exceed + 1) / 1000
            )
            table_residuals[field + "_permutation_count"] = residual(
                record["permutation_replicates"], 999
            )
            total_permutations += 999
    requirements["REQ-1A"] = check_record(
        table_residuals,
        {
            "Q1.association_table_count": len(q1["weathering_associations"]),
            "Q1.total_permutation_replicates": total_permutations,
        },
    )
    strata = weathering_centers(known)
    paired = 0
    effect_residuals = {}
    backcast_expected = {}
    for typ in sorted(TYPES):
        u = strata.get((typ, "无风化"), {})
        w = strata.get((typ, "风化"), {})
        paired += len(set(u) & set(w))
        if not u or not w:
            continue
        first = np.mean(list(u.values()), axis=0)
        second = np.mean(list(w.values()), axis=0)
        effect = q1["type_conditioned_composition_effects"][typ]
        for state, center in [("unweathered", first), ("weathered", second)]:
            observed = np.asarray([effect[state + "_center_percent"][name] for name in COMPONENTS])
            effect_residuals[typ + state + "_center_max_error"] = residual(
                float(np.max(np.abs(observed - close(np.exp(center)) * 100)))
            )
        effect_residuals[typ + "_unweathered_group_count_error"] = residual(
            effect["unweathered_artifact_n"] - len(u)
        )
        effect_residuals[typ + "_weathered_group_count_error"] = residual(
            effect["weathered_artifact_n"] - len(w)
        )
        for artifact, value in w.items():
            backcast_expected[artifact] = (
                close(
                    np.exp(
                        first if candidate == "BASELINE_RAW_CENTROID" else value - (second - first)
                    )
                )
                * 100
            )
    requirements["REQ-1B"] = check_record(
        effect_residuals,
        {
            "Q1.type_effect_count": len(q1["type_conditioned_composition_effects"]),
            "Q1.paired_artifact_count": paired,
        },
    )
    backcast_error = 0.0
    backcast_ids = []
    for row in q1["preweathering_backcasts"]:
        backcast_ids.append(row["artifact_id"])
        values = np.asarray([row["predicted_preweather_percent"][name] for name in COMPONENTS])
        backcast_error = max(
            backcast_error, float(np.max(np.abs(values - backcast_expected[row["artifact_id"]])))
        )
    requirements["REQ-1C"] = check_record(
        {
            "backcast_max_absolute_percent_error": residual(backcast_error),
            "backcast_artifact_set_match": residual(
                float(sorted(backcast_ids) == sorted(backcast_expected)), 1
            ),
        },
        {"Q1.backcast_artifact_count": len(backcast_expected), "Q1.paired_artifact_count": paired},
        scope="CONDITIONAL_TYPE_EXCHANGEABILITY;NO_CAUSAL_IDENTIFICATION",
    )
    cv = output["scientific_diagnostics"]["grouped_repeated_cv"]
    allowed = sorted(set(ids) & set(split["train"] + split["validation"]))
    evaluated = Counter()
    overlap = 0
    fold_bad = 0
    for fold in cv["folds"]:
        fit = set(fold["fit_artifact_ids"])
        evaluate = set(fold["evaluation_artifact_ids"])
        fold_bad += int(bool(fit & evaluate) or fit | evaluate != set(allowed))
        overlap += len((fit | evaluate) & set(split["test"]))
        evaluated.update(fold["evaluation_artifact_ids"])
    prediction_count = sum(evaluated.values())
    cv_checks = {
        "cv_effective_group_set_match": residual(float(cv["effective_artifact_ids"] == allowed), 1),
        "cv_internal_test_group_overlap": residual(overlap),
        "fold_group_leakage_count": residual(fold_bad),
        "each_artifact_evaluated_two_times": residual(
            max([abs(evaluated[i] - 2) for i in allowed] + [0])
        ),
        "cv_reported_prediction_count_error": residual(
            cv["held_out_group_predictions"] - prediction_count
        ),
        "historical_contamination_acknowledged": residual(
            float(cv["historical_internal_test_labels_exposed"]), 1
        ),
    }
    requirements["REQ-2A"] = check_record(
        cv_checks,
        {
            "Q2.cv_artifact_count": len(allowed),
            "Q2.cv_prediction_count": prediction_count,
            "Q2.cv_internal_test_overlap_count": overlap,
        },
        classifier_scores_independently_refit=False,
        scope="FINITE_COLLECTION_DEVELOPMENT_DIAGNOSTIC",
    )
    q2 = output["question_2"]
    assignments = {}
    cluster_count = 0
    cluster_error = 0.0
    membership_bad = 0
    by_id = defaultdict(list)
    type_map = {r["artifact_id"]: r["glass_type"] for r in metadata}
    for row in valid:
        by_id[row["artifact_id"]].append(close(row["composition"]))
    composition = {i: np.mean(values, axis=0) for i, values in by_id.items()}
    bootstrap_bad = 0
    for typ, record in q2["within_type_subtypes"].items():
        if "artifact_assignments" not in record:
            continue
        assigned = record["artifact_assignments"]
        assignments.update(assigned)
        membership_bad += int(set(assigned) != {i for i in ids if type_map[i] == typ})
        cluster_count += len(record["clusters"])
        cluster_members = Counter()
        for cluster in record["clusters"]:
            members = cluster["artifact_ids"]
            cluster_members.update(members)
            membership_bad += int(
                cluster["artifact_n"] != len(members)
                or any(assigned[i] != cluster["subtype_id"] for i in members)
            )
            center = np.mean([composition[i] for i in members], axis=0)
            center = center / center.sum() * 100
            actual = np.asarray([cluster["chemical_center_percent"][name] for name in COMPONENTS])
            cluster_error = max(cluster_error, float(np.max(np.abs(actual - center))))
        membership_bad += int(cluster_members != Counter({artifact: 1 for artifact in assigned}))
        boot = record["bootstrap"]
        seen = np.asarray(boot["pair_observation_counts"])
        same = np.asarray(boot["pair_same_cluster_counts"])
        bootstrap_bad += int(
            seen.shape != same.shape or np.any(same < 0) or np.any(same > seen) or np.any(seen > 20)
        )
        bootstrap_bad += int(len(boot["replicates"]) != 20)
        recomputed_seen = np.zeros_like(seen)
        artifact_order = boot["artifact_order"]
        for replicate in boot["replicates"]:
            if replicate["status"] != "SUCCESS":
                continue
            sampled = replicate["sampled_artifact_ids"]
            bootstrap_bad += int(len(sampled) != len(artifact_order))
            positions = sorted({artifact_order.index(artifact) for artifact in sampled})
            recomputed_seen[np.ix_(positions, positions)] += 1
        bootstrap_bad += int(np.any(recomputed_seen != seen))
        for i in range(len(seen)):
            for j in range(len(seen)):
                actual = boot["co_membership_fraction"][i][j]
                bootstrap_bad += int(
                    (actual is not None)
                    if seen[i, j] == 0
                    else (actual is None or abs(actual - same[i, j] / seen[i, j]) > 1e-9)
                )
    requirements["REQ-2B"] = check_record(
        {
            "subtype_membership_failures": residual(membership_bad),
            "chemical_center_max_error": residual(cluster_error),
        },
        {
            "Q2.subtype_assigned_artifact_count": len(assignments),
            "Q2.subtype_cluster_count": cluster_count,
        },
        clustering_fits_independently_refit=False,
        scope="ASSIGNMENT_AND_CHEMICAL_CENTER_CONSISTENCY_ONLY",
    )
    perturb_error = 0.0
    norm_error = 0.0
    norms = []
    for perturbation in q2["subtype_sensitivity"]["replicates"]:
        rng = np.random.default_rng(perturbation["perturbation_seed"])
        changed = []
        for row in known:
            new = dict(row)
            if row["valid"]:
                changed_vector = (
                    close(np.maximum(close(row["composition"]) + rng.normal(0, 0.01, 14), 1e-8))
                    * 100
                )
                emitted = np.asarray(
                    perturbation["perturbed_compositions_percent"][row["sample_id"]]
                )
                perturb_error = max(perturb_error, float(np.max(np.abs(emitted - changed_vector))))
                new["composition"] = changed_vector
            changed.append(new)
        altered_ids, altered = group_matrix(changed, candidate)
        if altered_ids != ids:
            raise ValueError("PERTURBATION_GROUP_SET_CHANGED")
        change = float(np.linalg.norm(altered - matrix))
        norms.append(change)
        norm_error = max(norm_error, abs(change - perturbation["feature_matrix_change_frobenius"]))
    requirements["REQ-2C"] = check_record(
        {
            "actual_perturbation_composition_error": residual(perturb_error),
            "feature_change_norm_error": residual(norm_error),
            "nonzero_feature_change": residual(min(norms), 1e-12, "GE", 0),
            "bootstrap_count_ratio_consistency_failures": residual(bootstrap_bad),
        },
        {
            "Q2.minimum_perturbation_feature_change": min(norms),
            "Q2.subtype_bootstrap_requested_per_type": 20,
        },
        bootstrap_clustering_fits_independently_refit=False,
    )
    predictions = output["question_3"]["unknown_type_predictions"]
    unknown_map = {r["sample_id"]: r for r in unknown}
    pred_bad = 0
    unknown_true_labels = 0
    distance_error = 0.0
    flip_rates = []
    within = 0
    noise_error = 0.0
    fit_ids = allowed
    fit_positions = [ids.index(i) for i in fit_ids]
    fitting = matrix[fit_positions]
    fit_labels = np.asarray([TYPES[type_map[i]] for i in fit_ids])
    rng = np.random.default_rng(seed)
    for prediction in predictions:
        raw = unknown_map[prediction["sample_id"]]
        score = float(prediction["decision_score_lead_barium"])
        label = int(score >= 0.5)
        pred_bad += int(
            prediction["predicted_type"] != ["高钾", "铅钡"][label] or not 0 <= score <= 1
        )
        unknown_true_labels += int(prediction.get("unknown_true_label_available") is not False)
        feature = features(raw["composition"], candidate)
        local = fitting[fit_labels == label]
        distances = np.linalg.norm(local[:, None, :] - local[None, :, :], axis=2)
        np.fill_diagonal(distances, np.inf)
        threshold = float(np.quantile(distances.min(axis=1), 0.95))
        distance = float(np.linalg.norm(local - feature, axis=1).min())
        distance_error = max(
            distance_error,
            abs(distance - prediction["nearest_predicted_class_training_distance"]),
            abs(threshold - prediction["training_class_distance_threshold_95pct"]),
        )
        expected_domain = distance <= threshold
        within += int(expected_domain)
        pred_bad += int(prediction["within_training_distance_domain"] != expected_domain)
        altered_scores = prediction["perturbation_decision_scores"]
        if len(altered_scores) != 50:
            raise ValueError("UNKNOWN_PERTURBATION_COUNT_INVALID")
        pred_bad += sum(
            not math.isfinite(float(value)) or not 0 <= value <= 1 for value in altered_scores
        )
        flip = sum(int(value >= 0.5) != label for value in altered_scores)
        flip_rates.append(flip / 50)
        pred_bad += int(
            flip != prediction["perturbation_flip_count"]
            or abs(flip / 50 - prediction["perturbation_flip_rate"]) > 1e-9
        )
        for i in range(50):
            altered = close(np.maximum(close(raw["composition"]) + rng.normal(0, 0.01, 14), 1e-8))
            change = float(np.linalg.norm(features(altered, candidate) - feature))
            noise_error = max(
                noise_error, abs(change - prediction["perturbation_feature_change_norms"][i])
            )
    requirements["REQ-3A"] = check_record(
        {
            "unknown_prediction_id_set_match": residual(
                float(sorted(r["sample_id"] for r in predictions) == sorted(unknown_map)), 1
            ),
            "class_threshold_consistency_failures": residual(pred_bad),
            "unavailable_unknown_labels_not_fabricated": residual(unknown_true_labels),
            "unknown_accuracy_evaluation_available": residual(0, 1, "GE", 0),
        },
        {"Q3.unknown_prediction_count": len(predictions), "Q3.unknown_available_label_count": 0},
        scope="PREDICTIONS_AVAILABLE;UNKNOWN_ACCURACY_INSUFFICIENT",
        classifier_fit_independently_refit=False,
    )
    requirements["REQ-3B"] = check_record(
        {
            "training_distance_max_error": residual(distance_error),
            "perturbation_feature_change_norm_error": residual(noise_error),
            "class_flip_count_consistency_failures": residual(pred_bad),
        },
        {
            "Q3.mean_unknown_perturbation_flip_rate": float(np.mean(flip_rates)),
            "Q3.unknown_in_training_domain_count": within,
        },
        score_function_independently_refit=False,
        scope="RAW_DISTANCE_AND_PERTURBATION_FEATURES_RECOMPUTED;FLIPS_FROM_REPORTED_MODEL_SCORES",
    )
    grouped = defaultdict(lambda: defaultdict(list))
    for row in valid:
        grouped[row["glass_type"]][row["artifact_id"]].append(logratio(row["composition"]))
    vectors = {
        typ: np.asarray([np.mean(points, axis=0) for _, points in sorted(artifacts.items())])
        for typ, artifacts in grouped.items()
    }
    correlation = {typ: spearman(values) for typ, values in vectors.items()}
    nulls = 0
    corr_error = 0.0
    bad_nulls = 0
    for typ in sorted(TYPES):
        error, bad = matrix_error(
            output["question_4"][typ]["spearman_clr_matrix"], correlation[typ]
        )
        corr_error = max(corr_error, error)
        bad_nulls += bad
        nulls += int(np.sum(~np.isfinite(correlation[typ])))
    requirements["REQ-4A"] = check_record(
        {
            "artifact_equal_weight_correlation_max_error": residual(corr_error),
            "undefined_correlation_null_mismatch_count": residual(bad_nulls),
        },
        {
            "Q4.association_artifact_count": sum(len(v) for v in vectors.values()),
            "Q4.undefined_correlation_entry_count": nulls,
        },
    )
    first, second = sorted(TYPES)
    rng = np.random.default_rng(seed)
    boots = []
    for _ in range(30):
        a = vectors[first][rng.choice(len(vectors[first]), len(vectors[first]), replace=True)]
        b = vectors[second][rng.choice(len(vectors[second]), len(vectors[second]), replace=True)]
        boots.append(spearman(a) - spearman(b))
    boots = np.asarray(boots)
    between = output["question_4"]["between_type_difference"]
    difference = correlation[first] - correlation[second]
    error, bad = matrix_error(between["difference_matrix"], difference)
    range_error = 0.0
    range_null_bad = 0
    reported_pairs = []
    finite_count_errors = 0
    for pair in between["pair_difference_ranges"]:
        i = COMPONENTS.index(pair["left"])
        j = COMPONENTS.index(pair["right"])
        values = boots[:, i, j]
        values = values[np.isfinite(values)]
        reported_pairs.append((i, j))
        finite_count_errors += int(pair["bootstrap_finite_replicates"] != len(values))
        actual = pair["bootstrap_percentile_range"]
        if not len(values):
            range_null_bad += int(actual is not None)
        elif actual is None:
            range_null_bad += 1
        else:
            range_error = max(
                range_error,
                float(np.max(np.abs(np.asarray(actual) - np.quantile(values, [0.025, 0.975])))),
            )
    requirements["REQ-4B"] = check_record(
        {
            "association_difference_matrix_error": residual(error),
            "association_difference_null_mismatch": residual(bad),
            "bootstrap_range_max_error": residual(range_error),
            "bootstrap_range_null_mismatch": residual(range_null_bad),
            "bootstrap_replicate_count": residual(between["bootstrap_replicates"], 30),
            "all_component_pairs_present_once": residual(
                float(
                    sorted(reported_pairs) == [(i, j) for i in range(14) for j in range(i + 1, 14)]
                ),
                1,
            ),
            "bootstrap_finite_replicate_count_errors": residual(finite_count_errors),
        },
        {"Q4.association_bootstrap_replicates": 30, "Q4.component_pair_count": 91},
        scope="EXPLORATORY_BOOTSTRAP;NO_MULTIPLICITY_ADJUSTMENT",
    )
    requirements["REQ-EVIDENCE"] = check_record(
        {"unknown_primary_accuracy_gap": residual(0, 1, "GE", 0)},
        {"DATA.scientific_requirement_count": 13},
        scope="PARTIAL_SCIENTIFIC_COVERAGE",
    )
    metrics = {
        key: value
        for record in requirements.values()
        for key, value in record["metric_values"].items()
    }
    metric_bad = 0
    for req, fact in output["scientific_evidence"].items():
        for name, value in fact["metric_values"].items():
            expected = requirements[req]["metric_values"].get(name)
            metric_bad += int(
                expected is None
                or abs(float(value) - expected) > 1e-7
                or abs(float(output["final_metrics"][name]) - expected) > 1e-7
            )
    return {
        "schema_version": "c2022-independent-check/v1",
        "run_id": run_id,
        "output_sha256": digest(path),
        "checker_sha256": digest(Path(__file__)),
        "input_hashes": {WORKBOOK: WORKBOOK_HASH},
        "assumption_artifact_sha256": assumptions,
        "requirements": requirements,
        "metric_values": metrics,
        "scientific_metric_binding_mismatch_count": metric_bad,
        "scientific_coverage_status": "PARTIAL_SCIENTIFIC_COVERAGE",
        "whole_problem_scientifically_complete": False,
        "independence": "SEPARATE_IMPLEMENTATION_NO_PRODUCER_OR_HELPER_IMPORT",
        "limitations": [
            (
                "Classifier and clustering fits are not independently rerun; decision "
                "scores and subtype assignments get consistency checks."
            ),
            (
                "Count, composition, permutation, centers, feature perturbations, CV "
                "membership, raw distances and association bootstrap are independently "
                "recomputed."
            ),
            (
                "Independent Python checking is not an independent agent audit; unknown"
                " accuracy and causal backcasts are not established."
            ),
        ],
    }


def legacy_check(
    path: Path,
    run_id: str,
    output: dict[str, Any],
    metadata: list[dict[str, Any]],
    known: list[dict[str, Any]],
    unknown: list[dict[str, Any]],
    ids: list[str],
    split: dict[str, list[str]],
) -> dict[str, Any]:
    overlap = sorted(set(ids) & set(split["test"]))
    cv = output["scientific_diagnostics"]["grouped_repeated_cv"]
    expected = len(ids) * 5
    metric = {
        "Q2.legacy_cv_artifact_count": len(ids),
        "Q2.legacy_cv_prediction_count": expected,
        "Q2.legacy_internal_test_artifact_count": len(overlap),
        "DATA.known_rows": len(known),
        "Q3.unknown_prediction_count": len(unknown),
    }
    return {
        "schema_version": "c2022-independent-check/v1",
        "run_id": run_id,
        "output_sha256": digest(path),
        "checker_sha256": digest(Path(__file__)),
        "input_hashes": {WORKBOOK: WORKBOOK_HASH},
        "legacy_control": True,
        "requirements": {
            "REQ-2A": check_record(
                {
                    "legacy_cv_prediction_count_error": residual(
                        cv["held_out_group_predictions"] - expected
                    ),
                    "sealed_internal_test_group_violation_count": residual(len(overlap)),
                },
                metric,
            )
        },
        "metric_values": metric,
        "historical_internal_test_groups_used": overlap,
        "legacy_reported_test_labels_accessed": cv.get("test_labels_accessed"),
        "scientific_coverage_status": "PARTIAL_SCIENTIFIC_COVERAGE",
        "whole_problem_scientifically_complete": False,
        "limitation": (
            "Old algorithm control; 56x5 repeats are not 280 independent artifacts "
            "or sealed external validation."
        ),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--case-root", required=True)
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--model-output")
    args = parser.parse_args()
    root = Path(args.case_root).resolve()
    path = Path(args.output)
    if not path.is_absolute():
        path = root / path
    path = path.resolve()
    path.relative_to(root)
    model_path = None if args.model_output is None else (root / args.model_output).resolve()
    result = verify_output(root, args.run_id, model_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "run_id": args.run_id,
                "output_sha256": result["output_sha256"],
                "scientific_coverage_status": result["scientific_coverage_status"],
            }
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
