#!/usr/bin/env python3
"""First-party 2022 C Development scientific repair.

The deterministic case runner supplies ``--case-root``, ``--candidate-id``,
``--seed`` and ``--output``.  This program writes exactly one JSON output and
does not modify raw inputs.  Candidate selection uses only the predeclared
validation score; the unknown samples have no answer labels in this workspace.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from scipy.stats import chi2_contingency, rankdata
from sklearn.cluster import AgglomerativeClustering
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import adjusted_rand_score, silhouette_score
from sklearn.model_selection import RepeatedStratifiedKFold
from sklearn.neighbors import KNeighborsClassifier
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler

PIPELINE_VERSION = "2022-c-rc8-development-v6"
CANDIDATES = (
    "BASELINE_RAW_CENTROID",
    "CLR_RIDGE_WARD",
    "HELLINGER_KNN_COMPLETE",
)
COMPONENTS = (
    "SiO2",
    "Na2O",
    "K2O",
    "CaO",
    "MgO",
    "Al2O3",
    "Fe2O3",
    "CuO",
    "PbO",
    "BaO",
    "P2O5",
    "SrO",
    "SnO2",
    "SO2",
)
TYPE_TO_INT = {"高钾": 0, "铅钡": 1}
INT_TO_TYPE = {value: key for key, value in TYPE_TO_INT.items()}
SPLIT_SALT = "C-TARGET-BATCH-001-POSITION-1-SPLIT-V1"
OFFICIAL_WORKBOOK_SHA256 = "ffb82a8e209a005f26883e115de3ddea42ab6e0a34986d312a52a3cea6b1063c"
PERMUTATIONS = 999
CV_REPEATS = 2
SUBTYPE_BOOTSTRAPS = 20
ASSOCIATION_BOOTSTRAPS = 30
UNKNOWN_PERTURBATIONS = 50


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--case-root", required=True)
    parser.add_argument("--candidate-id", choices=CANDIDATES, required=True)
    parser.add_argument("--seed", type=int, required=True)
    parser.add_argument("--output", required=True)
    return parser.parse_args()


def safe_path(root: Path, relative: str) -> Path:
    value = Path(relative)
    if value.is_absolute() or ".." in value.parts:
        raise ValueError("CASE_RELATIVE_PATH_REQUIRED")
    resolved = (root / value).resolve()
    resolved.relative_to(root.resolve())
    return resolved


def base_artifact_id(sample_id: str) -> str:
    match = re.match(r"^(\d+)", sample_id)
    if match is None:
        raise ValueError("SAMPLE_ID_WITHOUT_ARTIFACT_PREFIX")
    return match.group(1).zfill(2)


def local_weathering(sample_id: str, surface_status: str) -> str:
    if "未风化点" in sample_id:
        return "无风化"
    if "严重风化点" in sample_id:
        return "严重风化"
    return surface_status


def finite_float(value: Any) -> float:
    result = float(value)
    if not math.isfinite(result):
        raise ValueError("NONFINITE_RESULT")
    return result


def json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(item) for item in value]
    if isinstance(value, np.ndarray):
        return json_safe(value.tolist())
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return finite_float(value)
    if isinstance(value, (str, int, bool)) or value is None:
        return value
    raise TypeError(f"unsupported JSON value: {type(value).__name__}")


def load_official_workbook(
    case_root: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    workbook_path = safe_path(case_root, "raw/case_files/附件.xlsx")
    if hashlib.sha256(workbook_path.read_bytes()).hexdigest() != OFFICIAL_WORKBOOK_SHA256:
        raise ValueError("OFFICIAL_WORKBOOK_HASH_MISMATCH")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)

    metadata: list[dict[str, Any]] = []
    for row in workbook["表单1"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        artifact_id = str(row[0]).zfill(2)
        metadata.append(
            {
                "artifact_id": artifact_id,
                "pattern": row[1],
                "glass_type": row[2],
                "color": row[3] if row[3] is not None else "未记录",
                "surface_weathering": row[4],
            }
        )
    by_id = {record["artifact_id"]: record for record in metadata}

    known: list[dict[str, Any]] = []
    for row in workbook["表单2"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        sample_id = str(row[0])
        artifact_id = base_artifact_id(sample_id)
        values = np.array(
            [0.0 if value is None else float(value) for value in row[1:]], dtype=float
        )
        observed = np.array([value is not None for value in row[1:]], dtype=bool)
        total = float(values.sum())
        meta = by_id[artifact_id]
        known.append(
            {
                **meta,
                "sample_id": sample_id,
                "composition": values,
                "observed": observed,
                "total": total,
                "valid": 85.0 <= total <= 105.0,
                "local_weathering": local_weathering(sample_id, meta["surface_weathering"]),
            }
        )

    unknown: list[dict[str, Any]] = []
    for row in workbook["表单3"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        values = np.array(
            [0.0 if value is None else float(value) for value in row[2:]], dtype=float
        )
        observed = np.array([value is not None for value in row[2:]], dtype=bool)
        total = float(values.sum())
        unknown.append(
            {
                "sample_id": str(row[0]),
                "surface_weathering": row[1],
                "composition": values,
                "observed": observed,
                "total": total,
                "valid": 85.0 <= total <= 105.0,
            }
        )
    workbook.close()
    if len(metadata) != 58 or len(unknown) != 8:
        raise ValueError("OFFICIAL_WORKBOOK_ENTITY_COUNTS_INVALID")
    if any(len(row["composition"]) != 14 for row in known + unknown):
        raise ValueError("OFFICIAL_WORKBOOK_COMPONENT_COUNT_INVALID")
    return metadata, known, unknown


def close_composition(values: np.ndarray) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    if np.any(values < 0) or not np.isfinite(values).all() or values.sum() <= 0:
        raise ValueError("INVALID_COMPOSITION")
    return values / values.sum()


def replace_zeros(values: np.ndarray, fraction: float = 0.50) -> np.ndarray:
    closed = close_composition(values)
    positive = closed[closed > 0]
    if positive.size == 0:
        raise ValueError("EMPTY_COMPOSITION")
    replacement = float(positive.min()) * fraction
    replaced = np.where(closed > 0, closed, replacement)
    return close_composition(replaced)


def clr(values: np.ndarray, fraction: float = 0.50) -> np.ndarray:
    logged = np.log(replace_zeros(values, fraction))
    return logged - logged.mean()


def transform(values: np.ndarray, candidate_id: str, fraction: float = 0.50) -> np.ndarray:
    if candidate_id == "BASELINE_RAW_CENTROID":
        return close_composition(values)
    if candidate_id == "CLR_RIDGE_WARD":
        return clr(values, fraction)
    if candidate_id == "HELLINGER_KNN_COMPLETE":
        return np.sqrt(close_composition(values))
    raise ValueError("UNKNOWN_CANDIDATE")


def deterministic_split(metadata: list[dict[str, Any]]) -> dict[str, list[str]]:
    by_type: dict[str, list[str]] = defaultdict(list)
    for record in metadata:
        by_type[record["glass_type"]].append(record["artifact_id"])
    assignment = {"train": [], "validation": [], "test": []}
    for glass_type, identifiers in sorted(by_type.items()):
        ordered = sorted(
            identifiers,
            key=lambda item: hashlib.sha256(
                f"{SPLIT_SALT}:{glass_type}:{item}".encode()
            ).hexdigest(),
        )
        count = len(ordered)
        validation_count = max(2, round(count * 0.20))
        test_count = max(2, round(count * 0.20))
        train_count = count - validation_count - test_count
        if train_count < 2:
            raise ValueError("INSUFFICIENT_STRATIFIED_GROUPS")
        assignment["train"].extend(ordered[:train_count])
        assignment["validation"].extend(ordered[train_count : train_count + validation_count])
        assignment["test"].extend(ordered[train_count + validation_count :])
    return {key: sorted(values) for key, values in assignment.items()}


def artifact_level_dataset(
    metadata: list[dict[str, Any]],
    known: list[dict[str, Any]],
    candidate_id: str,
    fraction: float = 0.50,
) -> tuple[list[str], np.ndarray, np.ndarray]:
    grouped: dict[str, list[np.ndarray]] = defaultdict(list)
    for sample in known:
        if sample["valid"]:
            grouped[sample["artifact_id"]].append(
                transform(sample["composition"], candidate_id, fraction)
            )
    type_by_id = {record["artifact_id"]: TYPE_TO_INT[record["glass_type"]] for record in metadata}
    ids = sorted(grouped)
    matrix = np.vstack([np.median(np.vstack(grouped[item]), axis=0) for item in ids])
    labels = np.array([type_by_id[item] for item in ids], dtype=int)
    return ids, matrix, labels


def fit_type_classifier(candidate_id: str, x: np.ndarray, y: np.ndarray, seed: int) -> Any:
    if candidate_id == "BASELINE_RAW_CENTROID":
        centers = {label: np.mean(x[y == label], axis=0) for label in sorted(set(y.tolist()))}

        class CentroidModel:
            def predict_proba(self, values: np.ndarray) -> np.ndarray:
                distances = np.column_stack(
                    [np.linalg.norm(values - centers[label], axis=1) for label in (0, 1)]
                )
                scale = max(float(np.median(distances)), 1e-9)
                weights = np.exp(-distances / scale)
                return weights / weights.sum(axis=1, keepdims=True)

        return CentroidModel()
    if candidate_id == "CLR_RIDGE_WARD":
        model = make_pipeline(
            StandardScaler(),
            LogisticRegression(C=0.5, solver="liblinear", random_state=seed, max_iter=2000),
        )
        return model.fit(x, y)
    neighbors = max(1, min(5, int(np.sqrt(len(y)))))
    model = KNeighborsClassifier(n_neighbors=neighbors, weights="distance", metric="euclidean")
    return model.fit(x, y)


def brier_loss(y_true: np.ndarray, probability_one: np.ndarray) -> float:
    return finite_float(np.mean((probability_one - y_true.astype(float)) ** 2))


def contingency_result(metadata: list[dict[str, Any]], field: str, seed: int) -> dict[str, Any]:
    rows = [r for r in metadata if r[field] not in [None, "未记录"]]
    levels = sorted({r[field] for r in rows})
    category = np.asarray([levels.index(r[field]) for r in rows])
    weather = np.asarray([int(r["surface_weathering"] == "风化") for r in rows])
    table = np.zeros((len(levels), 2), dtype=int)
    np.add.at(table, (category, weather), 1)
    valid = min(table.shape) >= 2 and np.all(table.sum(axis=0) > 0)
    p_value = permutation_p = chi2 = effect = None
    expected = np.zeros_like(table, dtype=float)
    exceed = 0
    if valid:
        chi2, p_value, _, expected = chi2_contingency(table, correction=False)
        effect = math.sqrt(float(chi2) / (len(rows) * min(table.shape[0] - 1, 1)))
        rng = np.random.default_rng(seed)
        for _ in range(PERMUTATIONS):
            counts = np.zeros_like(table)
            np.add.at(counts, (category, rng.permutation(weather)), 1)
            statistic = float(np.sum((counts - expected) ** 2 / expected))
            exceed += int(statistic >= float(chi2) - 1e-12)
        permutation_p = (exceed + 1) / (PERMUTATIONS + 1)
    return {
        "field": field,
        "levels": levels,
        "weathering_levels": ["无风化", "风化"],
        "counts": table.tolist(),
        "n": int(table.sum()),
        "missing_excluded": len(metadata) - len(rows),
        "chi_square": None if chi2 is None else float(chi2),
        "asymptotic_p_value_diagnostic_only": None if p_value is None else float(p_value),
        "cells_expected_below_5": int(np.sum(expected < 5)),
        "cramers_v": effect,
        "permutation_p_value": permutation_p,
        "permutation_replicates": PERMUTATIONS if valid else 0,
        "permutation_exceedance_count": exceed,
        "permutation_seed": seed,
        "status": "DESCRIPTIVE_ASSOCIATION" if valid else "INSUFFICIENT_VARIATION",
        "interpretation_guard": (
            "ASSOCIATION_NOT_CAUSATION;999_SHUFFLES_PLUS_ONE_CORRECTION;MULTIPLE_TABLES_EXPLORATORY"
        ),
    }


def artifact_compositions(known: list[dict[str, Any]]) -> dict[str, np.ndarray]:
    grouped = defaultdict(list)
    for row in known:
        if row["valid"]:
            grouped[row["artifact_id"]].append(close_composition(row["composition"]))
    return {
        artifact: close_composition(np.mean(values, axis=0)) for artifact, values in grouped.items()
    }


def weathering_group_vectors(
    known: list[dict[str, Any]],
) -> dict[tuple[str, str], dict[str, np.ndarray]]:
    points = defaultdict(lambda: defaultdict(list))
    for row in known:
        if row["valid"]:
            weather = "无风化" if row["local_weathering"] == "无风化" else "风化"
            points[(row["glass_type"], weather)][row["artifact_id"]].append(clr(row["composition"]))
    return {
        key: {artifact: np.mean(values, axis=0) for artifact, values in group.items()}
        for key, group in points.items()
    }


def percent_map(values: np.ndarray) -> dict[str, float]:
    return {
        name: float(value * 100)
        for name, value in zip(COMPONENTS, close_composition(values), strict=True)
    }


def weathering_analysis(
    known: list[dict[str, Any]], candidate_id: str
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    groups = weathering_group_vectors(known)
    effects = {}
    backcasts = []
    for glass_type in sorted(TYPE_TO_INT):
        u = groups.get((glass_type, "无风化"), {})
        w = groups.get((glass_type, "风化"), {})
        paired = sorted(set(u) & set(w))
        if not u or not w:
            effects[glass_type] = {"status": "INSUFFICIENT_GROUPS", "paired_artifact_ids": paired}
            continue
        u_log = np.mean(list(u.values()), axis=0)
        w_log = np.mean(list(w.values()), axis=0)
        delta = w_log - u_log
        effects[glass_type] = {
            "status": "DESCRIPTIVE_ARTIFACT_EQUAL_WEIGHT",
            "unweathered_artifact_ids": sorted(u),
            "weathered_artifact_ids": sorted(w),
            "unweathered_artifact_n": len(u),
            "weathered_artifact_n": len(w),
            "paired_artifact_ids": paired,
            "paired_artifact_n": len(paired),
            "unweathered_center_percent": percent_map(np.exp(u_log)),
            "weathered_center_percent": percent_map(np.exp(w_log)),
            "weathered_minus_unweathered_clr": dict(
                zip(COMPONENTS, map(float, delta), strict=True)
            ),
            "weighting": "EACH_ARTIFACT_EQUAL_WITHIN_WEATHERING_STRATUM;POINTS_AVERAGED_FIRST",
        }
        for artifact, observed_log in sorted(w.items()):
            restored = np.exp(
                u_log if candidate_id == "BASELINE_RAW_CENTROID" else observed_log - delta
            )
            backcasts.append(
                {
                    "artifact_id": artifact,
                    "glass_type": glass_type,
                    "weathered_sample_ids": [
                        r["sample_id"]
                        for r in known
                        if r["artifact_id"] == artifact
                        and r["valid"]
                        and r["local_weathering"] != "无风化"
                    ],
                    "method_scope": "CONDITIONAL_TYPE_EXCHANGEABILITY_BACKCAST_NOT_CAUSAL",
                    "predicted_preweather_percent": percent_map(restored),
                }
            )
    return effects, backcasts


def cluster_subtypes(
    metadata: list[dict[str, Any]],
    known: list[dict[str, Any]],
    candidate_id: str,
    seed: int,
    bootstrap: bool = True,
) -> dict[str, Any]:
    ids, matrix, labels = artifact_level_dataset(metadata, known, candidate_id)
    composition = artifact_compositions(known)
    output = {}
    rng = np.random.default_rng(seed)
    for label, glass_type in sorted(INT_TO_TYPE.items()):
        positions = np.flatnonzero(labels == label)
        local_ids = [ids[i] for i in positions]
        if len(positions) < 4:
            output[glass_type] = {"status": "INSUFFICIENT_ARTIFACTS", "artifact_n": len(positions)}
            continue
        all_values = matrix[positions]
        selected = sorted(np.argsort(np.var(all_values, axis=0), kind="stable")[-6:].tolist())
        values = all_values[:, selected]
        linkage = "complete" if candidate_id == "HELLINGER_KNN_COMPLETE" else "ward"
        diagnostics = []
        fits = {}
        for k in range(2, min(4, len(values) - 1) + 1):
            fitted = AgglomerativeClustering(n_clusters=k, linkage=linkage).fit_predict(values)
            score = float(silhouette_score(values, fitted)) if len(np.unique(fitted)) > 1 else -1.0
            diagnostics.append({"k": k, "silhouette_internal_geometry_only": score})
            fits[k] = fitted
        k = max(diagnostics, key=lambda d: (d["silhouette_internal_geometry_only"], -d["k"]))["k"]
        chosen = fits[k]
        assignments = {
            artifact: f"{glass_type}-S{int(group) + 1}"
            for artifact, group in zip(local_ids, chosen, strict=True)
        }
        clusters = []
        for number in sorted(np.unique(chosen)):
            members = [local_ids[i] for i in np.flatnonzero(chosen == number)]
            clusters.append(
                {
                    "subtype_id": f"{glass_type}-S{int(number) + 1}",
                    "artifact_ids": members,
                    "artifact_n": len(members),
                    "chemical_center_percent": percent_map(
                        np.mean([composition[i] for i in members], axis=0)
                    ),
                }
            )
        seen = np.zeros((len(values), len(values)), int)
        same = np.zeros_like(seen)
        replicates = []
        if bootstrap:
            for replicate in range(SUBTYPE_BOOTSTRAPS):
                sample = rng.choice(len(values), size=len(values), replace=True)
                available = sorted(set(sample.tolist()))
                if len(available) < k:
                    replicates.append(
                        {"replicate": replicate, "status": "INSUFFICIENT_UNIQUE_GROUPS"}
                    )
                    continue
                boot = AgglomerativeClustering(n_clusters=k, linkage=linkage).fit_predict(
                    values[sample]
                )
                collapsed = {i: int(np.argmax(np.bincount(boot[sample == i]))) for i in available}
                indices = np.ix_(available, available)
                labels_in = np.asarray([collapsed[i] for i in available])
                seen[indices] += 1
                same[indices] += (labels_in[:, None] == labels_in[None, :]).astype(int)
                replicates.append(
                    {
                        "replicate": replicate,
                        "status": "SUCCESS",
                        "sampled_artifact_ids": [local_ids[i] for i in sample],
                        "unique_artifact_n": len(available),
                        "adjusted_rand_index": float(
                            adjusted_rand_score(chosen[available], labels_in)
                        ),
                    }
                )
        consensus = [
            [
                None if seen[i, j] == 0 else float(same[i, j] / seen[i, j])
                for j in range(len(values))
            ]
            for i in range(len(values))
        ]
        output[glass_type] = {
            "status": "EXPLORATORY_SUBTYPES",
            "artifact_n": len(local_ids),
            "candidate_k_diagnostics": diagnostics,
            "selected_k": k,
            "linkage": linkage,
            "selection_rule": "MAX_INTERNAL_SILHOUETTE_THEN_SMALLEST_K_NOT_SCIENTIFIC_VALIDITY",
            "artifact_assignments": assignments,
            "selected_components": [COMPONENTS[i] for i in selected],
            "feature_space": candidate_id,
            "clusters": clusters,
            "bootstrap": {
                "design": "20_ARTIFACT_RESAMPLES_WITH_REPLACEMENT_FIXED_K_AND_FEATURES",
                "requested_replicates": SUBTYPE_BOOTSTRAPS if bootstrap else 0,
                "replicates": replicates,
                "artifact_order": local_ids,
                "pair_observation_counts": seen.tolist(),
                "pair_same_cluster_counts": same.tolist(),
                "co_membership_fraction": consensus,
            },
        }
    return output


def perturb_known(
    known: list[dict[str, Any]], seed: int
) -> tuple[list[dict[str, Any]], dict[str, list[float]]]:
    rng = np.random.default_rng(seed)
    changed = []
    payload = {}
    for row in known:
        new = dict(row)
        if row["valid"]:
            original = close_composition(row["composition"])
            new["composition"] = (
                close_composition(np.maximum(original + rng.normal(0, 0.01, len(COMPONENTS)), 1e-8))
                * 100
            )
            payload[row["sample_id"]] = new["composition"].tolist()
        changed.append(new)
    return changed, payload


def validation_score(
    metadata: list[dict[str, Any]], known: list[dict[str, Any]], candidate_id: str, seed: int
) -> tuple[float, dict[str, Any]]:
    ids, matrix, labels = artifact_level_dataset(metadata, known, candidate_id)
    split = deterministic_split(metadata)
    train = [i for i, artifact in enumerate(ids) if artifact in split["train"]]
    validation = [i for i, artifact in enumerate(ids) if artifact in split["validation"]]
    if not train or not validation:
        raise ValueError("EMPTY_EFFECTIVE_SPLIT")
    model = fit_type_classifier(candidate_id, matrix[train], labels[train], seed)
    scores = model.predict_proba(matrix[validation])[:, 1]
    loss = brier_loss(labels[validation], scores)
    return loss, {
        "group_unit": "ARTIFACT_ID",
        "split_assignment": split,
        "effective_train_ids": [ids[i] for i in train],
        "effective_validation_ids": [ids[i] for i in validation],
        "validation_brier_loss": loss,
        "score_interpretation": "UNCALIBRATED_DECISION_SCORE_SQUARED_ERROR_DIAGNOSTIC",
        "validation_accuracy_diagnostic": float(np.mean((scores >= 0.5) == labels[validation])),
        "validation_predictions": [
            {
                "artifact_id": ids[i],
                "known_class": int(labels[i]),
                "decision_score_lead_barium": float(score),
            }
            for i, score in zip(validation, scores, strict=True)
        ],
        "current_classifier_fit_uses_internal_test_groups": False,
        "historical_internal_test_labels_exposed": True,
        "internal_test_split_is_sealed": False,
        "case_evidence_status": "DEVELOPMENT_ALREADY_CONTAMINATED_NOT_EXTERNAL_VALIDATION",
    }


def grouped_repeated_cv_diagnostic(
    metadata: list[dict[str, Any]], known: list[dict[str, Any]], candidate_id: str, seed: int
) -> dict[str, Any]:
    ids, matrix, labels = artifact_level_dataset(metadata, known, candidate_id)
    split = deterministic_split(metadata)
    allowed = set(split["train"] + split["validation"])
    positions = [i for i, artifact in enumerate(ids) if artifact in allowed]
    ids = [ids[i] for i in positions]
    matrix = matrix[positions]
    labels = labels[positions]
    n_splits = min(4, int(np.min(np.bincount(labels, minlength=2))))
    if n_splits < 2:
        raise ValueError("INSUFFICIENT_GROUPS_FOR_CV")
    cv = RepeatedStratifiedKFold(n_splits=n_splits, n_repeats=CV_REPEATS, random_state=seed)
    folds = []
    losses = []
    accuracies = []
    prediction_count = 0
    for fold, (fit, validate) in enumerate(cv.split(matrix, labels)):
        model = fit_type_classifier(candidate_id, matrix[fit], labels[fit], seed)
        scores = model.predict_proba(matrix[validate])[:, 1]
        loss = brier_loss(labels[validate], scores)
        accuracy = float(np.mean((scores >= 0.5) == labels[validate]))
        losses.append(loss)
        accuracies.append(accuracy)
        prediction_count += len(validate)
        folds.append(
            {
                "fold": fold,
                "fit_artifact_ids": [ids[i] for i in fit],
                "evaluation_artifact_ids": [ids[i] for i in validate],
                "score_squared_error": loss,
                "accuracy": accuracy,
            }
        )
    return {
        "group_unit": "ARTIFACT_ID",
        "effective_group_count": len(ids),
        "effective_artifact_ids": ids,
        "n_splits": n_splits,
        "n_repeats": CV_REPEATS,
        "held_out_group_predictions": prediction_count,
        "folds": folds,
        "brier_loss": {
            "mean": float(np.mean(losses)),
            "minimum": min(losses),
            "maximum": max(losses),
        },
        "accuracy": {
            "mean": float(np.mean(accuracies)),
            "minimum": min(accuracies),
            "maximum": max(accuracies),
        },
        "current_cv_internal_test_group_overlap": sorted(set(ids) & set(split["test"])),
        "historical_internal_test_labels_exposed": True,
        "internal_test_split_is_sealed": False,
        "prior_v5_diagnostic_group_count": 56,
        "prior_v5_repeats": 5,
        "prior_v5_prediction_count": 280,
        "selection_role": "DEVELOPMENT_DIAGNOSTIC_ONLY_NOT_NEW_INDEPENDENT_ARTIFACTS",
        "external_validity_status": "UNESTABLISHED",
    }


def composition_sensitivity(
    metadata: list[dict[str, Any]],
    known: list[dict[str, Any]],
    candidate_id: str,
    seed: int,
    reference: dict[str, Any],
) -> dict[str, Any]:
    ids, matrix, _ = artifact_level_dataset(metadata, known, candidate_id)
    records = []
    for offset in [101, 202]:
        altered, payload = perturb_known(known, seed + offset)
        alternative_ids, alternative_matrix, _ = artifact_level_dataset(
            metadata, altered, candidate_id
        )
        if ids != alternative_ids:
            raise ValueError("PERTURBATION_CHANGED_GROUP_IDENTITY")
        loss, _ = validation_score(metadata, altered, candidate_id, seed)
        clusters = cluster_subtypes(metadata, altered, candidate_id, seed, bootstrap=False)
        ari = {}
        for glass_type in sorted(TYPE_TO_INT):
            if (
                "artifact_assignments" not in reference[glass_type]
                or "artifact_assignments" not in clusters[glass_type]
            ):
                continue
            base = reference[glass_type]["artifact_assignments"]
            other = clusters[glass_type]["artifact_assignments"]
            common = sorted(set(base) & set(other))
            ari[glass_type] = float(
                adjusted_rand_score([base[i] for i in common], [other[i] for i in common])
            )
        records.append(
            {
                "perturbation_seed": seed + offset,
                "simplex_noise_standard_deviation": 0.01,
                "perturbed_compositions_percent": payload,
                "artifact_order": ids,
                "feature_matrix_change_frobenius": float(
                    np.linalg.norm(alternative_matrix - matrix)
                ),
                "validation_composite_loss": loss,
                "within_type_adjusted_rand_indices": ari,
            }
        )
    return {
        "design": "TWO_SEEDED_ONE_PERCENTAGE_POINT_ADDITIVE_COMPOSITION_PERTURBATIONS_THEN_CLOSURE",
        "candidate_feature_transform_recomputed": True,
        "replicates": records,
    }


def predict_unknown(
    metadata: list[dict[str, Any]],
    known: list[dict[str, Any]],
    unknown: list[dict[str, Any]],
    candidate_id: str,
    seed: int,
) -> list[dict[str, Any]]:
    ids, matrix, labels = artifact_level_dataset(metadata, known, candidate_id)
    split = deterministic_split(metadata)
    allowed = set(split["train"] + split["validation"])
    positions = [i for i, artifact in enumerate(ids) if artifact in allowed]
    fitting = matrix[positions]
    fitting_labels = labels[positions]
    model = fit_type_classifier(candidate_id, fitting, fitting_labels, seed)
    thresholds = {}
    for label in [0, 1]:
        vectors = fitting[fitting_labels == label]
        distances = np.linalg.norm(vectors[:, None, :] - vectors[None, :, :], axis=2)
        np.fill_diagonal(distances, np.inf)
        thresholds[label] = float(np.quantile(distances.min(axis=1), 0.95))
    rng = np.random.default_rng(seed)
    output = []
    for row in unknown:
        feature = transform(row["composition"], candidate_id)
        score = float(model.predict_proba(feature.reshape(1, -1))[0, 1])
        label = int(score >= 0.5)
        altered_scores = []
        changes = []
        for _ in range(UNKNOWN_PERTURBATIONS):
            altered = close_composition(
                np.maximum(
                    close_composition(row["composition"]) + rng.normal(0, 0.01, len(COMPONENTS)),
                    1e-8,
                )
            )
            altered_feature = transform(altered, candidate_id)
            altered_scores.append(float(model.predict_proba(altered_feature.reshape(1, -1))[0, 1]))
            changes.append(float(np.linalg.norm(altered_feature - feature)))
        distance = float(np.linalg.norm(fitting[fitting_labels == label] - feature, axis=1).min())
        flips = sum(int(value >= 0.5) != label for value in altered_scores)
        output.append(
            {
                "sample_id": row["sample_id"],
                "input_total_percent": row["total"],
                "input_valid": row["valid"],
                "predicted_type": INT_TO_TYPE[label],
                "decision_score_lead_barium": score,
                "score_semantics": (
                    "UNCALIBRATED_DECISION_SCORE_OR_NEIGHBOR_VOTE_NOT_CALIBRATED_PROBABILITY"
                ),
                "training_artifact_ids": [ids[i] for i in positions],
                "nearest_predicted_class_training_distance": distance,
                "training_class_distance_threshold_95pct": thresholds[label],
                "within_training_distance_domain": distance <= thresholds[label],
                "domain_definition": (
                    "95_PERCENTILE_WITHIN_CLASS_LEAVE_ONE_ARTIFACT_OUT_NEAREST_DISTANCE"
                ),
                "perturbation_replicates": UNKNOWN_PERTURBATIONS,
                "perturbation_flip_count": flips,
                "perturbation_flip_rate": flips / UNKNOWN_PERTURBATIONS,
                "perturbation_decision_scores": altered_scores,
                "perturbation_feature_change_norms": changes,
                "perturbation_score_range": [
                    float(np.quantile(altered_scores, 0.025)),
                    float(np.quantile(altered_scores, 0.975)),
                ],
                "unknown_true_label_available": False,
                "unknown_accuracy_established": False,
            }
        )
    return output


def correlation_matrix(values: np.ndarray) -> np.ndarray:
    ranks = np.column_stack(
        [rankdata(values[:, i], method="average") for i in range(values.shape[1])]
    )
    centered = ranks - ranks.mean(axis=0)
    norm = np.linalg.norm(centered, axis=0)
    denominator = norm[:, None] * norm[None, :]
    return np.divide(
        centered.T @ centered,
        denominator,
        out=np.full(denominator.shape, np.nan),
        where=denominator > 1e-12,
    )


def nullable_matrix(matrix: np.ndarray) -> list[list[float | None]]:
    return [[float(v) if np.isfinite(v) else None for v in row] for row in matrix]


def association_matrices(known: list[dict[str, Any]], seed: int) -> dict[str, Any]:
    grouped = defaultdict(lambda: defaultdict(list))
    for row in known:
        if row["valid"]:
            grouped[row["glass_type"]][row["artifact_id"]].append(clr(row["composition"]))
    values = {
        typ: np.vstack([np.mean(points, axis=0) for _, points in sorted(group.items())])
        for typ, group in grouped.items()
    }
    matrices = {typ: correlation_matrix(matrix) for typ, matrix in values.items()}
    output = {
        typ: {
            "artifact_ids": sorted(grouped[typ]),
            "artifact_n": len(matrix),
            "components": list(COMPONENTS),
            "spearman_clr_matrix": nullable_matrix(matrices[typ]),
            "undefined_entry_count": int(np.sum(~np.isfinite(matrices[typ]))),
            "weighting": "EQUAL_ARTIFACT;SAMPLING_POINT_CLR_VECTORS_AVERAGED_FIRST",
        }
        for typ, matrix in values.items()
    }
    first, second = sorted(TYPE_TO_INT)
    difference = matrices[first] - matrices[second]
    rng = np.random.default_rng(seed)
    bootstrap = []
    for _ in range(ASSOCIATION_BOOTSTRAPS):
        a = values[first][rng.choice(len(values[first]), len(values[first]), replace=True)]
        b = values[second][rng.choice(len(values[second]), len(values[second]), replace=True)]
        bootstrap.append(correlation_matrix(a) - correlation_matrix(b))
    boot = np.asarray(bootstrap)
    pairs = []
    for i in range(len(COMPONENTS)):
        for j in range(i + 1, len(COMPONENTS)):
            observed = float(difference[i, j]) if np.isfinite(difference[i, j]) else None
            samples = boot[:, i, j]
            samples = samples[np.isfinite(samples)]
            pairs.append(
                {
                    "left": COMPONENTS[i],
                    "right": COMPONENTS[j],
                    "spearman_difference": observed,
                    "bootstrap_finite_replicates": int(len(samples)),
                    "bootstrap_percentile_range": None
                    if not len(samples)
                    else [float(np.quantile(samples, 0.025)), float(np.quantile(samples, 0.975))],
                }
            )
    finite = difference[np.isfinite(difference)]
    output["between_type_difference"] = {
        "type_order": [first, second],
        "difference_matrix": nullable_matrix(difference),
        "frobenius_norm_defined_entries": float(np.linalg.norm(finite)) if len(finite) else None,
        "pair_difference_ranges": pairs,
        "bootstrap_replicates": ASSOCIATION_BOOTSTRAPS,
        "bootstrap_seed": seed,
        "guard": (
            "DESCRIPTIVE;30_BOOTSTRAP_EXPLORATORY_RANGES;NO_CAUSAL_DIRECTION;NO_MUL"
            "TIPLE_COMPARISON_ADJUSTMENT;NO_COVERAGE_GUARANTEE"
        ),
    }
    return output


def requirement_claims(output_path: str) -> dict[str, dict[str, Any]]:
    statements = {
        "REQ-VALIDITY": "Original composition totals and invalid sample exclusions are audited.",
        "REQ-COMPOSITION": (
            "Valid composition vectors close to one under the declared nondetection assumption."
        ),
        "REQ-1A": (
            "Weathering associations use bounded 999-shuffle permutation checks "
            "with sparse-table limitations."
        ),
        "REQ-1B": (
            "Type and weathering centers weight each artifact equally within "
            "strata; associations are descriptive."
        ),
        "REQ-1C": (
            "Artifact-level pre-weathering backcasts assume type exchangeability; "
            "sparse pairs do not identify a causal effect."
        ),
        "REQ-2A": (
            "Two-repeat CV covers only effective train-plus-validation artifacts; "
            "previously exposed Development labels do not establish external "
            "performance."
        ),
        "REQ-2B": (
            "Exploratory subtypes include chemical centers and artifact counts; "
            "silhouette does not prove scientific validity."
        ),
        "REQ-2C": (
            "Subtype stability uses bounded artifact resampling and actual "
            "composition perturbations with nonzero feature changes."
        ),
        "REQ-3A": (
            "Eight unknown samples receive predictions; their unavailable labels "
            "prevent a claim of unknown-sample accuracy."
        ),
        "REQ-3B": (
            "Unknown predictions include training-distance domain checks and "
            "perturbation flips; model scores are uncalibrated."
        ),
        "REQ-4A": (
            "Within-type chemical associations weight artifacts equally and "
            "preserve undefined correlations as null."
        ),
        "REQ-4B": (
            "Between-type association differences have exploratory 30-bootstrap "
            "ranges without multiplicity or causal guarantees."
        ),
        "REQ-EVIDENCE": (
            "Development outputs distinguish measured diagnostics, conditional "
            "analyses and unresolved scientific requirements."
        ),
    }
    return {
        key: {
            "claim_id": f"CLAIM-C2022-V6-{i:02d}",
            "claim_text": value,
            "evidence_artifact_ids": [output_path],
        }
        for i, (key, value) in enumerate(statements.items(), 1)
    }


def scientific_evidence(payload: dict[str, Any], assumption_hash: str) -> dict[str, Any]:
    q1 = payload["question_1"]
    q2 = payload["question_2"]
    q3 = payload["question_3"]
    q4 = payload["question_4"]
    cv = payload["scientific_diagnostics"]["grouped_repeated_cv"]
    paired = sum(
        r.get("paired_artifact_n", 0) for r in q1["type_conditioned_composition_effects"].values()
    )
    metrics = {
        "REQ-VALIDITY": {
            "DATA.known_rows": payload["data_scope"]["known_rows"],
            "DATA.valid_known_rows": payload["data_scope"]["valid_known_rows"],
        },
        "REQ-COMPOSITION": {
            "DATA.maximum_closed_sum_error": payload["data_scope"]["maximum_closed_sum_error"]
        },
        "REQ-1A": {
            "Q1.association_table_count": len(q1["weathering_associations"]),
            "Q1.total_permutation_replicates": sum(
                r["permutation_replicates"] for r in q1["weathering_associations"]
            ),
        },
        "REQ-1B": {
            "Q1.type_effect_count": len(q1["type_conditioned_composition_effects"]),
            "Q1.paired_artifact_count": paired,
        },
        "REQ-1C": {
            "Q1.backcast_artifact_count": len(q1["preweathering_backcasts"]),
            "Q1.paired_artifact_count": paired,
        },
        "REQ-2A": {
            "Q2.cv_artifact_count": cv["effective_group_count"],
            "Q2.cv_prediction_count": cv["held_out_group_predictions"],
            "Q2.cv_internal_test_overlap_count": len(cv["current_cv_internal_test_group_overlap"]),
        },
        "REQ-2B": {
            "Q2.subtype_assigned_artifact_count": sum(
                len(r.get("artifact_assignments", {})) for r in q2["within_type_subtypes"].values()
            ),
            "Q2.subtype_cluster_count": sum(
                len(r.get("clusters", [])) for r in q2["within_type_subtypes"].values()
            ),
        },
        "REQ-2C": {
            "Q2.minimum_perturbation_feature_change": min(
                r["feature_matrix_change_frobenius"]
                for r in q2["subtype_sensitivity"]["replicates"]
            ),
            "Q2.subtype_bootstrap_requested_per_type": SUBTYPE_BOOTSTRAPS,
        },
        "REQ-3A": {
            "Q3.unknown_prediction_count": len(q3["unknown_type_predictions"]),
            "Q3.unknown_available_label_count": 0,
        },
        "REQ-3B": {
            "Q3.mean_unknown_perturbation_flip_rate": float(
                np.mean([r["perturbation_flip_rate"] for r in q3["unknown_type_predictions"]])
            ),
            "Q3.unknown_in_training_domain_count": sum(
                r["within_training_distance_domain"] for r in q3["unknown_type_predictions"]
            ),
        },
        "REQ-4A": {
            "Q4.association_artifact_count": sum(q4[t]["artifact_n"] for t in TYPE_TO_INT),
            "Q4.undefined_correlation_entry_count": sum(
                q4[t]["undefined_entry_count"] for t in TYPE_TO_INT
            ),
        },
        "REQ-4B": {
            "Q4.association_bootstrap_replicates": q4["between_type_difference"][
                "bootstrap_replicates"
            ],
            "Q4.component_pair_count": len(q4["between_type_difference"]["pair_difference_ranges"]),
        },
        "REQ-EVIDENCE": {"DATA.scientific_requirement_count": 13},
    }
    methods = {
        "REQ-VALIDITY": "DESCRIPTIVE_STATISTIC",
        "REQ-COMPOSITION": "DESCRIPTIVE_STATISTIC",
        "REQ-1A": "DESCRIPTIVE_STATISTIC",
        "REQ-1B": "DESCRIPTIVE_STATISTIC",
        "REQ-1C": "CONDITIONAL_SIMULATION",
        "REQ-2A": "DEVELOPMENT_DIAGNOSTIC",
        "REQ-2B": "DEVELOPMENT_DIAGNOSTIC",
        "REQ-2C": "DEVELOPMENT_DIAGNOSTIC",
        "REQ-3A": "PREDICTION",
        "REQ-3B": "DEVELOPMENT_DIAGNOSTIC",
        "REQ-4A": "DESCRIPTIVE_STATISTIC",
        "REQ-4B": "DESCRIPTIVE_STATISTIC",
        "REQ-EVIDENCE": "DEVELOPMENT_DIAGNOSTIC",
    }
    output = {}
    for req, values in metrics.items():
        entities = ["UNKNOWN_8_SAMPLES"] if req.startswith("REQ-3") else ["KNOWN_58_ARTIFACTS"]
        scope = {
            "fields": [
                "artifact_id",
                "glass_type",
                "pattern",
                "color",
                "weathering",
                "composition_14",
            ],
            "time": ["ARCHAEOLOGICAL_COLLECTION_UNDATED"],
            "entities": entities,
        }
        output[req] = {
            "generation_method": methods[req],
            "source_ids": ["SRC-2022-GLASS-WORKBOOK"],
            "scope": scope,
            "conditional_scope": dict(scope),
            "metric_values": values,
            "assumption_artifact_sha256": assumption_hash,
            "status": "INSUFFICIENT" if req in ["REQ-3A", "REQ-EVIDENCE"] else "SUPPORTED",
            "scope_limitation": (
                "Previously exposed Development collection; conditional backcast, "
                "exploratory statistics and uncalibrated predictions do not establish "
                "unknown accuracy or causality."
            ),
        }
    return output


def solve(case_root: Path, candidate_id: str, seed: int) -> dict[str, Any]:
    assumption_hash = hashlib.sha256(
        (case_root / "models/assumptions_and_symbols.json").read_bytes()
    ).hexdigest()
    metadata, known, unknown = load_official_workbook(case_root)
    loss, validation = validation_score(metadata, known, candidate_id, seed)
    cv = grouped_repeated_cv_diagnostic(metadata, known, candidate_id, seed)
    effects, backcasts = weathering_analysis(known, candidate_id)
    subtypes = cluster_subtypes(metadata, known, candidate_id, seed)
    sensitivity = composition_sensitivity(metadata, known, candidate_id, seed, subtypes)
    predictions = predict_unknown(metadata, known, unknown, candidate_id, seed)
    associations = association_matrices(known, seed)
    payload = {
        "schema_version": "1.0.0",
        "pipeline_version": PIPELINE_VERSION,
        "candidate_id": candidate_id,
        "seed": seed,
        "status": "SUCCESS",
        "answer_access_status": "UNLOCKED_AFTER_FIRST_RUN",
        "test_accessed": False,
        "final_evaluation_accessed": False,
        "development_contamination": {
            "historical_internal_test_labels_exposed": True,
            "prior_v5_cv_used_all_56_valid_artifacts_including_12_internal_test_groups": True,
            "current_case_is_blind": False,
        },
        "assumption_artifact_sha256": assumption_hash,
        "input_hashes": {"raw/case_files/附件.xlsx": OFFICIAL_WORKBOOK_SHA256},
        "validation_metrics": {"validation_composite_loss": loss},
        "validation_detail": validation,
        "scientific_diagnostics": {"grouped_repeated_cv": cv},
        "data_scope": {
            "known_rows": len(known),
            "known_metadata_artifact_count": len(metadata),
            "valid_known_rows": sum(r["valid"] for r in known),
            "valid_known_artifact_count": len(artifact_compositions(known)),
            "invalid_known_sample_ids": [r["sample_id"] for r in known if not r["valid"]],
            "unknown_rows": len(unknown),
            "unknown_valid_rows": sum(r["valid"] for r in unknown),
            "maximum_closed_sum_error": max(
                abs(float(close_composition(r["composition"]).sum()) - 1)
                for r in known + unknown
                if r["valid"]
            ),
            "blank_cell_policy": "NONDETECTION_ZERO_REPLACEMENT_IS_AN_ASSUMPTION",
            "artifact_grouping": True,
        },
        "question_1": {
            "weathering_associations": [
                contingency_result(metadata, field, seed + i)
                for i, field in enumerate(["glass_type", "pattern", "color"])
            ],
            "type_conditioned_composition_effects": effects,
            "preweathering_backcasts": backcasts,
        },
        "question_2": {
            "type_classifier": {
                "candidate_id": candidate_id,
                "validation": validation,
                "decision_threshold": 0.5,
            },
            "within_type_subtypes": subtypes,
            "subtype_sensitivity": sensitivity,
        },
        "question_3": {
            "unknown_type_predictions": predictions,
            "answer_labels_available": False,
            "sensitivity_design": (
                "50_SEEDED_ONE_PERCENTAGE_POINT_SIMPLEX_PERTURBATIONS_PER_UNKNOWN"
            ),
        },
        "question_4": associations,
        "final_metrics": {"validation_composite_loss": loss},
        "figure_ready_data": [{"figure_id": "UNKNOWN-TYPE-PREDICTIONS", "rows": predictions}],
        "uncertainty": {
            "classification": (
                "finite-group Development diagnosis; no unknown true labels or calibration evidence"
            ),
            "backcast": "equal-artifact type exchangeability; only sparse within-artifact pairing",
            "association": "30 bootstrap exploratory ranges; no multiplicity adjustment",
        },
        "robustness_evidence": {
            "metric": "validation_composite_loss",
            "metric_direction": "MIN",
            "perturbations": [
                {
                    "perturbation_id": (
                        f"COMPOSITION-ONE-PERCENTAGE-POINT-SEED-{r['perturbation_seed']}"
                    ),
                    "metric": "validation_composite_loss",
                    "result": r["validation_composite_loss"],
                    "evidence": "DETERMINISTIC_RECOMPUTATION_FROM_BOUND_INPUTS",
                }
                for r in sensitivity["replicates"]
            ],
            "failure_cases": [
                "Repeated predictions are not independent artifacts.",
                "Sparse paired artifacts cannot identify causal weathering.",
                (
                    "Training-distance applicability does not establish unknown "
                    "classification accuracy."
                ),
            ],
        },
        "scientific_coverage_status": "PARTIAL_SCIENTIFIC_COVERAGE",
        "whole_problem_scientifically_complete": False,
        "limitations": [
            (
                "Case is Development and historically contaminated; old internal test "
                "labels are exposed."
            ),
            "Blank/nondetection handling is an assumption.",
            (
                "CV, subtype geometry and bootstrap stability do not establish external"
                " scientific truth."
            ),
            "Unknown decision scores and KNN votes are not calibrated probabilities.",
            "Permutation and bootstrap outputs are exploratory with no multiplicity adjustment.",
            (
                "Backcasts are conditional and not causal estimates; unknown labels "
                "remain unavailable."
            ),
        ],
    }
    payload["scientific_evidence"] = scientific_evidence(payload, assumption_hash)
    for evidence in payload["scientific_evidence"].values():
        payload["final_metrics"].update(evidence["metric_values"])
    return payload


def main() -> int:
    args = parse_args()
    root = Path(args.case_root).resolve()
    output = safe_path(root, args.output)
    result = solve(root, args.candidate_id, args.seed)
    result["requirement_claims"] = requirement_claims(str(Path(args.output)))
    result["claim_scope"] = result["requirement_claims"]["REQ-VALIDITY"]["claim_text"]
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(json_safe(result), ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False)
        + "\n",
        encoding="utf-8",
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
