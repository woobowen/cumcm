"""Verify all dual/primal inequalities using exact fractions, without any solver.

Raw parameter checks independently read only W001-W168. The exact certificate is
for the explicitly serialized decimal-parameter LP, not unknown enterprise capacity.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from fractions import Fraction
from pathlib import Path

from openpyxl import load_workbook

RAW = {
    "raw/case_files/附件1 近5年402家供应商的相关数据.xlsx": (
        "1b93a7598b8f0489e3c054439967bca654a83654837908c833752617cf950e1b"
    ),
    "raw/case_files/附件2 近5年8家转运商的相关数据.xlsx": (
        "29e1499b133aa88d765902081bfb827f78ffb4c091f5f75341fc37d844344685"
    ),
}


def rational(value) -> Fraction:
    if isinstance(value, bool):
        raise ValueError("BOOLEAN_NOT_NUMERIC")
    return Fraction(str(value))


def verify_exact(output: dict, certificate: dict) -> dict:
    p = output["parameters"]
    suppliers = sorted(p["supplier_types"])
    carriers = sorted(p["carrier_loss_rates"])
    u = {k: rational(v) for k, v in certificate["supplier_multipliers"].items()}
    v = {k: rational(x) for k, x in certificate["carrier_multipliers"].items()}
    if set(u) != set(suppliers) or set(v) != set(carriers):
        raise ValueError("MULTIPLIER_COVERAGE_MISMATCH")
    if min([*u.values(), *v.values()]) < 0:
        raise ValueError("NEGATIVE_DUAL_MULTIPLIER")
    rows = certificate["primal_flows"]
    flow = {(r["supplier_id"], r["carrier_id"]): rational(r["volume_m3"]) for r in rows}
    if len(flow) != len(rows) or any(s not in suppliers or t not in carriers for s, t in flow):
        raise ValueError("INVALID_PRIMAL_IDENTITIES")
    if any(x < 0 for x in flow.values()):
        raise ValueError("NEGATIVE_PRIMAL_FLOW")
    scale = rational(certificate["primal_scale"])
    original_rows = output["question_4"]["weekly_plan_repeated_for_24_weeks"]["transport"]
    original = {
        (r["supplier_id"], r["carrier_id"]): rational(r["volume_m3"]) for r in original_rows
    }
    if len(original) != len(original_rows) or not 0 < scale <= 1:
        raise ValueError("ORIGINAL_FLOW_OR_SCALE_INVALID")
    if flow != {key: value * scale for key, value in original.items()}:
        raise ValueError("PRIMAL_NOT_SCALED_ORIGINAL_PLAN")
    supplier_residuals = {}
    carrier_residuals = {}
    dual_residuals = {}
    lower = Fraction(0)
    for s in suppliers:
        cap = rational(p["supplier_supply_capacity_m3"][s])
        total = sum((flow.get((s, t), Fraction(0)) for t in carriers), Fraction(0))
        supplier_residuals[s] = str(cap - total)
        if total > cap:
            raise ValueError(f"SUPPLIER_CAPACITY_EXCEEDED:{s}")
        for t in carriers:
            c = (1 - rational(p["carrier_loss_rates"][t])) / rational(
                p["consumption"][p["supplier_types"][s]]
            )
            slack = u[s] + v[t] - c
            dual_residuals[f"{s}:{t}"] = str(slack)
            if slack < 0:
                raise ValueError(f"DUAL_INEQUALITY_VIOLATED:{s}:{t}")
            lower += c * flow.get((s, t), Fraction(0))
    cap = rational(p["carrier_capacity_m3"])
    for t in carriers:
        total = sum((flow.get((s, t), Fraction(0)) for s in suppliers), Fraction(0))
        carrier_residuals[t] = str(cap - total)
        if total > cap:
            raise ValueError(f"CARRIER_CAPACITY_EXCEEDED:{t}")
    upper = sum(
        (rational(p["supplier_supply_capacity_m3"][s]) * u[s] for s in suppliers), Fraction(0)
    ) + cap * sum(v.values(), Fraction(0))
    gap = upper - lower
    if not 0 <= gap <= Fraction(1, 1000):
        raise ValueError(f"ABSOLUTE_OPTIMALITY_GAP_EXCEEDS_PREREGISTERED_0_001:{float(gap)}")
    return {
        "status": "PASS",
        "proof_kind": "EXACT_RATIONAL_WEAK_DUALITY_INTERVAL",
        "scope": "FIXED_DECIMAL_PARAMETER_STATIONARY_TRANSPORT_LP",
        "lower_bound_exact": str(lower),
        "upper_bound_exact": str(upper),
        "absolute_gap_exact": str(gap),
        "lower_bound_m3": float(lower),
        "upper_bound_m3": float(upper),
        "absolute_gap_m3": float(gap),
        "relative_gap": float(gap / upper),
        "original_reported_attainment_m3": output["question_4"]["weekly_capacity_product_m3"],
        "primal_scale": str(scale),
        "supplier_slacks_exact": supplier_residuals,
        "carrier_slacks_exact": carrier_residuals,
        "dual_slacks_exact": dual_residuals,
        "strict_inequality_verification_tolerance": 0,
        "gap_acceptance_tolerance_m3": 0.001,
        "independent_solver_invocations": 0,
        "whole_problem_scientifically_complete": False,
    }


def verify_raw(root: Path, output: dict) -> dict:
    for name, expected in RAW.items():
        if hashlib.sha256((root / name).read_bytes()).hexdigest() != expected:
            raise ValueError("RAW_WORKBOOK_HASH_MISMATCH")
    names = list(RAW)
    book = load_workbook(root / names[0], read_only=True, data_only=True)
    orders = list(book["企业的订货量（m³）"].iter_rows(max_col=170, values_only=True))
    supply = list(book["供应商的供货量（m³）"].iter_rows(max_col=170, values_only=True))
    book.close()
    book = load_workbook(root / names[1], read_only=True, data_only=True)
    losses = list(book["运输损耗率（%）"].iter_rows(max_col=169, values_only=True))
    book.close()
    weeks = [f"W{i:03d}" for i in range(1, 169)]
    if list(orders[0][2:]) != weeks or list(supply[0][2:]) != weeks or list(losses[0][1:]) != weeks:
        raise ValueError("RAW_TRAINING_HEADERS_INVALID")
    orders, supply, losses = [
        [r for r in rows[1:] if r[0] is not None] for rows in [orders, supply, losses]
    ]
    if (len(orders), len(supply), len(losses)) != (402, 402, 8):
        raise ValueError("RAW_ENTITY_COUNTS_INVALID")
    p = output["parameters"]
    residuals = {}
    for ordered, supplied in zip(orders, supply, strict=True):
        sid, typ = ordered[:2]
        if ordered[:2] != supplied[:2] or p["supplier_types"][sid] != typ:
            raise ValueError("RAW_SUPPLIER_IDENTITY_MISMATCH")
        pairs = [
            (float(o), float(x)) for o, x in zip(ordered[2:], supplied[2:], strict=True) if o > 0
        ]
        if pairs:
            positive = sorted(o for o, _ in pairs)
            index = (len(positive) - 1) * 0.75
            lo = math.floor(index)
            hi = math.ceil(index)
            q = positive[lo] + (index - lo) * (positive[hi] - positive[lo])
            ratio = min(1.5, math.fsum(min(2.0, max(0.0, x / o)) for o, x in pairs) / len(pairs))
        else:
            q, ratio = 0.0, 0.0
        for label, expected in [
            ("supplier_order_capacity_m3", q),
            ("supplier_delivery_ratios", ratio),
            ("supplier_supply_capacity_m3", q * ratio),
        ]:
            residuals[f"{sid}:{label}"] = float(p[label][sid]) - expected
    for row in losses:
        values = [float(x) / 100 for x in row[1:] if x > 0]
        residuals[f"{row[0]}:loss"] = float(p["carrier_loss_rates"][row[0]]) - math.fsum(
            values
        ) / len(values)
    if p["consumption"] != {"A": 0.6, "B": 0.66, "C": 0.72} or p["carrier_capacity_m3"] != 6000:
        raise ValueError("PHYSICAL_CONSTANTS_MISMATCH")
    maximum = max(abs(x) for x in residuals.values())
    if maximum > 1e-9:
        raise ValueError("RAW_PARAMETER_DERIVATION_MISMATCH")
    return {
        "input_hashes": RAW,
        "numeric_columns_read": "W001-W168",
        "parameter_residuals": residuals,
        "maximum_absolute_parameter_residual": maximum,
        "parameter_comparison_tolerance": 1e-9,
        "exact_certificate_targets_serialized_parameters": True,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-output", type=Path, required=True)
    parser.add_argument("--certificate", type=Path, required=True)
    parser.add_argument("--case-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        raise FileExistsError(args.output)
    raw = args.source_output.read_bytes()
    output, certificate = json.loads(raw), json.loads(args.certificate.read_text())
    if certificate["source_output_sha256"] != hashlib.sha256(raw).hexdigest():
        raise ValueError("SOURCE_OUTPUT_HASH_MISMATCH")
    result = verify_exact(output, certificate)
    result["raw_parameter_check"] = verify_raw(args.case_root, output)
    result["certificate_sha256"] = hashlib.sha256(args.certificate.read_bytes()).hexdigest()
    result["source_output_sha256"] = hashlib.sha256(raw).hexdigest()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(
        json.dumps(
            {
                k: result[k]
                for k in [
                    "status",
                    "lower_bound_m3",
                    "upper_bound_m3",
                    "absolute_gap_m3",
                    "relative_gap",
                ]
            }
        )
    )


if __name__ == "__main__":
    main()
